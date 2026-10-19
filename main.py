from __future__ import annotations

import calendar
from copy import copy
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
import hashlib
from io import BytesIO
import logging
import os
from pathlib import Path
import re
import secrets
from itertools import chain
from typing import Any
from zipfile import BadZipFile, ZipFile
from urllib.parse import quote, unquote
from uuid import uuid4

import pandas as pd
from dotenv import load_dotenv
from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.exception_handlers import http_exception_handler
from fastapi.responses import JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict
from sqlalchemy import URL, create_engine, text
from sqlalchemy.engine import Engine
from sqlalchemy.exc import IntegrityError, SQLAlchemyError


BASE_DIR = Path(__file__).resolve().parent
PUBLIC_DIR = BASE_DIR / "public"
MAX_UPLOAD_BYTES = 200 * 1024 * 1024
ATTENDANCE_PREVIEW_ROWS = 50
SESSION_LIFETIME = timedelta(hours=4)

EXPECTED_COLUMNS = [
    "Legend", "Team", "Zone", "Region", "Dept. Territory Code",
    "Position_Code", "SAS_ID", "SPO Name", "Base Town",
    "Emp. Type", "Emp. Code", "Payroll Status", "D.O.J",
]
EDITABLE_COLUMNS = [
    "Zone", "Region", "Team", "SPO Name", "Base Town",
    "Emp. Type", "D.O.J", "Payroll Status",
]

DATABASE_FIELDS = {
    "Legend": "designation_name",
    "Team": "team_name",
    "Zone": "sub_zone",
    "Region": "region_name",
    "Dept. Territory Code": "dept_territory_code",
    "Position_Code": "position_code",
    "SAS_ID": "sas_id",
    "SPO Name": "employe_name",
    "Base Town": "base_town",
    "Emp. Type": "nature_type",
    "Emp. Code": "e_code",
    "Payroll Status": "payroll_type",
    "D.O.J": "doj",
}

HEADER_ALIASES = {
    "Legend": ("Legend", "Designation_Name"),
    "Team": ("Team", "Team_Name"),
    "Zone": ("Zone", "Zone_Name", "Sub_Zone"),
    "Region": ("Region", "Region_Name"),
    "Dept. Territory Code": ("Dept. Territory Code", "Dept_Territory_Code"),
    "Position_Code": ("Position_Code", "Employee_Position_Code"),
    "SAS_ID": ("SAS_ID",),
    "SPO Name": ("SPO Name", "Emp_Name", "Employe_Name", "Employee_Name"),
    "Base Town": ("Base Town", "Base_Town"),
    "Emp. Type": ("Emp. Type", "Nature Type", "Nature_Type"),
    "Emp. Code": ("Emp. Code", "E-CODE", "E_Code"),
    "Payroll Status": ("Payroll Status", "Payroll Type", "Payroll_Type"),
    "D.O.J": ("D.O.J", "DOJ", "D_O_J"),
}

load_dotenv(BASE_DIR / ".env", override=True)
logger = logging.getLogger("uvicorn.error")
app = FastAPI(title="IKON Call", version="2.0.0")


class UploadSession(BaseModel):
    content: bytes
    filename: str
    created_at: datetime
    excel_content: bytes | None = None
    csv_content: bytes | None = None
    source_content: bytes | None = None
    source_filename: str | None = None


class ExportRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    sheetName: str
    records: list[dict[str, Any]]
    rowSources: list[int | None]
    originalRows: list[int]
    columnNumbers: list[int]
    insertAtRow: int


upload_sessions: dict[str, UploadSession] = {}
attendance_sessions: dict[str, UploadSession] = {}
fmcg_attendance_sessions: dict[str, UploadSession] = {}
processed_sessions: dict[str, UploadSession] = {}
database_engine: Engine | None = None


def normalized_name(value: object) -> str:
    return re.sub(r"\s+", " ", display_value(value)).strip().casefold()


def header_key(value: object) -> str:
    """Compare Excel headers independent of spaces, underscores, or punctuation."""
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


column_key = header_key


def display_value(value: object) -> str:
    if value is None or (not isinstance(value, (date, datetime)) and pd.isna(value)):
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def normalize_activity_employee_id(value: object) -> str:
    """Normalize FF.Id/SAS_ID values so source files with different padding match."""
    value_text = display_value(value).strip()
    if not value_text:
        return ""
    if re.fullmatch(r"\d+\.0+", value_text):
        value_text = value_text.split(".", 1)[0]
    normalized = value_text.lstrip("0")
    return normalized or "0"


def clean_filename(filename: str | None) -> str:
    name = Path(filename or "updated-data.xlsx").name.replace("\r", "").replace("\n", "")
    return name if name.casefold().endswith(".xlsx") else f"{name}.xlsx"


def cleanup_sessions() -> None:
    cutoff = datetime.now(timezone.utc) - SESSION_LIFETIME
    expired = [session_id for session_id, session in upload_sessions.items() if session.created_at < cutoff]
    for session_id in expired:
        upload_sessions.pop(session_id, None)
    for session_id, session in list(attendance_sessions.items()):
        if session.created_at < cutoff:
            attendance_sessions.pop(session_id, None)
    for session_id, session in list(fmcg_attendance_sessions.items()):
        if session.created_at < cutoff:
            fmcg_attendance_sessions.pop(session_id, None)
    for session_id, session in list(processed_sessions.items()):
        if session.created_at < cutoff:
            processed_sessions.pop(session_id, None)


def get_upload_session(session_id: str) -> UploadSession:
    cleanup_sessions()
    session = upload_sessions.get(session_id)
    if session is None:
        raise HTTPException(status_code=404, detail="This upload has expired. Please upload the workbook again.")
    return session


def get_attendance_session(session_id: str) -> UploadSession:
    cleanup_sessions()
    session = attendance_sessions.get(session_id)
    if session is None:
        raise HTTPException(status_code=404, detail="This summary upload has expired. Please upload it again.")
    return session


def get_processed_session(session_id: str) -> UploadSession:
    cleanup_sessions()
    session = processed_sessions.get(session_id)
    if session is None:
        raise HTTPException(status_code=404, detail="This processed file has expired. Upload and process it again.")
    return session


def get_database_engine() -> Engine:
    global database_engine
    if database_engine is not None:
        return database_engine

    required = ("DB_HOST", "DB_PORT", "DB_NAME", "DB_USER", "DB_PASSWORD")
    missing = [name for name in required if not os.getenv(name)]
    if missing:
        raise ValueError(f"Missing database setting(s): {', '.join(missing)}")

    port_match = re.match(r"\d+", os.environ["DB_PORT"].strip())
    if not port_match:
        raise ValueError("DB_PORT must start with a port number.")

    database_engine = create_engine(
        URL.create(
            "postgresql+psycopg2",
            username=os.environ["DB_USER"],
            password=os.environ["DB_PASSWORD"],
            host=os.environ["DB_HOST"],
            port=int(port_match.group()),
            database=os.environ["DB_NAME"],
        ),
        pool_pre_ping=True,
        connect_args={"connect_timeout": 5},
    )
    return database_engine


def read_editable_data(content: bytes, sheet_name: str) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise HTTPException(status_code=404, detail="Worksheet not found.")

        worksheet = workbook[sheet_name]
        wanted = {
            header_key(alias): column
            for column, aliases in HEADER_ALIASES.items()
            for alias in aliases
        }
        header_row: int | None = None
        positions: dict[str, int] = {}

        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=1,
                max_row=min(worksheet.max_row, 200),
                max_col=min(worksheet.max_column, 100),
                values_only=True,
            ),
            start=1,
        ):
            found: dict[str, int] = {}
            for column_number, value in enumerate(row, start=1):
                name = wanted.get(header_key(value))
                if name and name not in found:
                    found[name] = column_number
            if all(column in found for column in EDITABLE_COLUMNS):
                header_row = row_number
                positions = found
                break

        if header_row is None:
            raise HTTPException(
                status_code=400,
                detail="Could not find all 8 editable headers in the first 200 rows of this worksheet. Check the worksheet and header names.",
            )

        records: list[dict[str, str]] = []
        row_sources: list[int] = []
        max_column = max(positions.values())
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, max_col=max_column, values_only=True),
            start=header_row + 1,
        ):
            record = {
                column: display_value(row[positions[column] - 1]) if column in positions else ""
                for column in EXPECTED_COLUMNS
            }
            if any(record.values()):
                records.append(record)
                row_sources.append(row_number)

        # pandas provides a stable, JSON-safe record shape for the API.
        frame = pd.DataFrame(records, columns=EXPECTED_COLUMNS, dtype=str).fillna("")
        insert_at_row = max([header_row, *row_sources]) + 1
        return {
            "records": frame.to_dict(orient="records"),
            "rowSources": row_sources,
            "originalRows": row_sources.copy(),
            "columnNumbers": [positions[column] for column in EDITABLE_COLUMNS],
            "insertAtRow": insert_at_row,
        }
    finally:
        workbook.close()


def worksheet_has_editable_headers(worksheet) -> bool:
    """Return whether a worksheet contains all eight editable headers."""
    # Build the matched canonical field set so different aliases count as one field.
    alias_to_field = {
        header_key(alias): column
        for column in EDITABLE_COLUMNS
        for alias in HEADER_ALIASES[column]
    }
    for row in worksheet.iter_rows(
        min_row=1,
        max_row=min(worksheet.max_row, 200),
        max_col=min(worksheet.max_column, 100),
        values_only=True,
    ):
        found = {alias_to_field.get(header_key(value)) for value in row}
        if EDITABLE_COLUMNS and all(column in found for column in EDITABLE_COLUMNS):
            return True
    return False


def matching_format_row(
    worksheet,
    original_rows: list[int],
    column_numbers: list[int],
    values: list[object],
) -> int | None:
    best_row = None
    best_score = -1
    for row_number in original_rows:
        score = 0
        for index, column_number in enumerate(column_numbers[:4]):
            existing = normalized_name(worksheet.cell(row_number, column_number).value)
            incoming = normalized_name(values[index])
            if existing and incoming and existing == incoming:
                score += 10 if index == 0 else 1
        if score > best_score:
            best_row = row_number
            best_score = score
    return best_row


def copy_row_format(worksheet, source_row: int, target_row: int) -> None:
    for column_number in range(1, worksheet.max_column + 1):
        source = worksheet.cell(source_row, column_number)
        target = worksheet.cell(target_row, column_number)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format

    source_dimension = worksheet.row_dimensions[source_row]
    target_dimension = worksheet.row_dimensions[target_row]
    target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden


def write_record_cell(cell, value: object, column_name: str) -> None:
    text_value = display_value(value).strip()
    if column_name == "D.O.J":
        try:
            cell.value = date.fromisoformat(text_value)
            cell.number_format = "yyyy-mm-dd"
            return
        except ValueError:
            pass
    cell.value = text_value
    cell.data_type = "s"


@app.exception_handler(HTTPException)
async def json_http_exception(request: Request, exception: HTTPException):
    if request.url.path.startswith("/api/"):
        return JSONResponse({"error": str(exception.detail)}, status_code=exception.status_code)
    return await http_exception_handler(request, exception)


@app.exception_handler(Exception)
async def json_exception(request: Request, exception: Exception):
    logger.exception("Unhandled application error", exc_info=exception)
    if request.url.path.startswith("/api/"):
        return JSONResponse({"error": str(exception) or "Something went wrong."}, status_code=400)
    return JSONResponse({"error": "Something went wrong."}, status_code=500)


@app.get("/api/health")
def health() -> dict[str, str]:
    return {"status": "ok", "backend": "FastAPI"}


EXPECTED_MOBILE_LIVE_CALLS_COLUMNS = [
    "Team",
    "FF.Id",
    "FF.Name",
    "Call Date",
    "Manager",
    "Base Town",
    "PMDC",
    "Doctorname",
    "CallType",
    "WW",
    "LNG",
    "LAT",
    "Total Sum of NOOFCALLS",
    "Total Sum of NOOFPRODUCTS",
    "GeoMeter Range",
]

MOBILE_LIVE_CALLS_CANONICAL_MAP = {
    column_key(col): col for col in EXPECTED_MOBILE_LIVE_CALLS_COLUMNS
}
MOBILE_LIVE_CALLS_CANONICAL_MAP.update({
    "ffid": "FF.Id",
    "ffname": "FF.Name",
    "calldate": "Call Date",
    "basetown": "Base Town",
    "doctorname": "Doctorname",
    "doctor": "Doctorname",
    "calltype": "CallType",
    "type": "CallType",
    "totalsumofnoofcalls": "Total Sum of NOOFCALLS",
    "totalsumofnoofproducts": "Total Sum of NOOFPRODUCTS",
    "geometerrange": "GeoMeter Range",
})


def find_spreadsheet_header_row(
    sample: pd.DataFrame,
    canonical_by_key: dict[str, str],
) -> tuple[int, int]:
    """Return the best header-row index and number of distinct matched fields."""
    best_row = 0
    best_matches = 0
    for row_index, row in enumerate(sample.itertuples(index=False, name=None)):
        matched_fields = {
            canonical_by_key.get(header_key(value))
            for value in row
            if display_value(value).strip()
        }
        matched_fields.discard(None)
        if len(matched_fields) > best_matches:
            best_row = row_index
            best_matches = len(matched_fields)
    return best_row, best_matches


def read_excel_table_with_detected_header(
    content: bytes,
    canonical_by_key: dict[str, str],
    required_match_count: int,
    error_message: str,
) -> pd.DataFrame:
    """Find a table header in the first 100 rows of every Excel worksheet."""
    try:
        with pd.ExcelFile(BytesIO(content)) as workbook:
            best_sheet: str | None = None
            best_header_row = 0
            best_match_count = 0

            for sheet_name in workbook.sheet_names:
                sample = pd.read_excel(
                    workbook,
                    sheet_name=sheet_name,
                    header=None,
                    dtype=object,
                    nrows=100,
                ).fillna("")
                if sample.empty:
                    continue
                header_row, match_count = find_spreadsheet_header_row(
                    sample,
                    canonical_by_key,
                )
                if match_count > best_match_count:
                    best_sheet = sheet_name
                    best_header_row = header_row
                    best_match_count = match_count

            if best_sheet is None or best_match_count < required_match_count:
                raise ValueError(error_message)

            return pd.read_excel(
                workbook,
                sheet_name=best_sheet,
                header=best_header_row,
                dtype=object,
            ).fillna("")
    except ValueError:
        raise
    except Exception as error:
        raise ValueError(error_message) from error


def clean_and_validate_mobile_live_calls(frame: pd.DataFrame) -> pd.DataFrame:
    if frame.empty:
        raise HTTPException(
            status_code=400,
            detail="Unable to process file. Please check the Mobile Live Calls Summary format and required columns.",
        )

    cleaned_cols = []
    seen: dict[str, int] = {}
    for raw_col in frame.columns:
        norm_text = re.sub(r"\s+", " ", str(raw_col or "")).strip()
        key = column_key(norm_text)
        canonical = MOBILE_LIVE_CALLS_CANONICAL_MAP.get(key, norm_text)
        count = seen.get(canonical, 0)
        seen[canonical] = count + 1
        cleaned_cols.append(canonical if count == 0 else f"{canonical} ({count})")

    frame.columns = cleaned_cols

    missing = [col for col in EXPECTED_MOBILE_LIVE_CALLS_COLUMNS if col not in frame.columns]
    if missing:
        raise HTTPException(
            status_code=400,
            detail="Unable to process file. Please check the Mobile Live Calls Summary format and required columns.",
        )
    return frame


def read_mobile_live_calls_frame(content: bytes, extension: str) -> pd.DataFrame:
    """Read CSV/Excel data, locating its header even after title rows."""
    error_message = (
        "Unable to process file. Please check the Mobile Live Calls Summary "
        "format and required columns."
    )
    if extension == ".csv":
        for encoding in ("utf-8-sig", "utf-8", "latin1"):
            try:
                sample = pd.read_csv(
                    BytesIO(content),
                    header=None,
                    dtype=object,
                    nrows=100,
                    encoding=encoding,
                ).fillna("")
                break
            except Exception:
                continue
        else:
            raise HTTPException(status_code=400, detail=error_message)

        best_row, best_matches = find_spreadsheet_header_row(
            sample,
            MOBILE_LIVE_CALLS_CANONICAL_MAP,
        )
        if best_matches < len(EXPECTED_MOBILE_LIVE_CALLS_COLUMNS):
            raise HTTPException(status_code=400, detail=error_message)

        try:
            frame = pd.read_csv(
                BytesIO(content),
                header=best_row,
                dtype=object,
                encoding=encoding,
            ).fillna("")
        except Exception:
            frame = pd.read_csv(
                BytesIO(content),
                header=best_row,
                dtype=object,
                encoding="latin1",
            ).fillna("")
        return frame

    try:
        return read_excel_table_with_detected_header(
            content,
            MOBILE_LIVE_CALLS_CANONICAL_MAP,
            len(EXPECTED_MOBILE_LIVE_CALLS_COLUMNS),
            error_message,
        )
    except ValueError as error:
        raise HTTPException(status_code=400, detail=error_message) from error


@app.post("/api/attendance/mobile-live-call-summary")
async def upload_mobile_live_call_summary(file: UploadFile = File(...)) -> dict[str, Any]:
    """Process Call Nature and validate Mobile Live Calls Summary dataset."""
    filename = file.filename or ""
    extension = Path(filename).suffix.casefold()
    if extension not in {".csv", ".xlsx", ".xls"}:
        raise HTTPException(
            status_code=400,
            detail="Unable to process file. Please check the Mobile Live Calls Summary format and required columns.",
        )

    content = await file.read(MAX_UPLOAD_BYTES + 1)
    await file.close()
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="The summary file exceeds the 200 MB upload limit.")

    try:
        source_frame = read_mobile_live_calls_frame(content, extension)
        validated_frame = clean_and_validate_mobile_live_calls(source_frame)
        processed_frame, validation = add_call_nature(validated_frame)

        # Build processed Excel workbook for download
        excel_output = BytesIO()
        with pd.ExcelWriter(excel_output, engine="openpyxl") as writer:
            processed_frame.to_excel(writer, index=False, sheet_name="Mobile Live Calls")
        excel_bytes = excel_output.getvalue()
        csv_bytes = processed_frame.to_csv(index=False).encode("utf-8-sig")

        if extension == ".csv":
            stored_content = csv_bytes
            stored_filename = f"{Path(filename).stem}-updated.csv"
        else:
            stored_content = excel_bytes
            stored_filename = f"{Path(filename).stem}-updated.xlsx"

        call_nature = {
            "validation": validation,
            "filename": f"{Path(filename).stem}-updated.xlsx",
        }
        page = read_attendance_page(stored_content, extension, 0)
        database_upload = store_mobile_live_call_details(
            processed_frame,
            content,
            filename,
        )
    except HTTPException:
        raise
    except Exception as error:
        logger.warning("Mobile Live Calls Summary processing error: %s", error)
        raise HTTPException(
            status_code=400,
            detail="Unable to process file. Please check the Mobile Live Calls Summary format and required columns.",
        ) from error

    session_id = str(uuid4())
    attendance_sessions[session_id] = UploadSession(
        content=stored_content,
        filename=stored_filename,
        created_at=datetime.now(timezone.utc),
        excel_content=excel_bytes,
        csv_content=csv_bytes,
        source_content=content,
        source_filename=Path(filename).name,
    )
    return {
        "id": session_id,
        "filename": stored_filename,
        "callNature": call_nature,
        "databaseUpload": database_upload,
        **page,
    }


CALL_ATTENDENCE_FILE_COLUMNS = {
    "Team": "team",
    "FF.Id": "ff_id",
    "FF.Name": "ff_name",
    "Manager": "manager",
    "Base Town": "base_town",
    "Call Date": "call_date",
    "CallType": "call_type",
    "Total Sum of DONE": "total_sum_of_done",
    "Total Distinct Count of FF.Id": "total_distinct_count_of_ff_id",
}
CALL_ATTENDENCE_HISTORY_TABLE = "call_attendence_upload_history"


def read_call_attendence_frame(content: bytes, extension: str) -> pd.DataFrame:
    """Read a call-attendance CSV/Excel file and locate its header row."""
    canonical_by_key = {
        header_key(column): column for column in CALL_ATTENDENCE_FILE_COLUMNS
    }
    encoding: str | None = None

    try:
        if extension == ".csv":
            sample = None
            for candidate in ("utf-8-sig", "utf-8", "latin1"):
                try:
                    sample = pd.read_csv(
                        BytesIO(content), header=None, dtype=object, nrows=100, encoding=candidate,
                    ).fillna("")
                    encoding = candidate
                    break
                except UnicodeDecodeError:
                    continue
            if sample is None or encoding is None:
                raise ValueError("The CSV encoding could not be read.")
        else:
            return read_excel_table_with_detected_header(
                content,
                canonical_by_key,
                len(CALL_ATTENDENCE_FILE_COLUMNS),
                "The file header does not match the nine required call_attendence columns.",
            )
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("The file could not be read as a CSV or Excel workbook.") from error

    if sample.empty:
        raise ValueError("The uploaded file is empty.")

    best_header_row, best_match_count = find_spreadsheet_header_row(
        sample,
        canonical_by_key,
    )

    if best_match_count != len(CALL_ATTENDENCE_FILE_COLUMNS):
        raise ValueError(
            "The file header does not match the nine required call_attendence columns."
        )

    return pd.read_csv(
        BytesIO(content),
        header=best_header_row,
        dtype=object,
        encoding=encoding,
    )


def clean_call_attendence_frame(frame: pd.DataFrame) -> pd.DataFrame:
    """Validate and order the upload columns exactly as call_attendence expects."""
    canonical_by_key = {
        header_key(column): column for column in CALL_ATTENDENCE_FILE_COLUMNS
    }
    renamed: dict[object, str] = {}
    seen: set[str] = set()
    unsupported: list[str] = []

    for raw_column in frame.columns:
        source_name = re.sub(r"\s+", " ", str(raw_column or "")).strip()
        canonical = canonical_by_key.get(header_key(source_name))
        if canonical is None:
            # Spreadsheet programs sometimes retain a completely empty formatted column.
            if frame[raw_column].map(lambda value: not _mobile_live_call_value_is_empty(value)).any():
                unsupported.append(source_name or "unnamed column")
            continue
        if canonical in seen:
            raise ValueError(f"The file contains the '{canonical}' column more than once.")
        seen.add(canonical)
        renamed[raw_column] = canonical

    missing = [column for column in CALL_ATTENDENCE_FILE_COLUMNS if column not in seen]
    if missing:
        raise ValueError(f"The file is missing required column(s): {', '.join(missing)}.")
    if unsupported:
        raise ValueError(
            f"The file contains column(s) not present in call_attendence: {', '.join(unsupported)}."
        )

    cleaned = frame[list(renamed)].rename(columns=renamed)
    cleaned = cleaned.dropna(how="all")
    nonblank_rows = cleaned.apply(
        lambda row: any(not _mobile_live_call_value_is_empty(value) for value in row), axis=1,
    )
    cleaned = cleaned.loc[nonblank_rows, list(CALL_ATTENDENCE_FILE_COLUMNS)].reset_index(drop=True)
    if cleaned.empty:
        raise ValueError("The uploaded file contains no data rows.")
    return cleaned


def prepare_call_attendence_rows(
    frame: pd.DataFrame,
    schema: list[dict[str, Any]],
) -> tuple[list[dict[str, object]], int]:
    """Coerce values to database types and remove duplicate rows within the file."""
    schema_by_name = {str(column["column_name"]).casefold(): column for column in schema}
    missing_database_columns = [
        database_column
        for database_column in CALL_ATTENDENCE_FILE_COLUMNS.values()
        if database_column.casefold() not in schema_by_name
    ]
    if missing_database_columns:
        raise ValueError(
            "Database table public.call_attendence is missing column(s): "
            + ", ".join(missing_database_columns)
            + "."
        )

    prepared: list[dict[str, object]] = []
    seen_row_keys: set[str] = set()
    duplicate_rows = 0
    for row_number, (_, row) in enumerate(frame.iterrows(), start=2):
        values: dict[str, object] = {}
        for file_column, database_column in CALL_ATTENDENCE_FILE_COLUMNS.items():
            schema_column = schema_by_name[database_column.casefold()]
            values[database_column] = _coerce_mobile_live_call_value(
                row[file_column], schema_column, row_number, file_column,
            )
        row_key = _mobile_live_call_record_key(values)
        if row_key in seen_row_keys:
            duplicate_rows += 1
            continue
        seen_row_keys.add(row_key)
        prepared.append(values)
    return prepared, duplicate_rows


def store_call_attendence_upload(
    content: bytes,
    filename: str,
    extension: str,
) -> dict[str, Any] | JSONResponse:
    """Insert new call-attendance rows atomically while rejecting duplicate files/rows."""
    frame = clean_call_attendence_frame(read_call_attendence_frame(content, extension))
    file_hash = hashlib.sha256(content).hexdigest()
    target_ref = "public.call_attendence"
    history_ref = f"public.{CALL_ATTENDENCE_HISTORY_TABLE}"

    with get_database_engine().begin() as connection:
        schema = _mobile_live_call_schema(connection, "call_attendence")
        if not schema:
            raise ValueError("Database table public.call_attendence was not found.")
        prepared_rows, duplicate_in_file = prepare_call_attendence_rows(frame, schema)

        connection.execute(text(f"""
            CREATE TABLE IF NOT EXISTS {history_ref} (
                file_sha256 VARCHAR(64) PRIMARY KEY,
                original_filename TEXT NOT NULL,
                uploaded_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
                total_row_count INTEGER NOT NULL,
                inserted_row_count INTEGER NOT NULL DEFAULT 0,
                skipped_row_count INTEGER NOT NULL DEFAULT 0
            )
        """))
        # Serialize imports so simultaneous uploads cannot insert the same logical row.
        connection.execute(text("SELECT pg_advisory_xact_lock(hashtext('public.call_attendence.upload'))"))
        reserved = connection.execute(
            text(f"""
                INSERT INTO {history_ref} (file_sha256, original_filename, total_row_count)
                VALUES (:file_sha256, :original_filename, :total_row_count)
                ON CONFLICT (file_sha256) DO NOTHING
                RETURNING file_sha256
            """),
            {
                "file_sha256": file_hash,
                "original_filename": Path(filename).name,
                "total_row_count": len(frame),
            },
        ).scalar_one_or_none()
        if reserved is None:
            return JSONResponse(
                status_code=409,
                content={
                    "status": "already_uploaded",
                    "message": "File Already Uploaded: no rows were added to call_attendence.",
                    "insertedRows": 0,
                    "skippedRows": len(frame),
                    "totalRows": len(frame),
                },
            )

        database_columns = list(CALL_ATTENDENCE_FILE_COLUMNS.values())
        insert_columns = ", ".join(quote_pg_identifier(column) for column in database_columns)
        selected_columns = ", ".join(
            f"staged.{quote_pg_identifier(column)}" for column in database_columns
        )
        duplicate_predicate = " AND ".join(
            f"existing.{quote_pg_identifier(column)} "
            f"IS NOT DISTINCT FROM staged.{quote_pg_identifier(column)}"
            for column in database_columns
        )
        connection.execute(text(f"""
            CREATE TEMP TABLE call_attendence_upload_stage
            ON COMMIT DROP
            AS SELECT {insert_columns} FROM {target_ref} WITH NO DATA
        """))
        if prepared_rows:
            stage_parameters = ", ".join(f":{column}" for column in database_columns)
            connection.execute(
                text(f"""
                    INSERT INTO call_attendence_upload_stage ({insert_columns})
                    VALUES ({stage_parameters})
                """),
                prepared_rows,
            )
            insert_result = connection.execute(text(f"""
                INSERT INTO {target_ref} ({insert_columns})
                SELECT {selected_columns}
                FROM call_attendence_upload_stage AS staged
                WHERE NOT EXISTS (
                    SELECT 1
                    FROM {target_ref} AS existing
                    WHERE {duplicate_predicate}
                )
            """))
            inserted_rows = max(int(insert_result.rowcount or 0), 0)
        else:
            inserted_rows = 0
        skipped_rows = duplicate_in_file + len(prepared_rows) - inserted_rows

        connection.execute(
            text(f"""
                UPDATE {history_ref}
                SET inserted_row_count = :inserted_rows,
                    skipped_row_count = :skipped_rows
                WHERE file_sha256 = :file_sha256
            """),
            {
                "inserted_rows": inserted_rows,
                "skipped_rows": skipped_rows,
                "file_sha256": file_hash,
            },
        )

    return {
        "status": "success",
        "filename": Path(filename).name,
        "message": (
            f"Upload complete: {inserted_rows:,} new row(s) inserted into call_attendence; "
            f"{skipped_rows:,} duplicate row(s) skipped."
        ),
        "insertedRows": inserted_rows,
        "skippedRows": skipped_rows,
        "totalRows": len(frame),
    }


@app.post("/api/attendance/fmcg-mobile-live-call-summary")
async def upload_fmcg_mobile_live_call_summary(file: UploadFile = File(...)) -> Any:
    """Validate a daily call-attendance file and immediately store its new rows."""
    filename = file.filename or ""
    extension = Path(filename).suffix.casefold()
    if extension not in {".csv", ".xlsx", ".xls"}:
        raise HTTPException(
            status_code=400,
            detail="Choose a valid CSV, XLSX, or XLS file for the Mobile Live Call Attendence.",
        )

    content = await file.read(MAX_UPLOAD_BYTES + 1)
    await file.close()
    if not content:
        raise HTTPException(status_code=400, detail="The Mobile Live Call Attendence file is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="The Mobile Live Call Attendence file exceeds the 200 MB upload limit.")

    if extension == ".xlsx":
        try:
            with ZipFile(BytesIO(content)) as archive:
                if "[Content_Types].xml" not in archive.namelist():
                    raise ValueError("missing XLSX content types")
        except (BadZipFile, ValueError) as error:
            raise HTTPException(status_code=400, detail="The Mobile Live Call Attendence XLSX file is not valid.") from error

    try:
        return store_call_attendence_upload(content, filename, extension)
    except (SQLAlchemyError, ValueError, TypeError) as error:
        logger.exception("call_attendence upload failed", exc_info=error)
        raise HTTPException(
            status_code=400,
            detail=f"Upload Failed: no rows were added. {error}",
        ) from error


def unique_headers(values: list[object]) -> list[str]:
    """Give blank and duplicate spreadsheet headers stable display names."""
    counts: dict[str, int] = {}
    headers = []
    for index, value in enumerate(values):
        base = display_value(value).strip() or f"Column {index + 1}"
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base} ({counts[base]})")
    return headers


def first_summary_header(sample_rows: list[tuple[Any, ...]]) -> int:
    """Find the first row that looks like a real multi-column header."""
    for index, row in enumerate(sample_rows):
        if sum(bool(display_value(value).strip()) for value in row) >= 2:
            return index
    return 0


def read_call_nature_frame(content: bytes, extension: str) -> pd.DataFrame:
    """Read call data using its actual header row, even when a title row precedes it."""
    sample = pd.read_excel(BytesIO(content), header=None, dtype=object, nrows=200).fillna("")
    if sample.empty:
        return pd.DataFrame()
    header_index = first_summary_header([tuple(row) for row in sample.itertuples(index=False, name=None)])
    headers = unique_headers(sample.iloc[header_index].tolist())
    return pd.read_excel(
        BytesIO(content),
        header=None,
        names=headers,
        skiprows=range(header_index + 1),
        dtype=object,
    )


def read_attendance_page(content: bytes, extension: str, offset: int) -> dict[str, Any]:
    """Read one 50-row page and its total count without returning all rows."""
    limit = ATTENDANCE_PREVIEW_ROWS
    if extension == ".xlsx":
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        sample_rows = []
        for _ in range(200):
            try:
                sample_rows.append(next(rows))
            except StopIteration:
                break
        header_index = first_summary_header(sample_rows)
        if not sample_rows:
            workbook.close()
            return {"columns": [], "rows": [], "total": 0, "offset": offset, "limit": limit, "hasNext": False}
        header_values = list(sample_rows[header_index])
        last_header = max((index for index, value in enumerate(header_values) if display_value(value).strip()), default=0)
        headers = unique_headers(header_values[: last_header + 1])
        data_rows = []
        total = 0
        data_iter = chain(sample_rows[header_index + 1:], rows)
        for row_index, row in enumerate(data_iter):
            values = list(row[:len(headers)])
            if not any(display_value(value).strip() for value in values):
                continue
            total += 1
            if offset <= total - 1 < offset + limit:
                data_rows.append({headers[index]: display_value(values[index]) if index < len(values) else "" for index in range(len(headers))})
        workbook.close()
        return {
            "columns": headers,
            "rows": data_rows,
            "total": total,
            "offset": offset,
            "limit": limit,
            "hasNext": offset + len(data_rows) < total,
        }

    if extension == ".csv":
        sample = pd.read_csv(BytesIO(content), header=None, dtype=str, keep_default_na=False, nrows=200)
        header_index = first_summary_header([tuple(row) for row in sample.itertuples(index=False, name=None)])
        headers = unique_headers(sample.iloc[header_index].tolist())
        frame = pd.read_csv(
            BytesIO(content),
            header=None,
            names=headers,
            dtype=str,
            keep_default_na=False,
            skiprows=range(header_index + 1 + offset),
            skipinitialspace=True,
            nrows=limit,
        )
        total = sum(
            len(chunk)
            for chunk in pd.read_csv(
                BytesIO(content),
                header=None,
                skiprows=header_index + 1,
                chunksize=10000,
            )
        )
    else:
        sample = pd.read_excel(BytesIO(content), header=None, dtype=str, nrows=200).fillna("")
        header_index = first_summary_header([tuple(row) for row in sample.itertuples(index=False, name=None)])
        headers = unique_headers(sample.iloc[header_index].tolist())
        frame = pd.read_excel(BytesIO(content), header=None, names=headers, skiprows=range(header_index + 1 + offset), dtype=str, nrows=limit).fillna("")
        total = len(pd.read_excel(BytesIO(content), header=None, skiprows=header_index + 1, usecols=[0]))

    frame = frame.fillna("").astype(str)
    return {
        "columns": headers,
        "rows": frame.to_dict(orient="records"),
        "total": total,
        "offset": offset,
        "limit": limit,
        "hasNext": offset + len(frame) < total,
    }


@app.get("/api/attendance/mobile-live-call-summary/{session_id}")
def mobile_live_call_summary_page(session_id: str, offset: int = 0) -> dict[str, Any]:
    if offset < 0:
        raise HTTPException(status_code=400, detail="The page offset cannot be negative.")
    session = get_attendance_session(session_id)
    extension = Path(session.filename).suffix.casefold()
    return read_attendance_page(session.content, extension, offset)


@app.get("/api/attendance/mobile-live-call-summary/{session_id}/download")
def download_mobile_live_call_summary(session_id: str) -> StreamingResponse:
    session = get_attendance_session(session_id)
    download_bytes = session.excel_content if session.excel_content is not None else session.content
    download_name = session.filename if session.filename.casefold().endswith(".xlsx") else f"{Path(session.filename).stem}.xlsx"
    return StreamingResponse(
        BytesIO(download_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(download_name)}"},
    )


@app.get("/api/attendance/mobile-live-call-summary/{session_id}/download-csv")
def download_mobile_live_call_summary_csv(session_id: str) -> StreamingResponse:
    session = get_attendance_session(session_id)
    csv_bytes = session.csv_content
    if csv_bytes is None:
        source_bytes = session.excel_content or session.content
        frame = pd.read_excel(BytesIO(source_bytes), dtype=object).fillna("")
        csv_bytes = frame.to_csv(index=False).encode("utf-8-sig")
    download_name = f"{Path(session.filename).stem}.csv"
    return StreamingResponse(
        BytesIO(csv_bytes),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(download_name)}"},
    )


def column_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def find_frame_column(frame: pd.DataFrame, requested_name: str) -> str | None:
    requested_key = column_key(requested_name)
    return next((str(column) for column in frame.columns if column_key(column) == requested_key), None)


def add_call_nature(frame: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, Any]]:
    """Add/update Call Nature immediately after Doctorname using call business rules."""
    doctor_column = find_frame_column(frame, "Doctorname")
    call_type_column = find_frame_column(frame, "CallType")
    missing = [
        name for name, column in (("Doctorname", doctor_column), ("CallType", call_type_column))
        if column is None
    ]
    if missing:
        raise HTTPException(
            status_code=400,
            detail="Unable to process file. Please check the Mobile Live Calls Summary format and required columns.",
        )

    # Match the source Excel formula's values exactly after normalizing case,
    # leading/trailing whitespace, repeated internal whitespace, and nulls.
    call_types = frame[call_type_column].map(normalized_name)
    doctor_names = frame[doctor_column].map(normalized_name)
    call_nature = pd.Series("", index=frame.index, dtype="string")

    exception_rows = call_types.eq("exception")
    present_rows = exception_rows & doctor_names.isin({
        "present at schedule morning",
        "present at schedule evening",
    })
    leave_rows = exception_rows & doctor_names.isin({
        "non mandatory leave",
        "mandatory leave",
        "short leave",
    })
    call_nature.loc[call_types.eq("call")] = "Provider Calls"
    call_nature.loc[present_rows] = "Attendance"
    call_nature.loc[leave_rows] = "LEAVE"
    call_nature.loc[exception_rows & ~present_rows & ~leave_rows] = "Others Exc Calls"

    existing_column = find_frame_column(frame, "Call Nature")
    if existing_column is not None and existing_column in frame.columns:
        frame = frame.drop(columns=[existing_column])
    doctor_position = list(frame.columns).index(doctor_column)
    frame.insert(doctor_position + 1, "Call Nature", call_nature)

    categories = ("Attendance", "LEAVE", "Others Exc Calls", "Provider Calls")
    counts = {category: int(call_nature.eq(category).sum()) for category in categories}
    counts["Blank"] = int(call_nature.eq("").sum())
    validation = {
        "totalRows": len(frame),
        "counts": counts,
        "unrecognizedCallTypeRows": int((~call_types.isin(("call", "exception"))).sum()),
        "preview": [
            {
                "Doctorname": display_value(row[doctor_column]),
                "Call Nature": display_value(row["Call Nature"]),
                "CallType": display_value(row[call_type_column]),
            }
            for _, row in frame[[doctor_column, "Call Nature", call_type_column]].head(10).iterrows()
        ],
    }
    return frame, validation


MOBILE_LIVE_CALL_TARGET_TABLE = "live_call_details"
MOBILE_LIVE_CALL_HISTORY_TABLE = "mobile_live_call_upload_history"
MOBILE_LIVE_CALL_UNIQUE_INDEX = "uq_live_call_details_exact_row"

MOBILE_LIVE_CALL_DATABASE_COLUMNS: dict[str, str] = {
    "Team": "team",
    "FF.Id": "ff_id",
    "FF.Name": "ff_name",
    "Call Date": "call_date",
    "Manager": "manager",
    "Base Town": "base_town",
    "PMDC": "pmdc",
    "Doctorname": "doctorname",
    "Call Nature": "call_nature",
    "CallType": "calltype",
    "WW": "ww",
    "LNG": "lng",
    "LAT": "lat",
    "Total Sum of NOOFCALLS": "total_sum_noofcalls",
    "Total Sum of NOOFPRODUCTS": "total_sum_noofproducts",
    "GeoMeter Range": "geometer_range",
}


def _mobile_live_call_table_ref(table_name: str) -> str:
    return f"public.{quote_pg_identifier(table_name)}"


def _mobile_live_call_schema(connection, table_name: str) -> list[dict[str, Any]]:
    rows = connection.execute(
        text("""
            SELECT
                column_name,
                ordinal_position,
                data_type,
                udt_name,
                character_maximum_length,
                is_nullable,
                column_default,
                is_identity,
                is_generated
            FROM information_schema.columns
            WHERE table_schema = 'public'
              AND table_name = :table_name
            ORDER BY ordinal_position
        """),
        {"table_name": table_name},
    ).mappings().all()
    return [dict(row) for row in rows]


def _ensure_mobile_live_call_upload_support(connection) -> None:
    target_ref = _mobile_live_call_table_ref(MOBILE_LIVE_CALL_TARGET_TABLE)
    history_ref = _mobile_live_call_table_ref(MOBILE_LIVE_CALL_HISTORY_TABLE)
    connection.execute(text(f"""
        CREATE TABLE IF NOT EXISTS {history_ref} (
            file_sha256 VARCHAR(64) PRIMARY KEY,
            original_filename TEXT NOT NULL,
            uploaded_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
            total_row_count INTEGER NOT NULL DEFAULT 0,
            inserted_row_count INTEGER NOT NULL DEFAULT 0,
            skipped_row_count INTEGER NOT NULL DEFAULT 0
        )
    """))
    unique_columns = ", ".join(
        quote_pg_identifier(column)
        for column in MOBILE_LIVE_CALL_DATABASE_COLUMNS.values()
    )
    connection.execute(text(f"""
        CREATE UNIQUE INDEX IF NOT EXISTS
            {quote_pg_identifier(MOBILE_LIVE_CALL_UNIQUE_INDEX)}
        ON {target_ref} ({unique_columns}) NULLS NOT DISTINCT
    """))


def _mobile_live_call_value_is_empty(value: object) -> bool:
    if value is None:
        return True
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _coerce_mobile_live_call_value(
    value: object,
    column: dict[str, Any],
    row_number: int,
    source_column: str,
) -> object:
    if _mobile_live_call_value_is_empty(value):
        converted: object = None
    else:
        converted = value

    data_type = str(column.get("data_type") or "").casefold()
    udt_name = str(column.get("udt_name") or "").casefold()
    if converted is None:
        if column.get("is_nullable") != "YES":
            raise ValueError(f"row {row_number}, column '{source_column}' cannot be blank")
        return None

    raw_text = str(converted).strip()
    if data_type in {"character varying", "character", "text"}:
        max_length = column.get("character_maximum_length")
        if max_length and len(raw_text) > int(max_length):
            raise ValueError(f"row {row_number}, column '{source_column}' exceeds the {max_length}-character limit")
        return raw_text

    if data_type in {"smallint", "integer", "bigint"} or udt_name in {"int2", "int4", "int8"}:
        try:
            numeric = Decimal(raw_text.replace(",", ""))
            if numeric != numeric.to_integral_value():
                raise ValueError
            return int(numeric)
        except (InvalidOperation, ValueError) as error:
            raise ValueError(f"row {row_number}, column '{source_column}' must be an integer") from error

    if data_type in {"numeric", "decimal"} or udt_name == "numeric":
        try:
            return Decimal(raw_text.replace(",", ""))
        except InvalidOperation as error:
            raise ValueError(f"row {row_number}, column '{source_column}' must be numeric") from error

    if data_type in {"real", "double precision"} or udt_name in {"float4", "float8"}:
        try:
            return float(raw_text.replace(",", ""))
        except ValueError as error:
            raise ValueError(f"row {row_number}, column '{source_column}' must be numeric") from error

    if data_type == "boolean" or udt_name == "bool":
        normalized = raw_text.casefold()
        if normalized in {"true", "t", "yes", "y", "1"}:
            return True
        if normalized in {"false", "f", "no", "n", "0"}:
            return False
        raise ValueError(f"row {row_number}, column '{source_column}' must be boolean")

    if data_type == "date":
        try:
            parsed = pd.to_datetime(raw_text, dayfirst=True, errors="raise")
            return parsed.date()
        except (TypeError, ValueError) as error:
            raise ValueError(f"row {row_number}, column '{source_column}' must be a valid date") from error

    if data_type in {"timestamp without time zone", "timestamp with time zone"}:
        try:
            parsed = pd.to_datetime(raw_text, dayfirst=True, errors="raise")
            return parsed.to_pydatetime()
        except (TypeError, ValueError) as error:
            raise ValueError(f"row {row_number}, column '{source_column}' must be a valid timestamp") from error

    if data_type == "uuid":
        from uuid import UUID

        try:
            return str(UUID(raw_text))
        except ValueError as error:
            raise ValueError(f"row {row_number}, column '{source_column}' must be a valid UUID") from error

    return converted


def _mobile_live_call_record_key(values: dict[str, object]) -> str:
    parts = []
    for column_name, value in values.items():
        if isinstance(value, datetime):
            serialized = value.isoformat()
        elif isinstance(value, date):
            serialized = value.isoformat()
        elif value is None:
            serialized = "<NULL>"
        else:
            serialized = re.sub(r"\s+", " ", str(value).strip()).casefold()
        parts.append(f"{header_key(column_name)}={serialized}")
    return hashlib.sha256("\x1f".join(parts).encode("utf-8")).hexdigest()


def _prepare_mobile_live_call_rows(
    frame: pd.DataFrame,
    schema: list[dict[str, Any]],
) -> tuple[list[dict[str, object]], int, list[str]]:
    schema_by_name = {
        str(column["column_name"]).casefold(): column
        for column in schema
    }
    source_by_key: dict[str, str] = {}
    for raw_column in frame.columns:
        source_column = str(raw_column)
        key = header_key(source_column)
        if key in source_by_key:
            raise ValueError(
                f"The file contains more than one column matching '{source_column}'."
            )
        source_by_key[key] = source_column

    missing_database_columns = [
        database_column
        for database_column in MOBILE_LIVE_CALL_DATABASE_COLUMNS.values()
        if database_column.casefold() not in schema_by_name
    ]
    if missing_database_columns:
        raise ValueError(
            f"Database table public.{MOBILE_LIVE_CALL_TARGET_TABLE} is missing column(s): "
            + ", ".join(missing_database_columns)
            + "."
        )

    missing_source_columns = [
        source_column
        for source_column in MOBILE_LIVE_CALL_DATABASE_COLUMNS
        if header_key(source_column) not in source_by_key
    ]
    if missing_source_columns:
        raise ValueError(
            "The processed file is missing required column(s): "
            + ", ".join(missing_source_columns)
            + "."
        )

    prepared: list[dict[str, object]] = []
    seen_keys: set[str] = set()
    duplicate_rows = 0
    for row_number, (_, row) in enumerate(frame.iterrows(), start=2):
        values: dict[str, object] = {}
        for expected_source, database_name in MOBILE_LIVE_CALL_DATABASE_COLUMNS.items():
            source_column = source_by_key[header_key(expected_source)]
            database_column = schema_by_name[database_name.casefold()]
            values[database_name] = _coerce_mobile_live_call_value(
                row[source_column],
                database_column,
                row_number,
                source_column,
            )
        record_key = _mobile_live_call_record_key(values)
        if record_key in seen_keys:
            duplicate_rows += 1
            continue
        seen_keys.add(record_key)
        prepared.append(values)
    return prepared, duplicate_rows, list(MOBILE_LIVE_CALL_DATABASE_COLUMNS.values())


def store_mobile_live_call_details(
    frame: pd.DataFrame,
    source_content: bytes,
    source_filename: str,
) -> dict[str, Any]:
    """Atomically store processed live-call rows with file and row deduplication."""
    file_hash = hashlib.sha256(source_content).hexdigest()
    target_ref = _mobile_live_call_table_ref(MOBILE_LIVE_CALL_TARGET_TABLE)
    history_ref = _mobile_live_call_table_ref(MOBILE_LIVE_CALL_HISTORY_TABLE)

    try:
        with get_database_engine().begin() as connection:
            connection.execute(
                text("SELECT pg_advisory_xact_lock(hashtext('public.live_call_details.upload'))")
            )
            schema = _mobile_live_call_schema(connection, MOBILE_LIVE_CALL_TARGET_TABLE)
            if not schema:
                raise ValueError(
                    f"Database table public.{MOBILE_LIVE_CALL_TARGET_TABLE} was not found or has no columns."
                )

            prepared_rows, duplicate_rows, value_columns = _prepare_mobile_live_call_rows(
                frame,
                schema,
            )
            _ensure_mobile_live_call_upload_support(connection)

            reserved = connection.execute(
                text(f"""
                    INSERT INTO {history_ref}
                        (file_sha256, original_filename, total_row_count)
                    VALUES (:file_sha256, :original_filename, :total_row_count)
                    ON CONFLICT (file_sha256) DO NOTHING
                    RETURNING file_sha256
                """),
                {
                    "file_sha256": file_hash,
                    "original_filename": Path(source_filename).name,
                    "total_row_count": len(frame),
                },
            ).scalar_one_or_none()
            if reserved is None:
                return {
                    "status": "already_uploaded",
                    "message": (
                        "File Already Uploaded: no rows were added to "
                        f"public.{MOBILE_LIVE_CALL_TARGET_TABLE}."
                    ),
                    "insertedRows": 0,
                    "skippedRows": len(frame),
                    "totalRows": len(frame),
                }

            quoted_columns = ", ".join(
                quote_pg_identifier(column) for column in value_columns
            )
            parameter_sql = ", ".join(f":{column}" for column in value_columns)
            connection.execute(text(f"""
                CREATE TEMP TABLE mobile_live_call_upload_stage
                ON COMMIT DROP
                AS SELECT {quoted_columns} FROM {target_ref} WITH NO DATA
            """))
            if prepared_rows:
                connection.execute(
                    text(f"""
                        INSERT INTO mobile_live_call_upload_stage ({quoted_columns})
                        VALUES ({parameter_sql})
                    """),
                    prepared_rows,
                )
                selected_columns = ", ".join(
                    f"staged.{quote_pg_identifier(column)}"
                    for column in value_columns
                )
                insert_result = connection.execute(text(f"""
                    INSERT INTO {target_ref} ({quoted_columns})
                    SELECT {selected_columns}
                    FROM mobile_live_call_upload_stage AS staged
                    ON CONFLICT DO NOTHING
                """))
                inserted_rows = max(int(insert_result.rowcount or 0), 0)
            else:
                inserted_rows = 0

            skipped_rows = duplicate_rows + len(prepared_rows) - inserted_rows
            connection.execute(
                text(f"""
                    UPDATE {history_ref}
                    SET inserted_row_count = :inserted_row_count,
                        skipped_row_count = :skipped_row_count
                    WHERE file_sha256 = :file_sha256
                """),
                {
                    "inserted_row_count": inserted_rows,
                    "skipped_row_count": skipped_rows,
                    "file_sha256": file_hash,
                },
            )
    except (SQLAlchemyError, ValueError, TypeError) as error:
        logger.exception("Mobile Live Calls Summary database upload failed", exc_info=error)
        raise HTTPException(
            status_code=400,
            detail=(
                "Upload Failed: the database transaction was rolled back. "
                f"{error} Inserted: 0, skipped: 0."
            ),
        ) from error

    return {
        "status": "success",
        "message": (
            f"Database upload complete: {inserted_rows:,} row(s) inserted into "
            f"public.{MOBILE_LIVE_CALL_TARGET_TABLE}; {skipped_rows:,} duplicate row(s) skipped."
        ),
        "insertedRows": inserted_rows,
        "skippedRows": skipped_rows,
        "totalRows": len(frame),
    }


@app.post("/api/attendance/mobile-live-call-summary/{session_id}/upload-database")
def upload_mobile_live_call_summary_to_database(session_id: str) -> JSONResponse:
    """Backward-compatible manual import endpoint for an existing session."""
    session = get_attendance_session(session_id)
    extension = Path(session.filename).suffix.casefold()
    try:
        frame = clean_and_validate_mobile_live_calls(
            read_mobile_live_calls_frame(session.content, extension),
        )
        processed_frame, _ = add_call_nature(frame)
        result = store_mobile_live_call_details(
            processed_frame,
            session.source_content or session.content,
            session.source_filename or session.filename,
        )
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(
            status_code=400,
            detail=f"Upload Failed: {error} Inserted: 0, skipped: 0.",
        ) from error

    status_code = 409 if result["status"] == "already_uploaded" else 200
    return JSONResponse(status_code=status_code, content=result)


@app.post("/api/attendance/call-nature")
async def process_call_nature(file: UploadFile = File(...)) -> dict[str, Any]:
    """Create a downloadable Excel workbook containing the Call Nature column."""
    filename = file.filename or ""
    if Path(filename).suffix.casefold() not in {".xlsx", ".xls"}:
        raise HTTPException(status_code=400, detail="Upload an Excel workbook (.xlsx or .xls).")
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    await file.close()
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="The workbook exceeds the 200 MB upload limit.")

    try:
        frame = read_call_nature_frame(content, Path(filename).suffix.casefold())
        processed, validation = add_call_nature(frame)
    except Exception as error:
        raise HTTPException(status_code=400, detail=f"Unable to process Call Nature: {error}") from error

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        processed.to_excel(writer, index=False, sheet_name="Processed Data")
    output_filename = f"{Path(filename).stem}-updated.xlsx"
    session_id = str(uuid4())
    processed_sessions[session_id] = UploadSession(
        content=output.getvalue(),
        filename=output_filename,
        created_at=datetime.now(timezone.utc),
    )
    return {"id": session_id, "filename": output_filename, "validation": validation}


@app.get("/api/attendance/call-nature/{session_id}/download")
def download_call_nature(session_id: str) -> StreamingResponse:
    session = get_processed_session(session_id)
    return StreamingResponse(
        BytesIO(session.content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(session.filename)}"},
    )


@app.get("/api/staff-values")
def staff_values() -> dict[str, Any]:
    empty = {column: [] for column in EXPECTED_COLUMNS}
    try:
        selections = ", ".join(
            f'{database_column} AS "{column}"'
            for column, database_column in DATABASE_FIELDS.items()
        )
        with get_database_engine().connect() as connection:
            frame = pd.read_sql_query(
                text(f"SELECT {selections} FROM public.ff_master_staff"),
                connection,
            )

        values: dict[str, list[str]] = {}
        for column in EXPECTED_COLUMNS:
            unique = {
                display_value(value).strip().casefold(): display_value(value).strip()
                for value in frame[column]
                if display_value(value).strip()
            }
            values[column] = sorted(unique.values(), key=str.casefold)
        return {"values": values}
    except Exception as error:
        logger.warning("Database values could not be loaded: %s", error)
        return {
            "values": empty,
            "warning": "Database values could not be loaded; you can still type values.",
        }


@app.post("/api/workbooks")
async def upload_workbook(workbook: UploadFile = File(...)) -> dict[str, Any]:
    if not (workbook.filename or "").casefold().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Choose a valid .xlsx workbook.")

    content = await workbook.read(MAX_UPLOAD_BYTES + 1)
    await workbook.close()
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="The workbook exceeds the 200 MB upload limit.")

    try:
        source = load_workbook(BytesIO(content), read_only=True, data_only=False)
        sheet_names = source.sheetnames
        recommended_sheet = next(
            (
                worksheet.title
                for worksheet in source.worksheets
                if worksheet_has_editable_headers(worksheet)
            ),
            sheet_names[0] if sheet_names else None,
        )
        source.close()
    except Exception as error:
        raise HTTPException(status_code=400, detail=f"Unable to read this workbook: {error}") from error

    session_id = str(uuid4())
    upload_sessions[session_id] = UploadSession(
        content=content,
        filename=clean_filename(workbook.filename),
        created_at=datetime.now(timezone.utc),
    )
    cleanup_sessions()
    return {
        "id": session_id,
        "filename": workbook.filename,
        "sheets": sheet_names,
        "recommendedSheet": recommended_sheet,
    }


@app.get("/api/workbooks/{session_id}/sheets/{sheet_name}")
def worksheet_records(session_id: str, sheet_name: str) -> dict[str, Any]:
    session = get_upload_session(session_id)
    return read_editable_data(session.content, sheet_name)


@app.post("/api/workbooks/{session_id}/export")
def export_workbook(session_id: str, changes: ExportRequest) -> StreamingResponse:
    session = get_upload_session(session_id)
    if len(changes.records) != len(changes.rowSources):
        raise HTTPException(status_code=400, detail="Record and source-row counts do not match.")
    if len(changes.columnNumbers) not in (len(EDITABLE_COLUMNS), len(EXPECTED_COLUMNS)):
        raise HTTPException(status_code=400, detail="The editable column mapping is incomplete.")

    workbook = load_workbook(BytesIO(session.content), data_only=False)
    try:
        if changes.sheetName not in workbook.sheetnames:
            raise HTTPException(status_code=404, detail="Worksheet not found.")
        worksheet = workbook[changes.sheetName]

        retained_rows = {row for row in changes.rowSources if row is not None}
        for row_number in set(changes.originalRows) - retained_rows:
            for cell in worksheet[row_number]:
                cell.value = None

        new_record_count = sum(row is None for row in changes.rowSources)
        if new_record_count:
            worksheet.insert_rows(changes.insertAtRow, amount=new_record_count)

        next_new_row = changes.insertAtRow
        for record, source_row in zip(changes.records, changes.rowSources):
            if source_row is None:
                target_row = next_new_row
                next_new_row += 1
                values = [record.get(column, "") for column in EXPECTED_COLUMNS]
                format_row = matching_format_row(
                    worksheet,
                    changes.originalRows,
                    changes.columnNumbers,
                    values,
                )
                if format_row is not None:
                    copy_row_format(worksheet, format_row, target_row)
            else:
                target_row = source_row

            mapping_columns = (
                EDITABLE_COLUMNS
                if len(changes.columnNumbers) == len(EDITABLE_COLUMNS)
                else EXPECTED_COLUMNS
            )
            column_mapping = dict(zip(mapping_columns, changes.columnNumbers))
            for column in EDITABLE_COLUMNS:
                write_record_cell(
                    worksheet.cell(target_row, column_mapping[column]),
                    record.get(column, ""),
                    column,
                )

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
    finally:
        workbook.close()

    encoded_name = quote(session.filename)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_name}"},
    )


TEAM_SUMMARY_CONFIGS: dict[str, dict[str, str]] = {
    "FMCG": {
        "team_key": "FMCG",
        "region_label": "FMCG Region",
        "title_prefix": "FMCG Team – IKON MTD Call Performance Summary Till",
        "sheet_title": "FMCG IKON MTD Summary",
        "filename_prefix": "FMCG_IKON_MTD_Summary",
        "where_clause": """
            position_code IS NOT NULL
            AND (
                team_name ILIKE '%FAST MOVING CONSUMER GOODS%'
                OR team_name ILIKE '%FMCG%'
                OR team_id IN ('1', '2', '3', '15')
            )
            AND COALESCE(team_name, '') NOT ILIKE '%PHARMA%'
        """,
    },
    "PHARMA": {
        "team_key": "PHARMA",
        "region_label": "PHARMA Region",
        "title_prefix": "PHARMA Team – IKON MTD Call Performance Summary Till",
        "sheet_title": "PHARMA IKON MTD Summary",
        "filename_prefix": "PHARMA_IKON_MTD_Summary",
        "where_clause": """
            position_code IS NOT NULL
            AND (
                team_name ILIKE '%PHARMA%'
                OR team_id = '4'
            )
            AND COALESCE(team_name, '') NOT ILIKE '%FAST MOVING CONSUMER GOODS%'
            AND COALESCE(team_name, '') NOT ILIKE '%FMCG%'
        """,
    },
}

TEAM_DESIGNATION_ORDERS: dict[str, tuple[str, ...]] = {
    "FMCG": ("SPO", "SPOWS", "MIO", "PS", "TE", "ASM", "ASMTE", "CM", "RSM", "DGM", "HOD"),
    "PHARMA": ("MIO", "PS", "ASM", "RSM", "DGM", "HOD"),
}

# Reps are based on the real reporting tree, not on the number of rows in a
# region.  These role sets describe the field positions present in each team;
# the employee counts themselves are always derived from ff_master_staff.
TEAM_REP_RULES: dict[str, dict[str, set[str]]] = {
    "FMCG": {
        "direct_rep_designations": {"SPO"},
        "field_employee_designations": {"SPO", "SPOWS", "TE"},
    },
    "PHARMA": {
        "direct_rep_designations": {"MIO", "PS"},
        "field_employee_designations": {"MIO", "PS"},
    },
}


def get_team_summary_config(team_key: str) -> dict[str, str]:
    normalized = (team_key or "").strip().upper()
    if normalized not in TEAM_SUMMARY_CONFIGS:
        raise HTTPException(status_code=400, detail=f"Unsupported team summary key: '{team_key}'.")
    return TEAM_SUMMARY_CONFIGS[normalized]


def get_team_designation_order_case(team_key: str) -> str:
    order = TEAM_DESIGNATION_ORDERS.get((team_key or "").strip().upper(), ())
    clauses = " ".join(f"WHEN '{designation}' THEN {rank}" for rank, designation in enumerate(order, start=1))
    return f"CASE UPPER(COALESCE(designation_name, '')) {clauses} ELSE {len(order) + 1} END"


def get_team_rep_rules(team_key: str) -> dict[str, set[str]]:
    return TEAM_REP_RULES.get((team_key or "").strip().upper(), {
        "direct_rep_designations": set(),
        "field_employee_designations": set(),
    })


def get_team_summary_columns(team_key: str = "FMCG") -> list[tuple[str, str]]:
    cfg = get_team_summary_config(team_key)
    return [
        ("desg", "Legend"),
        ("zone", "Zone"),
        ("region", cfg["region_label"]),
        ("dept_territory_code", "Dept. Territory Code"),
        ("position_code", "Position_Code"),
        ("sas_id", "SAS_ID"),
        ("ff_name", "SPO Name"),
        ("base_town", "Base Town"),
        ("emp_type", "Emp. Type"),
        ("emp_code", "Emp. Code"),
        ("payroll_status", "Payroll Status"),
        ("doj", "D.O.J"),
        ("reps", "Reps"),
        ("vacant", "# Of Vacant"),
        ("day_att_used", "Active SPOs Today Attn"),
        ("mtd_att_used", "Active SPOs MTD Attn"),
        ("day_call_used", "Active SPOs Today Calls"),
        ("mtd_calls_used", "Active SPOs MTD Calls"),
        ("ikon_used", "IKON Used"),
    ]


FMCG_SUMMARY_COLUMNS = get_team_summary_columns("FMCG")
PHARMA_SUMMARY_COLUMNS = get_team_summary_columns("PHARMA")

REPORT_DATE_FALLBACK = date(2026, 10, 9)


def latest_call_attendence_date() -> date:
    """Return the newest reporting date available in the shared activity table."""
    try:
        with get_database_engine().connect() as connection:
            latest = connection.execute(
                text("SELECT MAX(call_date) FROM public.call_attendence")
            ).scalar_one_or_none()
        if isinstance(latest, datetime):
            return latest.date()
        if isinstance(latest, date):
            return latest
    except Exception as error:
        logger.warning("Unable to read the latest call_attendence date: %s", error)
    return REPORT_DATE_FALLBACK


def parse_report_date(report_date: str | None) -> date:
    if not report_date:
        return latest_call_attendence_date()
    try:
        return date.fromisoformat(report_date.strip())
    except ValueError as error:
        raise HTTPException(status_code=400, detail="Invalid report_date format. Expected YYYY-MM-DD.") from error


def resolve_ff_name(employe_name: object, last_employee: object) -> str:
    primary = display_value(employe_name).strip()
    secondary = display_value(last_employee).strip()
    invalid = {"", "N/A", "#N/A", "NONE", "NULL"}
    if primary.upper() not in invalid:
        return primary
    if secondary.upper() not in invalid:
        return secondary
    return ""


def build_team_dynamic_columns(
    selected_date: date,
    team_key: str = "FMCG",
) -> tuple[int, str, list[tuple[str, str, str]]]:
    base_columns = get_team_summary_columns(team_key)
    days_in_month = calendar.monthrange(selected_date.year, selected_date.month)[1]
    month_label = selected_date.strftime("%b-%Y").upper()
    columns: list[tuple[str, str, str]] = [
        (key, label, "base") for key, label in base_columns
    ]
    for day in range(1, days_in_month + 1):
        columns.append((f"att_day_{day}", str(day), "attendance"))
    columns.append(("attendance_total", "Attendance", "attendance"))

    for day in range(1, days_in_month + 1):
        columns.append((f"call_day_{day}", str(day), "calls"))
    columns.append(("calls_total", "Calls", "calls"))
    columns.append(("avg_call", "Avg Call", "calls"))
    return days_in_month, month_label, columns


def build_fmcg_dynamic_columns(selected_date: date) -> tuple[int, str, list[tuple[str, str, str]]]:
    return build_team_dynamic_columns(selected_date, "FMCG")


def parse_multi_filter_values(value: str) -> set[str]:
    """Accept the Excel-style multi-select values sent by the report filters."""
    return {
        item.strip().casefold()
        for item in (value or "").split("|")
        if item.strip()
    }


def fetch_team_summary_data(
    team_key: str = "FMCG",
    report_date: str | None = None,
    search: str = "",
    designation: str = "",
    zone: str = "",
    region: str = "",
    base_town: str = "",
    dept_territory_code: str = "",
    position_code: str = "",
    sas_id: str = "",
    employee_name: str = "",
    emp_type: str = "",
    emp_code: str = "",
    payroll_status: str = "",
    doj: str = "",
    sort_by: str = "",
    sort_dir: str = "asc",
) -> dict[str, Any]:
    cfg = get_team_summary_config(team_key)
    selected_date = parse_report_date(report_date)
    formatted_date = selected_date.strftime("%d-%b-%Y").upper()
    days_in_month, month_label, all_columns = build_team_dynamic_columns(selected_date, cfg["team_key"])
    month_days = list(range(1, days_in_month + 1))
    workdays_to_report = sum(
        date(selected_date.year, selected_date.month, day).weekday() != 6
        for day in month_days
        if day <= selected_date.day
    )

    query = text(f"""
        SELECT
            designation_name,
            zone_name,
            region_name,
            dept_territory_code,
            position_code,
            sas_id,
            employe_name,
            last_employee,
            nature_type,
            e_code,
            doj,
            base_town,
            actual_field_force,
            remarks,
            payroll_type,
            team_name,
            team_id,
            report_to_position_code
        FROM public.ff_master_staff
        WHERE {cfg['where_clause']}
        ORDER BY
            zone_name ASC NULLS LAST,
            region_name ASC NULLS LAST,
            {get_team_designation_order_case(cfg['team_key'])},
            CASE
                WHEN actual_field_force = 0
                  OR LOWER(COALESCE(remarks, '')) = 'un-mapped'
                  OR UPPER(COALESCE(e_code, '')) = 'UN-DEFINED'
                  OR UPPER(COALESCE(employe_name, '')) LIKE '%UNMAPPED%'
                  OR UPPER(COALESCE(position_code, '')) LIKE '%-00'
                THEN 0
                ELSE 1
            END,
            position_code ASC NULLS LAST
    """)

    with get_database_engine().connect() as connection:
        raw_rows = connection.execute(query).mappings().all()

    designations_set: set[str] = set()
    zones_set: set[str] = set()
    regions_set: set[str] = set()
    base_towns_set: set[str] = set()
    dept_territory_codes_set: set[str] = set()
    position_codes_set: set[str] = set()
    sas_ids_set: set[str] = set()
    employee_names_set: set[str] = set()
    employee_types_set: set[str] = set()
    employee_codes_set: set[str] = set()
    payroll_statuses_set: set[str] = set()
    dojs_set: set[str] = set()

    all_records: list[dict[str, Any]] = []
    for row in raw_rows:
        desg_val = display_value(row["designation_name"]).strip()
        zone_val = display_value(row["zone_name"]).strip()
        region_val = display_value(row["region_name"]).strip()
        dept_territory_val = display_value(row["dept_territory_code"]).strip()
        pos_val = display_value(row["position_code"]).strip()
        sas_val = display_value(row["sas_id"]).strip()
        emp_name_raw = display_value(row["employe_name"]).strip()
        last_emp_raw = display_value(row["last_employee"]).strip()
        ff_name_val = resolve_ff_name(emp_name_raw, last_emp_raw)
        emp_type_val = display_value(row["nature_type"]).strip()
        emp_code_val = display_value(row["e_code"]).strip()
        doj_val = display_value(row["doj"]).strip()
        base_town_val = display_value(row["base_town"]).strip()
        remarks_val = display_value(row["remarks"]).strip()
        payroll_val = display_value(row["payroll_type"]).strip()
        actual_ff = row["actual_field_force"]

        if desg_val:
            designations_set.add(desg_val)
        if zone_val:
            zones_set.add(zone_val)
        if region_val:
            regions_set.add(region_val)
        if base_town_val:
            base_towns_set.add(base_town_val)
        if dept_territory_val:
            dept_territory_codes_set.add(dept_territory_val)
        if pos_val:
            position_codes_set.add(pos_val)
        if sas_val:
            sas_ids_set.add(sas_val)
        if ff_name_val:
            employee_names_set.add(ff_name_val)
        if emp_type_val:
            employee_types_set.add(emp_type_val)
        if emp_code_val:
            employee_codes_set.add(emp_code_val)
        if payroll_val:
            payroll_statuses_set.add(payroll_val)
        if doj_val:
            dojs_set.add(doj_val)

        is_unmapped = (
            actual_ff == 0
            or remarks_val.casefold() == "un-mapped"
            or emp_code_val.upper() == "UN-DEFINED"
            or "UNMAPPED" in ff_name_val.upper()
            or pos_val.upper().endswith("-00")
        )
        is_vacant = (
            not is_unmapped
            and (
                emp_code_val.upper() == "VACANT"
                or payroll_val.upper() == "VACANT"
                or doj_val.upper() == "RESIGNED"
                or "VACANT" in emp_name_raw.upper()
            )
        )

        report_to_pos_val = display_value(row["report_to_position_code"]).strip()
        is_active_assigned = (
            not is_unmapped
            and not is_vacant
            and bool(emp_code_val)
            and emp_code_val.upper() not in {"UN-DEFINED", "VACANT"}
            and payroll_val.upper() != "VACANT"
            and doj_val.upper() != "RESIGNED"
            and "VACANT" not in emp_name_raw.upper()
        )
        vacant_val = 1 if is_vacant else 0

        record_item: dict[str, Any] = {
            "desg": desg_val,
            "zone": zone_val,
            "region": region_val,
            "dept_territory_code": dept_territory_val,
            "position_code": pos_val,
            "sas_id": sas_val,
            "ff_name": ff_name_val,
            "last_employee": "" if last_emp_raw.upper() in {"N/A", "#N/A"} else last_emp_raw,
            "emp_type": emp_type_val,
            "emp_code": emp_code_val,
            "payroll_status": payroll_val,
            "doj": doj_val,
            "base_town": base_town_val,
            "reps": 0,
            "vacant": vacant_val,
            "management_reps": 0,
            "management_vacant": 0,
            "day_att_used": None,
            "mtd_att_used": None,
            "day_call_used": None,
            "mtd_calls_used": None,
            "ikon_used": None,
            "attendance_days": {str(day): None for day in month_days},
            "attendance_total": None,
            "call_days": {str(day): None for day in month_days},
            "calls_total": None,
            "avg_call": None,
            "is_vacant": is_vacant,
            "is_unmapped": is_unmapped,
            "_report_to_position_code": report_to_pos_val,
            "_employee_key": pos_val.casefold() if pos_val else f"row:{len(all_records)}",
            "_is_active_assigned": is_active_assigned,
        }
        for day in month_days:
            record_item[f"att_day_{day}"] = None
            record_item[f"call_day_{day}"] = None
        all_records.append(record_item)

    rep_rules = get_team_rep_rules(cfg["team_key"])
    direct_rep_designations = rep_rules["direct_rep_designations"]
    field_employee_designations = rep_rules["field_employee_designations"]
    records_by_key: dict[str, dict[str, Any]] = {}
    children_by_parent: dict[str, set[str]] = {}
    for record in all_records:
        employee_key = record["_employee_key"]
        # A position code identifies one employee position.  If a source row
        # is repeated, keep the first source record in the relationship map so
        # it cannot inflate a manager's Reps count.
        records_by_key.setdefault(employee_key, record)
        report_to_key = record["_report_to_position_code"].casefold()
        if report_to_key:
            children_by_parent.setdefault(report_to_key, set()).add(employee_key)

    def active_field_employee_keys(record: dict[str, Any]) -> set[str]:
        result: set[str] = set()
        visited: set[str] = set()
        pending = list(children_by_parent.get(record["_employee_key"], set()))
        while pending:
            child_key = pending.pop()
            if child_key in visited:
                continue
            visited.add(child_key)
            child = records_by_key.get(child_key)
            if child is None:
                continue
            child_designation = child["desg"].strip().upper()
            if (
                child_designation in field_employee_designations
                and child["_is_active_assigned"]
            ):
                result.add(child_key)
            pending.extend(children_by_parent.get(child_key, set()))
        return result

    def management_field_position_keys(record: dict[str, Any]) -> set[str]:
        """Return all field positions below a manager for Management Print headcount."""
        result: set[str] = set()
        visited: set[str] = set()
        pending = list(children_by_parent.get(record["_employee_key"], set()))
        while pending:
            child_key = pending.pop()
            if child_key in visited:
                continue
            visited.add(child_key)
            child = records_by_key.get(child_key)
            if child is None:
                continue
            if child["desg"].strip().upper() in field_employee_designations:
                result.add(child_key)
            pending.extend(children_by_parent.get(child_key, set()))
        return result

    rep_members_by_key: dict[str, set[str]] = {}
    for record in all_records:
        designation_value = record["desg"].strip().upper()
        employee_key = record["_employee_key"]
        if designation_value in field_employee_designations:
            members = {employee_key} if record["_is_active_assigned"] else set()
        elif designation_value == "ASM":
            members = {
                child_key
                for child_key in children_by_parent.get(employee_key, set())
                if (
                    records_by_key.get(child_key)
                    and records_by_key[child_key]["desg"].strip().upper() in direct_rep_designations
                    and records_by_key[child_key]["_is_active_assigned"]
                )
            }
        else:
            members = active_field_employee_keys(record)
        rep_members_by_key[employee_key] = members
        record["reps"] = len(members)
        record["_rep_member_keys"] = members

        # Management Print excludes unmapped and vacant field positions from
        # Reps.  Vacant field slots are shown separately in # Of Vacant.
        management_positions = management_field_position_keys(record)
        record["management_reps"] = sum(
            not records_by_key[position_key]["is_vacant"]
            and not records_by_key[position_key]["is_unmapped"]
            for position_key in management_positions
        )
        record["management_vacant"] = sum(
            records_by_key[position_key]["is_vacant"]
            for position_key in management_positions
        )

    activity_by_position = fetch_call_attendence_activity_by_position(
        cfg["team_key"],
        selected_date,
    )
    attendance_designations = {"SPO"} if cfg["team_key"] == "FMCG" else {"MIO", "PS"}
    attendance_rollup_designations = (
        {"SPO", "ASM", "RSM"}
        if cfg["team_key"] == "FMCG"
        else {"MIO", "PS", "ASM", "RSM"}
    )

    def active_team_attendance_positions(record: dict[str, Any]) -> set[str]:
        result: set[str] = set()
        visited: set[str] = set()
        if (
            record["desg"].strip().upper() in attendance_designations
            and record["_is_active_assigned"]
        ):
            result.add(record["_employee_key"])

        pending = list(children_by_parent.get(record["_employee_key"], set()))
        while pending:
            child_key = pending.pop()
            if child_key in visited:
                continue
            visited.add(child_key)
            child = records_by_key.get(child_key)
            if child is None:
                continue
            if (
                child["desg"].strip().upper() in attendance_designations
                and child["_is_active_assigned"]
            ):
                result.add(child_key)
            pending.extend(children_by_parent.get(child_key, set()))
        return result

    for record in all_records:
        attendance_member_keys = active_team_attendance_positions(record)
        record["_attendance_member_keys"] = attendance_member_keys
        designation_value = record["desg"].strip().upper()
        if designation_value not in attendance_rollup_designations:
            continue

        position_attendance = {
            day: sum(
                1
                for position_key in attendance_member_keys
                if date(selected_date.year, selected_date.month, day)
                in activity_by_position.get(position_key, {}).get("attendance_days", set())
            )
            for day in month_days
        }
        for day in month_days:
            present_count = int(position_attendance.get(day, 0) or 0)
            if designation_value == "SPO":
                record[f"att_day_{day}"] = 1 if present_count else None
            else:
                # Bound manager counts to active field positions in the
                # current team hierarchy.
                record[f"att_day_{day}"] = min(present_count, len(attendance_member_keys))
            record["attendance_days"][str(day)] = record[f"att_day_{day}"]

        record["attendance_total"] = sum(
            int(record[f"att_day_{day}"] or 0)
            for day in month_days
        )

    def active_team_activity_positions(record: dict[str, Any]) -> set[str]:
        result: set[str] = set()
        visited: set[str] = set()
        if (
            record["desg"].strip().upper() in attendance_designations
            and record["_is_active_assigned"]
            and record["_employee_key"] in activity_by_position
        ):
            result.add(record["_employee_key"])

        pending = list(children_by_parent.get(record["_employee_key"], set()))
        while pending:
            child_key = pending.pop()
            if child_key in visited:
                continue
            visited.add(child_key)
            child = records_by_key.get(child_key)
            if child is None:
                continue
            if (
                child["desg"].strip().upper() in attendance_designations
                and child["_is_active_assigned"]
                and child_key in activity_by_position
            ):
                result.add(child_key)
            pending.extend(children_by_parent.get(child_key, set()))
        return result

    def aggregate_activity(position_keys: set[str]) -> dict[str, Any]:
        seen_employee_ids: set[str] = set()
        totals = {
            "today_att": 0,
            "mtd_att": 0,
            "today_calls": 0,
            "mtd_calls": 0,
            "calls_total": 0,
            "call_days": {day: 0 for day in month_days},
        }
        for position_key in position_keys:
            activity = activity_by_position.get(position_key)
            if not activity or activity["employee_id"] in seen_employee_ids:
                continue
            seen_employee_ids.add(activity["employee_id"])
            for metric_key in ("today_att", "mtd_att", "today_calls", "mtd_calls"):
                totals[metric_key] += activity[metric_key]
            totals["calls_total"] += int(activity.get("calls_total", 0) or 0)
            for day, count in activity.get("call_days", {}).items():
                totals["call_days"][day] = totals["call_days"].get(day, 0) + int(count or 0)
        return totals

    for record in all_records:
        activity_member_keys = active_team_activity_positions(record) if activity_by_position else set()
        record["_activity_member_keys"] = activity_member_keys
        if activity_by_position:
            activity_totals = aggregate_activity(activity_member_keys)
            record["day_att_used"] = activity_totals["today_att"]
            record["mtd_att_used"] = activity_totals["mtd_att"]
            record["day_call_used"] = activity_totals["today_calls"]
            record["mtd_calls_used"] = activity_totals["mtd_calls"]
            record["ikon_used"] = int(bool(activity_totals["mtd_att"] or activity_totals["mtd_calls"]))
            record["calls_total"] = activity_totals["calls_total"]
            record["avg_call"] = (
                activity_totals["calls_total"] / workdays_to_report
                if workdays_to_report
                else 0
            )
            for day in month_days:
                record[f"call_day_{day}"] = activity_totals["call_days"].get(day, 0)

    desg_filter = parse_multi_filter_values(designation)
    zone_filter = parse_multi_filter_values(zone)
    region_filter = parse_multi_filter_values(region)
    town_filter = parse_multi_filter_values(base_town)
    dept_territory_filter = dept_territory_code.strip().casefold()
    position_filter = position_code.strip().casefold()
    sas_filter = sas_id.strip().casefold()
    employee_name_filter = employee_name.strip().casefold()
    emp_type_filter = emp_type.strip().casefold()
    emp_code_filter = emp_code.strip().casefold()
    payroll_status_filter = payroll_status.strip().casefold()
    doj_filter = doj.strip().casefold()
    search_query = search.strip().casefold()
    search_no_leading_zero = search_query.lstrip("0") if search_query.isdigit() else ""

    filtered: list[dict[str, Any]] = []
    for record in all_records:
        if desg_filter and record["desg"].casefold() not in desg_filter:
            continue
        if zone_filter and record["zone"].casefold() not in zone_filter:
            continue
        if region_filter and record["region"].casefold() not in region_filter:
            continue
        if town_filter and record["base_town"].casefold() not in town_filter:
            continue
        if dept_territory_filter and record["dept_territory_code"].casefold() != dept_territory_filter:
            continue
        if position_filter and record["position_code"].casefold() != position_filter:
            continue
        if sas_filter and record["sas_id"].casefold() != sas_filter:
            continue
        if employee_name_filter and record["ff_name"].casefold() != employee_name_filter:
            continue
        if emp_type_filter and record["emp_type"].casefold() != emp_type_filter:
            continue
        if emp_code_filter and record["emp_code"].casefold() != emp_code_filter:
            continue
        if payroll_status_filter and record["payroll_status"].casefold() != payroll_status_filter:
            continue
        if doj_filter and record["doj"].casefold() != doj_filter:
            continue
        if search_query:
            haystack = " ".join([
                record["ff_name"],
                record["last_employee"],
                record["position_code"],
                record["sas_id"],
                f"0{record['sas_id']}" if record["sas_id"].isdigit() and len(record["sas_id"]) == 6 else "",
                record["emp_code"],
                record["base_town"],
                record["dept_territory_code"],
                record["emp_type"],
                record["payroll_status"],
                record["region"],
                record["desg"],
            ]).casefold()
            if search_query not in haystack and (not search_no_leading_zero or search_no_leading_zero not in record["sas_id"].casefold()):
                continue
        filtered.append(record)

    valid_sort_keys = {key for key, _, _ in all_columns}
    if sort_by in valid_sort_keys:
        reverse = sort_dir.strip().casefold() == "desc"
        text_keys = {
            "desg", "zone", "region", "dept_territory_code", "position_code", "sas_id",
            "ff_name", "base_town", "emp_type", "emp_code", "payroll_status", "doj",
        }
        if sort_by not in text_keys:
            filtered.sort(
                key=lambda item: (item.get(sort_by) is None, item.get(sort_by) if item.get(sort_by) is not None else 0),
                reverse=reverse,
            )
        else:
            filtered.sort(
                key=lambda item: str(item.get(sort_by) or "").casefold(),
                reverse=reverse,
            )

    total_rep_members: set[str] = set()
    for item in filtered:
        total_rep_members.update(item["_rep_member_keys"])
    total_reps = len(total_rep_members)
    total_vacant = sum(item["vacant"] for item in filtered)
    total_unmapped = sum(1 for item in filtered if item["is_unmapped"])
    filtered_attendance_total = sum(
        int(item["attendance_total"] or 0)
        for item in filtered
        if item["attendance_total"] is not None
    )

    activity_total_members = {
        "today_att": set(),
        "mtd_att": set(),
        "today_calls": set(),
        "mtd_calls": set(),
        "ikon": set(),
    }
    if activity_by_position:
        for item in filtered:
            seen_employee_ids: set[str] = set()
            for position_key in item["_activity_member_keys"]:
                activity = activity_by_position.get(position_key)
                if not activity or activity["employee_id"] in seen_employee_ids:
                    continue
                seen_employee_ids.add(activity["employee_id"])
                if activity["today_att"]:
                    activity_total_members["today_att"].add(activity["employee_id"])
                if activity["mtd_att"]:
                    activity_total_members["mtd_att"].add(activity["employee_id"])
                if activity["today_calls"]:
                    activity_total_members["today_calls"].add(activity["employee_id"])
                if activity["mtd_calls"]:
                    activity_total_members["mtd_calls"].add(activity["employee_id"])
                if activity["mtd_att"] or activity["mtd_calls"]:
                    activity_total_members["ikon"].add(activity["employee_id"])

    activity_total_call_counts: dict[str, int] = {}
    activity_total_call_days = {day: 0 for day in month_days}
    if activity_by_position:
        for item in filtered:
            seen_employee_ids: set[str] = set()
            for position_key in item["_activity_member_keys"]:
                activity = activity_by_position.get(position_key)
                if not activity or activity["employee_id"] in seen_employee_ids:
                    continue
                seen_employee_ids.add(activity["employee_id"])
                activity_total_call_counts[activity["employee_id"]] = int(activity.get("calls_total", 0) or 0)
                for day, count in activity.get("call_days", {}).items():
                    activity_total_call_days[day] += int(count or 0)

    total_calls = sum(activity_total_call_counts.values()) if activity_by_position else None
    average_calls = (
        total_calls / workdays_to_report
        if total_calls is not None and workdays_to_report
        else None
    )

    activity_totals = {
        "active_spos_today_attn": len(activity_total_members["today_att"]) if activity_by_position else None,
        "active_spos_mtd_attn": len(activity_total_members["mtd_att"]) if activity_by_position else None,
        "active_spos_today_calls": len(activity_total_members["today_calls"]) if activity_by_position else None,
        "active_spos_mtd_calls": len(activity_total_members["mtd_calls"]) if activity_by_position else None,
        "ikon_used": len(activity_total_members["ikon"]) if activity_by_position else None,
        "calls_total": total_calls,
        "avg_call": average_calls,
        "call_days": activity_total_call_days if activity_by_position else None,
    }

    for item in all_records:
        item.pop("_report_to_position_code", None)
        item.pop("_employee_key", None)
        item.pop("_is_active_assigned", None)
        item.pop("_rep_member_keys", None)
        item.pop("_activity_member_keys", None)
        item.pop("_attendance_member_keys", None)

    return {
        "team": cfg["team_key"],
        "regionLabel": cfg["region_label"],
        "reportDate": selected_date.isoformat(),
        "formattedDate": formatted_date,
        "monthLabel": month_label,
        "daysInMonth": days_in_month,
        "monthDays": month_days,
        "totalSourceRows": len(all_records),
        "filters": {
            "designations": sorted(designations_set, key=str.casefold),
            "zones": sorted(zones_set, key=str.casefold),
            "regions": sorted(regions_set, key=str.casefold),
            "baseTowns": sorted(base_towns_set, key=str.casefold),
            "deptTerritoryCodes": sorted(dept_territory_codes_set, key=str.casefold),
            "positionCodes": sorted(position_codes_set, key=str.casefold),
            "sasIds": sorted(sas_ids_set, key=str.casefold),
            "employeeNames": sorted(employee_names_set, key=str.casefold),
            "employeeTypes": sorted(employee_types_set, key=str.casefold),
            "employeeCodes": sorted(employee_codes_set, key=str.casefold),
            "payrollStatuses": sorted(payroll_statuses_set, key=str.casefold),
            "dojs": sorted(dojs_set, key=str.casefold),
        },
        "totals": {
            "rows": len(filtered),
            "reps": total_reps,
            "vacant": total_vacant,
            "activeFilled": total_reps,
            "unmapped": total_unmapped,
            "day_att_used": activity_totals["active_spos_today_attn"],
            "mtd_att_used": activity_totals["active_spos_mtd_attn"],
            "day_call_used": activity_totals["active_spos_today_calls"],
            "mtd_calls_used": activity_totals["active_spos_mtd_calls"],
            "ikon_used": activity_totals["ikon_used"],
            "attendance_total": filtered_attendance_total,
            "calls_total": activity_totals["calls_total"],
            "avg_call": activity_totals["avg_call"],
            "call_days": activity_totals["call_days"],
        },
        "unimplementedMetrics": [],
        "active_spos_today_attn": activity_totals["active_spos_today_attn"],
        "active_spos_mtd_attn": activity_totals["active_spos_mtd_attn"],
        "active_spos_today_calls": activity_totals["active_spos_today_calls"],
        "active_spos_mtd_calls": activity_totals["active_spos_mtd_calls"],
        "ikon_used": activity_totals["ikon_used"],
        "activityMetrics": activity_totals,
        "rows": filtered,
    }


def fetch_fmcg_summary_data(
    report_date: str | None = None,
    search: str = "",
    designation: str = "",
    zone: str = "",
    region: str = "",
    base_town: str = "",
    dept_territory_code: str = "",
    position_code: str = "",
    sas_id: str = "",
    employee_name: str = "",
    emp_type: str = "",
    emp_code: str = "",
    payroll_status: str = "",
    doj: str = "",
    sort_by: str = "",
    sort_dir: str = "asc",
) -> dict[str, Any]:
    return fetch_team_summary_data(
        team_key="FMCG",
        report_date=report_date,
        search=search,
        designation=designation,
        zone=zone,
        region=region,
        base_town=base_town,
        dept_territory_code=dept_territory_code,
        position_code=position_code,
        sas_id=sas_id,
        employee_name=employee_name,
        emp_type=emp_type,
        emp_code=emp_code,
        payroll_status=payroll_status,
        doj=doj,
        sort_by=sort_by,
        sort_dir=sort_dir,
    )


def export_team_ikon_mtd_summary_workbook(
    team_key: str = "FMCG",
    report_date: str | None = None,
    search: str = "",
    designation: str = "",
    zone: str = "",
    region: str = "",
    base_town: str = "",
    dept_territory_code: str = "",
    position_code: str = "",
    sas_id: str = "",
    employee_name: str = "",
    emp_type: str = "",
    emp_code: str = "",
    payroll_status: str = "",
    doj: str = "",
    sort_by: str = "",
    sort_dir: str = "asc",
) -> StreamingResponse:
    cfg = get_team_summary_config(team_key)
    summary = fetch_team_summary_data(
        team_key=cfg["team_key"],
        report_date=report_date,
        search=search,
        designation=designation,
        zone=zone,
        region=region,
        base_town=base_town,
        dept_territory_code=dept_territory_code,
        position_code=position_code,
        sas_id=sas_id,
        employee_name=employee_name,
        emp_type=emp_type,
        emp_code=emp_code,
        payroll_status=payroll_status,
        doj=doj,
        sort_by=sort_by,
        sort_dir=sort_dir,
    )
    selected_date = parse_report_date(report_date)
    days_in_month, month_label, all_columns = build_team_dynamic_columns(selected_date, cfg["team_key"])
    formatted_date = summary["formattedDate"]
    base_columns = get_team_summary_columns(cfg["team_key"])

    wb = Workbook()
    ws = wb.active
    ws.title = cfg["sheet_title"]

    thin_side = Side(style="thin", color="B7DECF")
    dark_side = Side(style="thin", color="064E3B")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    header_border = Border(left=dark_side, right=dark_side, top=dark_side, bottom=dark_side)

    title_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    header_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    att_group_fill = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
    att_col_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    calls_group_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
    calls_col_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    even_fill = PatternFill(start_color="F8FFFB", end_color="F8FFFB", fill_type="solid")
    vacant_fill = PatternFill(start_color="FFF1F2", end_color="FFF1F2", fill_type="solid")
    footer_fill = PatternFill(start_color="064E3B", end_color="064E3B", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    designation_row_styles = {
        "SPO": (white_fill, "022C22"),
        "MIO": (white_fill, "022C22"),
        "ASM": (PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"), "92400E"),
        "ASMTE": (PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"), "92400E"),
        "RSM": (PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid"), "3730A3"),
        "TE": (PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid"), "1E3A8A"),
        "CM": (PatternFill(start_color="CFFAFE", end_color="CFFAFE", fill_type="solid"), "155E75"),
        "DGM": (PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid"), "9A3412"),
        "HOD": (PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"), "166534"),
        "PS": (PatternFill(start_color="FCE7F3", end_color="FCE7F3", fill_type="solid"), "9D174D"),
        "SPOWS": (PatternFill(start_color="F3E8FF", end_color="F3E8FF", fill_type="solid"), "6B21A8"),
    }

    base_count = len(base_columns)
    att_start = base_count + 1
    att_end = base_count + days_in_month + 1
    calls_start = att_end + 1
    calls_end = att_end + days_in_month + 2

    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=base_count)
    title_cell = ws.cell(1, 1, f"{cfg['title_prefix']} {formatted_date}")
    title_cell.font = Font(name="Calibri", size=12, bold=True, color="022C22")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    for col_idx in range(1, base_count + 1):
        ws.cell(1, col_idx).border = header_border
        ws.cell(1, col_idx).fill = title_fill

    ws.merge_cells(start_row=1, start_column=att_start, end_row=1, end_column=att_end)
    att_title_cell = ws.cell(1, att_start, f"Attendance ({month_label})")
    att_title_cell.font = Font(name="Calibri", size=11, bold=True, color="022C22")
    att_title_cell.alignment = Alignment(horizontal="center", vertical="center")
    for col_idx in range(att_start, att_end + 1):
        ws.cell(1, col_idx).border = header_border
        ws.cell(1, col_idx).fill = att_group_fill

    ws.merge_cells(start_row=1, start_column=calls_start, end_row=1, end_column=calls_end)
    calls_title_cell = ws.cell(1, calls_start, f"Calls + Exceptions Other Calls ({month_label})")
    calls_title_cell.font = Font(name="Calibri", size=11, bold=True, color="022C22")
    calls_title_cell.alignment = Alignment(horizontal="center", vertical="center")
    for col_idx in range(calls_start, calls_end + 1):
        ws.cell(1, col_idx).border = header_border
        ws.cell(1, col_idx).fill = calls_group_fill

    ws.row_dimensions[2].height = 24
    for col_idx, (_key, label, group) in enumerate(all_columns, start=1):
        cell = ws.cell(2, col_idx, label)
        cell.font = Font(name="Calibri", size=10, bold=True, color="022C22")
        cell.fill = att_col_fill if group == "attendance" else (calls_col_fill if group == "calls" else header_fill)
        cell.border = header_border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    text_keys = {
        "desg", "zone", "region", "dept_territory_code", "position_code", "sas_id",
        "ff_name", "base_town", "emp_type", "emp_code", "payroll_status", "doj",
    }

    for row_offset, record in enumerate(summary["rows"], start=3):
        is_vacant = record["is_vacant"]
        designation = str(record.get("desg") or "").strip().upper()
        designation_style = designation_row_styles.get(designation)
        row_fill = vacant_fill if is_vacant else (designation_style[0] if designation_style else (even_fill if row_offset % 2 == 0 else None))
        for col_idx, (key, _label, _group) in enumerate(all_columns, start=1):
            raw_val = record.get(key)
            if key not in text_keys:
                cell_val = "-" if raw_val is None else raw_val
            else:
                cell_val = raw_val if raw_val else "-"
            cell = ws.cell(row_offset, col_idx, cell_val)
            cell.font = Font(
                name="Calibri",
                size=10,
                bold=is_vacant and key in {"ff_name", "emp_code", "doj", "vacant"},
                color=(
                    "9F1239" if is_vacant and key in {"ff_name", "emp_code", "doj", "vacant"}
                    else "022C22" if is_vacant
                    else designation_style[1] if designation_style else "022C22"
                ),
            )
            if row_fill:
                cell.fill = row_fill
            cell.border = cell_border
            cell.alignment = Alignment(
                horizontal="left" if key in text_keys else "right",
                vertical="center",
            )

    footer_row = len(summary["rows"]) + 3
    ws.row_dimensions[footer_row].height = 22
    for col_idx, (key, _label, _group) in enumerate(all_columns, start=1):
        if col_idx == 1:
            val: Any = "Total"
        elif key == "ff_name":
            val = f"{summary['totals']['rows']} Record(s)"
        elif key == "reps":
            val = summary["totals"]["reps"]
        elif key == "vacant":
            val = summary["totals"]["vacant"]
        elif key not in text_keys:
            val = "-"
        else:
            val = ""
        cell = ws.cell(footer_row, col_idx, val)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = footer_fill
        cell.border = header_border
        cell.alignment = Alignment(
            horizontal="left" if key in text_keys else "right",
            vertical="center",
        )

    base_widths = [8, 10, 18, 22, 28, 12, 30, 22, 16, 14, 16, 14, 8, 12, 18, 18, 20, 20, 12]
    att_widths = [5] * days_in_month + [12]
    calls_widths = [5] * days_in_month + [10, 10]
    for col_idx, width in enumerate(base_widths + att_widths + calls_widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A3"

    # Both team exports include the same management-print worksheet.  The
    # underlying summary is already team-scoped, so this sheet cannot mix
    # FMCG and PHARMA staff or activity.
    if cfg["team_key"] in TEAM_SUMMARY_CONFIGS:
        management_ws = wb.create_sheet("Management Print")
        management_ws.sheet_view.showGridLines = False

        management_columns = [
            ("desg", "Legend", "base"),
            ("zone", "Zone", "base"),
            ("region", "Region", "base"),
            ("ff_name", "F.F Name", "base"),
            ("reps", "Reps", "base"),
            ("vacant", "# Of Vacant", "base"),
            ("not_used_ikon_today", "Not Used IKON Today", "base"),
            ("att_yesterday", "# of Attn Yesterday", "attendance"),
            ("day_att_used", "Active SPOs Today", "attendance"),
            ("attendance_total", "# of Attn MTD", "attendance"),
            ("mtd_att_used", "Active SPOs MTD", "attendance"),
            ("calls_yesterday", "# of Calls Yesterday", "calls"),
            ("calls_today", "# of Calls Today", "calls"),
            ("day_call_used", "Active SPOs Today", "calls"),
            ("calls_total", "# of Calls MTD", "calls"),
            ("mtd_calls_used", "Active SPOs MTD", "calls"),
            ("avg_call", "Average P/Rep/Day", "calls"),
            ("comments", "Comments", "comments"),
        ]
        management_designations = {"ASM", "ASMTE", "CM", "RSM", "DGM", "HOD", "TE"}
        management_rows = [
            record
            for record in summary["rows"]
            if str(record.get("desg") or "").strip().upper() in management_designations
        ]
        if not management_rows:
            management_rows = list(summary["rows"])

        report_day = selected_date.day
        yesterday_day = report_day - 1 if report_day > 1 else None

        def management_day_value(record: dict[str, Any], prefix: str, day: int | None) -> object:
            return record.get(f"{prefix}_{day}") if day else None

        def not_used_ikon_today(record: dict[str, Any]) -> int | None:
            values = [
                record.get("management_reps", record.get("reps")),
                record.get("day_att_used"),
            ]
            if any(value is None for value in values):
                return None
            try:
                return max(int(values[0]) - int(values[1]), 0)
            except (TypeError, ValueError):
                return None

        def management_value(record: dict[str, Any], key: str) -> object:
            if key == "reps":
                value = record.get("management_reps", record.get("reps"))
            elif key == "vacant":
                value = record.get("management_vacant", record.get("vacant"))
            elif key == "not_used_ikon_today":
                value = not_used_ikon_today(record)
            elif key == "att_yesterday":
                value = management_day_value(record, "att_day", yesterday_day)
            elif key == "calls_yesterday":
                value = management_day_value(record, "call_day", yesterday_day)
            elif key == "calls_today":
                value = management_day_value(record, "call_day", report_day)
            elif key == "comments":
                value = ""
            else:
                value = record.get(key)
            if value is None and key not in {"comments"}:
                return "-"
            return value

        management_data = [
            [management_value(record, key) for key, _label, _group in management_columns]
            for record in management_rows
        ]

        management_ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=7)
        management_ws.merge_cells(start_row=1, start_column=8, end_row=1, end_column=11)
        management_ws.merge_cells(start_row=2, start_column=8, end_row=2, end_column=11)
        management_ws.merge_cells(start_row=1, start_column=12, end_row=1, end_column=17)
        management_ws.merge_cells(start_row=2, start_column=12, end_row=2, end_column=17)
        management_ws.merge_cells(start_row=1, start_column=19, end_row=2, end_column=19)
        management_ws.row_dimensions[1].height = 31
        management_ws.row_dimensions[2].height = 25
        management_ws.row_dimensions[3].height = 42

        comments_fill = PatternFill(start_color="A9D18E", end_color="A9D18E", fill_type="solid")
        full_date = selected_date.strftime("%d-%b-%Y")
        short_date = selected_date.strftime("%d-%b-%y")
        group_headers = [
            (1, 7, f"{cfg['team_key']} Team - IKON MTD Call Performance Summary Till\n{full_date}", title_fill),
            (8, 11, "Attendance", att_group_fill),
            (12, 17, "Calls + Exception Other Calls", calls_group_fill),
            (19, 19, "Comments", comments_fill),
        ]
        for start_column, end_column, label, fill in group_headers:
            cell = management_ws.cell(1, start_column, label)
            cell.font = Font(name="Calibri", size=11 if start_column == 1 else 10, bold=True, color="022C22")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row_idx in range(1, 3):
                for col_idx in range(start_column, end_column + 1):
                    management_ws.cell(row_idx, col_idx).fill = fill
                    management_ws.cell(row_idx, col_idx).border = header_border

        for start_column, label in ((8, short_date), (12, short_date)):
            cell = management_ws.cell(2, start_column, label)
            cell.font = Font(name="Calibri", size=10, bold=True, color="022C22")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        management_group_fills = {
            "base": header_fill,
            "attendance": att_col_fill,
            "calls": calls_col_fill,
            "comments": PatternFill(start_color="A9D18E", end_color="A9D18E", fill_type="solid"),
        }
        for col_idx, (_key, label, group) in enumerate(management_columns[:17], start=1):
            cell = management_ws.cell(3, col_idx, label)
            cell.font = Font(name="Calibri", size=9, bold=True, color="022C22")
            cell.fill = management_group_fills[group]
            if col_idx == 4:
                cell.fill = PatternFill(start_color="08AADC", end_color="08AADC", fill_type="solid")
                cell.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
            cell.border = header_border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        comments_header = management_ws.cell(3, 19, "Comments")
        comments_header.font = Font(name="Calibri", size=9, bold=True, color="022C22")
        comments_header.fill = comments_fill
        comments_header.border = header_border
        comments_header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        text_columns = {1, 2, 3, 4, 19}
        for row_offset, (record, row_values) in enumerate(zip(management_rows, management_data), start=4):
            designation = str(row_values[0] or "").strip().upper()
            is_vacant = bool(record.get("is_vacant"))
            designation_style = designation_row_styles.get(designation)
            row_fill = vacant_fill if is_vacant else (designation_style[0] if designation_style else (even_fill if row_offset % 2 == 0 else None))
            row_text_color = "9F1239" if is_vacant else (designation_style[1] if designation_style else "022C22")
            for col_idx, value in enumerate(row_values[:17], start=1):
                cell = management_ws.cell(row_offset, col_idx, value)
                cell.font = Font(
                    name="Calibri",
                    size=9,
                    bold=designation not in {"SPO", "MIO", ""} and col_idx == 1,
                    color=row_text_color,
                )
                if row_fill:
                    cell.fill = row_fill
                cell.border = cell_border
                cell.alignment = Alignment(
                    horizontal="left" if col_idx in text_columns else "right",
                    vertical="center",
                )
            comments_cell = management_ws.cell(row_offset, 19, row_values[17])
            comments_cell.font = Font(name="Calibri", size=9, color=row_text_color)
            comments_cell.border = cell_border
            comments_cell.alignment = Alignment(horizontal="left", vertical="center")
            if row_fill:
                comments_cell.fill = row_fill

        management_footer_row = len(management_data) + 4

        def numeric_values(index: int) -> list[float]:
            values: list[float] = []
            for row_values in management_data:
                value = row_values[index]
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    values.append(float(value))
            return values

        management_totals: list[object] = ["TOTAL SUMMARY", "", "", ""]
        for index in range(4, 17):
            values = numeric_values(index)
            if index == 16:
                management_totals.append(sum(values) / len(values) if values else "-")
            else:
                management_totals.append(sum(values) if values else "-")
        management_totals.append("")
        for col_idx, value in enumerate(management_totals[:17], start=1):
            cell = management_ws.cell(management_footer_row, col_idx, value)
            cell.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
            cell.fill = footer_fill
            cell.border = header_border
            cell.alignment = Alignment(
                horizontal="left" if col_idx <= 4 else "right",
                vertical="center",
            )
        spacer_cell = management_ws.cell(management_footer_row, 18, "")
        spacer_cell.fill = footer_fill
        comments_footer = management_ws.cell(management_footer_row, 19, management_totals[17])
        comments_footer.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        comments_footer.fill = footer_fill
        comments_footer.border = header_border
        management_ws.row_dimensions[management_footer_row].height = 22

        management_widths = [8, 8, 14, 20, 6, 8, 12, 10, 10, 10, 10, 10, 10, 10, 11, 10, 12, 2, 16]
        for col_idx, width in enumerate(management_widths, start=1):
            management_ws.column_dimensions[get_column_letter(col_idx)].width = width
        management_ws.column_dimensions["R"].hidden = True
        management_ws.freeze_panes = "A4"
        management_ws.auto_filter.ref = f"A3:S{max(3, management_footer_row - 1)}"
        management_ws.sheet_view.zoomScale = 90
        management_ws.sheet_view.zoomScaleNormal = 90
        management_ws.sheet_properties.pageSetUpPr.fitToPage = True
        management_ws.page_setup.orientation = "landscape"
        management_ws.page_setup.fitToWidth = 1
        management_ws.page_setup.fitToHeight = 0

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    wb.close()

    export_filename = f"{cfg['filename_prefix']}_{formatted_date}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(export_filename)}"},
    )


# Call and attendance activity for both teams comes from one shared source.
LIVE_CALL_SOURCE_CONFIGS: dict[str, dict[str, str]] = {
    "FMCG": {
        "team_key": "FMCG",
        "table": "call_attendence",
        "label": "FMCG call attendance",
    },
    "PHARMA": {
        "team_key": "PHARMA",
        "table": "call_attendence",
        "label": "PHARMA call attendance",
    },
}

# Source labels such as "FMCG Team 02" all belong to the one FMCG report.
# Normalize punctuation and whitespace first so the match is stable for values
# such as FMCG-1, FMCG 1, and FMCG Team 01.
ACTIVITY_TEAM_ALIASES: dict[str, tuple[str, ...]] = {
    "FMCG": (
        "FMCG", "FMCG1", "FMCG2", "FMCG3",
        "FMCGTEAM1", "FMCGTEAM2", "FMCGTEAM3",
        "FMCGTEAM01", "FMCGTEAM02", "FMCGTEAM03",
    ),
    "PHARMA": (
        "PHARMA", "PHARMATEAM", "PHARMATEAM01", "PHARMATEAM02", "PHARMATEAM03",
    ),
}


def get_live_call_source_config(team_key: str) -> dict[str, str]:
    normalized = (team_key or "").strip().upper()
    if normalized not in LIVE_CALL_SOURCE_CONFIGS:
        raise HTTPException(status_code=400, detail=f"Unsupported live call source key: '{team_key}'.")
    return LIVE_CALL_SOURCE_CONFIGS[normalized]


def activity_team_sql_condition(column_expression: str, team_key: str) -> str:
    """Build the shared-table team condition using the approved report aliases."""
    normalized_team = (team_key or "").strip().upper()
    aliases = ACTIVITY_TEAM_ALIASES.get(normalized_team, (normalized_team,))
    alias_sql = ", ".join(f"'{alias}'" for alias in aliases)
    normalized_column = (
        f"REGEXP_REPLACE(UPPER(COALESCE(CAST({column_expression} AS text), '')), "
        "'[^A-Z0-9]', '', 'g')"
    )
    return f"{normalized_column} IN ({alias_sql})"


def normalize_activity_person_name(value: object) -> str:
    """Normalize a source name for a cautious one-to-one legacy ID fallback."""
    without_suffix = re.sub(r"\([^)]*\)", " ", display_value(value))
    return re.sub(r"\s+", " ", without_suffix).strip().casefold()


def quote_pg_identifier(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'


def inspect_live_call_source_columns(connection, table_name: str) -> list[dict[str, Any]]:
    rows = connection.execute(
        text("""
            SELECT
                column_name,
                ordinal_position,
                data_type,
                udt_name,
                character_maximum_length,
                is_nullable
            FROM information_schema.columns
            WHERE table_schema = 'public'
              AND table_name = :table_name
            ORDER BY ordinal_position
        """),
        {"table_name": table_name},
    ).mappings().all()
    return [
        {
            "name": row["column_name"],
            "ordinalPosition": row["ordinal_position"],
            "dataType": row["data_type"],
            "udtName": row["udt_name"],
            "maxLength": row["character_maximum_length"],
            "nullable": row["is_nullable"] == "YES",
        }
        for row in rows
    ]


def serialize_live_call_value(value: object) -> object:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None:
        return None
    if isinstance(value, float) and pd.isna(value):
        return None
    return value


def fetch_live_call_source_data(
    team_key: str,
    employee: str = "",
    position_code: str = "",
    region: str = "",
    date_from: str | None = None,
    date_to: str | None = None,
    offset: int = 0,
    limit: int = 50,
) -> dict[str, Any]:
    cfg = get_live_call_source_config(team_key)
    clean_offset = max(offset, 0)
    clean_limit = min(max(limit, 1), 100)

    parsed_from: date | None = None
    parsed_to: date | None = None
    try:
        if date_from:
            parsed_from = date.fromisoformat(date_from.strip())
        if date_to:
            parsed_to = date.fromisoformat(date_to.strip())
    except ValueError as error:
        raise HTTPException(status_code=400, detail="date_from and date_to must use YYYY-MM-DD format.") from error
    if parsed_from and parsed_to and parsed_from > parsed_to:
        raise HTTPException(status_code=400, detail="date_from cannot be later than date_to.")

    table_name = cfg["table"]
    table_ref = f"public.{quote_pg_identifier(table_name)}"
    try:
        with get_database_engine().connect() as connection:
            column_meta = inspect_live_call_source_columns(connection, table_name)
            if not column_meta:
                raise HTTPException(
                    status_code=503,
                    detail=f"Database table public.{table_name} was not found or has no columns.",
                )

            column_names = [column["name"] for column in column_meta]
            available = {column.casefold(): column for column in column_names}
            data_types = {column["name"].casefold(): column["dataType"] for column in column_meta}
            where_clauses: list[str] = []
            params: dict[str, Any] = {}
            unsupported_filters: list[str] = []

            team_column = available.get("team")
            if team_column:
                where_clauses.append(
                    activity_team_sql_condition(quote_pg_identifier(team_column), cfg["team_key"])
                )
            else:
                unsupported_filters.append("team")

            employee_columns = [available[name] for name in ("ff_id", "ff_name") if name in available]
            if employee.strip():
                if employee_columns:
                    employee_parts = [
                        f"CAST({quote_pg_identifier(column)} AS text) ILIKE :employee"
                        for column in employee_columns
                    ]
                    where_clauses.append(f"({' OR '.join(employee_parts)})")
                    params["employee"] = f"%{employee.strip()}%"
                else:
                    unsupported_filters.append("employee")

            position_columns = [
                available[name]
                for name in ("position_code", "position_code_secondary")
                if name in available
            ]
            if position_code.strip():
                if position_columns:
                    position_parts = [
                        f"BTRIM(CAST({quote_pg_identifier(column)} AS text)) = :position_code"
                        for column in position_columns
                    ]
                    where_clauses.append(f"({' OR '.join(position_parts)})")
                    params["position_code"] = position_code.strip()
                else:
                    unsupported_filters.append("position_code")

            region_columns = [
                available[name]
                for name in ("region", "region_name", "team_region", "team_region_name")
                if name in available
            ]
            if region.strip():
                if region_columns:
                    region_parts = [
                        f"BTRIM(CAST({quote_pg_identifier(column)} AS text)) = :region"
                        for column in region_columns
                    ]
                    where_clauses.append(f"({' OR '.join(region_parts)})")
                    params["region"] = region.strip()
                else:
                    unsupported_filters.append("region")

            date_column = available.get("call_date") or available.get("date")
            if date_column:
                date_type = data_types.get(date_column.casefold(), "")
                date_expression = quote_pg_identifier(date_column)
                if date_type not in {"date", "timestamp without time zone", "timestamp with time zone"}:
                    date_expression = f"CAST({date_expression} AS date)"
                if parsed_from:
                    where_clauses.append(f"{date_expression} >= :date_from")
                    params["date_from"] = parsed_from
                if parsed_to:
                    where_clauses.append(f"{date_expression} <= :date_to")
                    params["date_to"] = parsed_to
            elif date_from or date_to:
                unsupported_filters.append("date")

            where_sql = f"WHERE {' AND '.join(where_clauses)}" if where_clauses else ""
            count_query = text(f"SELECT COUNT(*) FROM {table_ref} {where_sql}")
            total = int(connection.execute(count_query, params).scalar_one())

            select_columns = ", ".join(
                f"{quote_pg_identifier(column)} AS {quote_pg_identifier(column)}"
                for column in column_names
            )
            data_query = text(f"""
                SELECT {select_columns}
                FROM {table_ref}
                {where_sql}
                ORDER BY
                    {quote_pg_identifier(date_column)} DESC NULLS LAST,
                    {quote_pg_identifier(available.get('ff_id') or column_names[0])} ASC NULLS LAST
                LIMIT :limit_value OFFSET :offset_value
            """) if date_column else text(f"""
                SELECT {select_columns}
                FROM {table_ref}
                {where_sql}
                ORDER BY {quote_pg_identifier(available.get('ff_id') or column_names[0])} ASC NULLS LAST
                LIMIT :limit_value OFFSET :offset_value
            """)
            data_params = {**params, "limit_value": clean_limit, "offset_value": clean_offset}
            rows = [
                {column: serialize_live_call_value(row[column]) for column in column_names}
                for row in connection.execute(data_query, data_params).mappings().all()
            ]
    except SQLAlchemyError as error:
        logger.exception("Unable to retrieve %s from public.%s", cfg["label"], table_name, exc_info=error)
        raise HTTPException(
            status_code=503,
            detail=f"The {cfg['label']} database source is temporarily unavailable.",
        ) from error

    return {
        "team": cfg["team_key"],
        "table": f"public.{table_name}",
        "columns": column_meta,
        "rows": rows,
        "total": total,
        "offset": clean_offset,
        "limit": clean_limit,
        "hasNext": clean_offset + len(rows) < total,
        "filterCapabilities": {
            "employee": bool(employee_columns),
            "position_code": bool(position_columns),
            "region": bool(region_columns),
            "date": bool(date_column),
        },
        "unsupportedFilters": sorted(set(unsupported_filters)),
    }


def fetch_call_attendence_activity_by_position(
    team_key: str,
    selected_date: date,
) -> dict[str, dict[str, Any]]:
    """Return call and attendance metrics from the shared call_attendence table."""
    cfg = get_team_summary_config(team_key)
    month_start = selected_date.replace(day=1)

    activity_by_employee: dict[str, dict[str, Any]] = {}
    staff_query = text(f"""
        SELECT
            position_code,
            sas_id,
            designation_name,
            actual_field_force,
            remarks,
            e_code,
            employe_name,
            payroll_type,
            doj
        FROM public.ff_master_staff
        WHERE {cfg['where_clause']}
    """)
    activity_query = text(f"""
        SELECT ff_id, ff_name, call_date, call_type, total_sum_of_done
        FROM public.call_attendence
        WHERE {activity_team_sql_condition(quote_pg_identifier('team'), cfg['team_key'])}
          AND call_date >= :month_start
          AND call_date <= :selected_date
    """)

    try:
        with get_database_engine().connect() as connection:
            activity_rows = connection.execute(
                activity_query,
                {
                    "month_start": month_start,
                    "selected_date": selected_date,
                },
            ).mappings().all()
            staff_rows = connection.execute(staff_query).mappings().all()
    except SQLAlchemyError as error:
        logger.exception("Unable to retrieve %s call attendance activity", cfg["team_key"], exc_info=error)
        raise HTTPException(
            status_code=503,
            detail=f"{cfg['team_key']} call and attendance data is temporarily unavailable.",
        ) from error

    for row in activity_rows:
        employee_id = normalize_activity_employee_id(row["ff_id"])
        source_date = row["call_date"]
        if not employee_id or not isinstance(source_date, date):
            continue
        activity = activity_by_employee.setdefault(
            employee_id,
            {"attendance_days": set(), "call_days": {}, "source_names": set()},
        )
        source_name = normalize_activity_person_name(row["ff_name"])
        if source_name:
            activity["source_names"].add(source_name)
        call_type = display_value(row["call_type"]).strip().upper()
        done_value = pd.to_numeric(row["total_sum_of_done"], errors="coerce")
        done = int(done_value) if pd.notna(done_value) else 0
        if done <= 0:
            continue
        if call_type == "EXCEPTION":
            activity["attendance_days"].add(source_date)
        elif call_type == "CALL":
            activity["call_days"][source_date] = activity["call_days"].get(source_date, 0) + done

    def is_active(row: Any) -> bool:
        actual_field_force = display_value(row["actual_field_force"]).strip().upper()
        return bool(
            actual_field_force not in {"", "0", "0.0", "FALSE", "NO"}
            and display_value(row["remarks"]).strip().casefold() != "un-mapped"
            and display_value(row["e_code"]).strip().upper() not in {"UN-DEFINED", "VACANT"}
            and display_value(row["payroll_type"]).strip().upper() != "VACANT"
            and display_value(row["doj"]).strip().upper() != "RESIGNED"
            and "UNMAPPED" not in display_value(row["employe_name"]).strip().upper()
            and "VACANT" not in display_value(row["employe_name"]).strip().upper()
            and not display_value(row["position_code"]).strip().upper().endswith("-00")
            and bool(normalize_activity_employee_id(row["sas_id"]))
        )

    active_staff_ids: set[str] = set()
    active_staff_ids_by_name: dict[str, set[str]] = {}
    for row in staff_rows:
        if not is_active(row):
            continue
        staff_id = normalize_activity_employee_id(row["sas_id"])
        staff_name = normalize_activity_person_name(row["employe_name"])
        if not staff_id:
            continue
        active_staff_ids.add(staff_id)
        if staff_name:
            active_staff_ids_by_name.setdefault(staff_name, set()).add(staff_id)

    # A live-call file can retain a prior FF.Id after an employee moves to a
    # new position.  Only remap that legacy activity when its source name maps
    # to exactly one active staff ID; ambiguous names are intentionally ignored.
    for source_id, source_activity in list(activity_by_employee.items()):
        if source_id in active_staff_ids:
            continue
        candidate_ids = {
            staff_id
            for source_name in source_activity["source_names"]
            for staff_id in active_staff_ids_by_name.get(source_name, set())
        }
        if len(candidate_ids) != 1:
            continue
        target_id = next(iter(candidate_ids))
        target_activity = activity_by_employee.setdefault(
            target_id,
            {"attendance_days": set(), "call_days": {}, "source_names": set()},
        )
        target_activity["attendance_days"].update(source_activity["attendance_days"])
        for activity_date, call_count in source_activity["call_days"].items():
            target_activity["call_days"][activity_date] = (
                target_activity["call_days"].get(activity_date, 0) + call_count
            )
        target_activity["source_names"].update(source_activity["source_names"])
        logger.info(
            "Mapped legacy %s activity to active %s by unique employee name.",
            source_id,
            target_id,
        )

    activity_by_position: dict[str, dict[str, Any]] = {}
    days_to_report = range(1, selected_date.day + 1)
    for row in staff_rows:
        if not is_active(row):
            continue
        position_key = display_value(row["position_code"]).strip().casefold()
        employee_id = normalize_activity_employee_id(row["sas_id"])
        if not position_key or not employee_id:
            continue
        employee_activity = activity_by_employee.get(
            employee_id,
            {"attendance_days": set(), "call_days": {}, "source_names": set()},
        )
        call_days = {
            day: int(employee_activity["call_days"].get(
                date(selected_date.year, selected_date.month, day),
                0,
            ))
            for day in days_to_report
        }
        calls_total = sum(call_days.values())
        attendance_days = employee_activity["attendance_days"]
        activity_by_position[position_key] = {
            "employee_id": employee_id,
            "today_att": int(selected_date in attendance_days),
            "mtd_att": int(bool(attendance_days)),
            "today_calls": int(call_days.get(selected_date.day, 0) > 0),
            "mtd_calls": int(calls_total > 0),
            "attendance_days": attendance_days,
            "call_days": call_days,
            "calls_total": calls_total,
        }
    return activity_by_position


def fetch_fmcg_activity_by_position(selected_date: date) -> dict[str, dict[str, Any]]:
    """Backward-compatible wrapper for the shared call attendance source."""
    return fetch_call_attendence_activity_by_position("FMCG", selected_date)


def live_call_source_endpoint(
    team_key: str,
    employee: str = "",
    position_code: str = "",
    region: str = "",
    date_from: str | None = None,
    date_to: str | None = None,
    offset: int = 0,
    limit: int = 50,
) -> dict[str, Any]:
    return fetch_live_call_source_data(
        team_key=team_key,
        employee=employee,
        position_code=position_code,
        region=region,
        date_from=date_from,
        date_to=date_to,
        offset=offset,
        limit=limit,
    )


@app.get("/api/fmcg/ikon-mtd-call-source")
def fmcg_ikon_mtd_call_source(
    employee: str = "",
    position_code: str = "",
    region: str = "",
    date_from: str | None = None,
    date_to: str | None = None,
    offset: int = 0,
    limit: int = 50,
) -> dict[str, Any]:
    return live_call_source_endpoint("FMCG", employee, position_code, region, date_from, date_to, offset, limit)


@app.get("/api/pharma/ikon-mtd-call-source")
def pharma_ikon_mtd_call_source(
    employee: str = "",
    position_code: str = "",
    region: str = "",
    date_from: str | None = None,
    date_to: str | None = None,
    offset: int = 0,
    limit: int = 50,
) -> dict[str, Any]:
    return live_call_source_endpoint("PHARMA", employee, position_code, region, date_from, date_to, offset, limit)


@app.get("/api/teams/{team_slug}/ikon-mtd-call-source")
def team_ikon_mtd_call_source(
    team_slug: str,
    employee: str = "",
    position_code: str = "",
    region: str = "",
    date_from: str | None = None,
    date_to: str | None = None,
    offset: int = 0,
    limit: int = 50,
) -> dict[str, Any]:
    return live_call_source_endpoint(team_slug, employee, position_code, region, date_from, date_to, offset, limit)


@app.get("/api/fmcg/ikon-mtd-summary")
def fmcg_ikon_mtd_summary(
    report_date: str | None = None,
    search: str = "",
    designation: str = "",
    zone: str = "",
    region: str = "",
    base_town: str = "",
    dept_territory_code: str = "",
    position_code: str = "",
    sas_id: str = "",
    employee_name: str = "",
    emp_type: str = "",
    emp_code: str = "",
    payroll_status: str = "",
    doj: str = "",
    sort_by: str = "",
    sort_dir: str = "asc",
) -> dict[str, Any]:
    return fetch_team_summary_data(
        team_key="FMCG",
        report_date=report_date,
        search=search,
        designation=designation,
        zone=zone,
        region=region,
        base_town=base_town,
        dept_territory_code=dept_territory_code,
        position_code=position_code,
        sas_id=sas_id,
        employee_name=employee_name,
        emp_type=emp_type,
        emp_code=emp_code,
        payroll_status=payroll_status,
        doj=doj,
        sort_by=sort_by,
        sort_dir=sort_dir,
    )


@app.get("/api/fmcg/ikon-mtd-summary/export")
def export_fmcg_ikon_mtd_summary(
    report_date: str | None = None,
    search: str = "",
    designation: str = "",
    zone: str = "",
    region: str = "",
    base_town: str = "",
    dept_territory_code: str = "",
    position_code: str = "",
    sas_id: str = "",
    employee_name: str = "",
    emp_type: str = "",
    emp_code: str = "",
    payroll_status: str = "",
    doj: str = "",
    sort_by: str = "",
    sort_dir: str = "asc",
) -> StreamingResponse:
    return export_team_ikon_mtd_summary_workbook(
        team_key="FMCG",
        report_date=report_date,
        search=search,
        designation=designation,
        zone=zone,
        region=region,
        base_town=base_town,
        dept_territory_code=dept_territory_code,
        position_code=position_code,
        sas_id=sas_id,
        employee_name=employee_name,
        emp_type=emp_type,
        emp_code=emp_code,
        payroll_status=payroll_status,
        doj=doj,
        sort_by=sort_by,
        sort_dir=sort_dir,
    )


@app.get("/api/pharma/ikon-mtd-summary")
def pharma_ikon_mtd_summary(
    report_date: str | None = None,
    search: str = "",
    designation: str = "",
    zone: str = "",
    region: str = "",
    base_town: str = "",
    dept_territory_code: str = "",
    position_code: str = "",
    sas_id: str = "",
    employee_name: str = "",
    emp_type: str = "",
    emp_code: str = "",
    payroll_status: str = "",
    doj: str = "",
    sort_by: str = "",
    sort_dir: str = "asc",
) -> dict[str, Any]:
    return fetch_team_summary_data(
        team_key="PHARMA",
        report_date=report_date,
        search=search,
        designation=designation,
        zone=zone,
        region=region,
        base_town=base_town,
        dept_territory_code=dept_territory_code,
        position_code=position_code,
        sas_id=sas_id,
        employee_name=employee_name,
        emp_type=emp_type,
        emp_code=emp_code,
        payroll_status=payroll_status,
        doj=doj,
        sort_by=sort_by,
        sort_dir=sort_dir,
    )


@app.get("/api/pharma/ikon-mtd-summary/export")
def export_pharma_ikon_mtd_summary(
    report_date: str | None = None,
    search: str = "",
    designation: str = "",
    zone: str = "",
    region: str = "",
    base_town: str = "",
    dept_territory_code: str = "",
    position_code: str = "",
    sas_id: str = "",
    employee_name: str = "",
    emp_type: str = "",
    emp_code: str = "",
    payroll_status: str = "",
    doj: str = "",
    sort_by: str = "",
    sort_dir: str = "asc",
) -> StreamingResponse:
    return export_team_ikon_mtd_summary_workbook(
        team_key="PHARMA",
        report_date=report_date,
        search=search,
        designation=designation,
        zone=zone,
        region=region,
        base_town=base_town,
        dept_territory_code=dept_territory_code,
        position_code=position_code,
        sas_id=sas_id,
        employee_name=employee_name,
        emp_type=emp_type,
        emp_code=emp_code,
        payroll_status=payroll_status,
        doj=doj,
        sort_by=sort_by,
        sort_dir=sort_dir,
    )


@app.get("/api/teams/{team_slug}/ikon-mtd-summary")
def team_ikon_mtd_summary(
    team_slug: str,
    report_date: str | None = None,
    search: str = "",
    designation: str = "",
    zone: str = "",
    region: str = "",
    base_town: str = "",
    dept_territory_code: str = "",
    position_code: str = "",
    sas_id: str = "",
    employee_name: str = "",
    emp_type: str = "",
    emp_code: str = "",
    payroll_status: str = "",
    doj: str = "",
    sort_by: str = "",
    sort_dir: str = "asc",
) -> dict[str, Any]:
    return fetch_team_summary_data(
        team_key=team_slug,
        report_date=report_date,
        search=search,
        designation=designation,
        zone=zone,
        region=region,
        base_town=base_town,
        dept_territory_code=dept_territory_code,
        position_code=position_code,
        sas_id=sas_id,
        employee_name=employee_name,
        emp_type=emp_type,
        emp_code=emp_code,
        payroll_status=payroll_status,
        doj=doj,
        sort_by=sort_by,
        sort_dir=sort_dir,
    )


@app.get("/api/teams/{team_slug}/ikon-mtd-summary/export")
def export_team_ikon_mtd_summary(
    team_slug: str,
    report_date: str | None = None,
    search: str = "",
    designation: str = "",
    zone: str = "",
    region: str = "",
    base_town: str = "",
    dept_territory_code: str = "",
    position_code: str = "",
    sas_id: str = "",
    employee_name: str = "",
    emp_type: str = "",
    emp_code: str = "",
    payroll_status: str = "",
    doj: str = "",
    sort_by: str = "",
    sort_dir: str = "asc",
) -> StreamingResponse:
    return export_team_ikon_mtd_summary_workbook(
        team_key=team_slug,
        report_date=report_date,
        search=search,
        designation=designation,
        zone=zone,
        region=region,
        base_town=base_town,
        dept_territory_code=dept_territory_code,
        position_code=position_code,
        sas_id=sas_id,
        employee_name=employee_name,
        emp_type=emp_type,
        emp_code=emp_code,
        payroll_status=payroll_status,
        doj=doj,
        sort_by=sort_by,
        sort_dir=sort_dir,
    )


# ---------------------------------------------------------------------------
# FF Master Staff – PostgreSQL Database Management (CRUD)
# ---------------------------------------------------------------------------

FF_STAFF_WRITE_TOKEN = os.getenv("FF_EDITOR_WRITE_TOKEN") or secrets.token_urlsafe(24)

FF_STAFF_COLUMN_META: list[tuple[str, str, str]] = [
    ("position_code", "Position Code", "Identity & Position"),
    ("position_code_secondary", "Secondary Position Code", "Identity & Position"),
    ("dept_territory_code", "Dept. Territory Code", "Identity & Position"),
    ("sas_id", "SAS ID", "Identity & Position"),
    ("sas_upload", "SAS Upload", "Identity & Position"),
    ("hawk_emp_id", "Hawk Emp ID", "Identity & Position"),
    ("employe_name", "Employee Name", "Employee Details"),
    ("last_employee", "Last Employee", "Employee Details"),
    ("e_code", "Emp. Code (E-Code)", "Employee Details"),
    ("doj", "Date of Joining (DOJ)", "Employee Details"),
    ("nature_type", "Nature Type", "Employee Details"),
    ("payroll_type", "Payroll Type", "Employee Details"),
    ("check_status", "Check Status", "Employee Details"),
    ("desg_id", "Designation ID", "Organization & Hierarchy"),
    ("designation_name", "Designation Name", "Organization & Hierarchy"),
    ("level", "Level", "Organization & Hierarchy"),
    ("dept_id", "Department ID", "Organization & Hierarchy"),
    ("depatment_name", "Department Name", "Organization & Hierarchy"),
    ("team_id", "Team ID", "Organization & Hierarchy"),
    ("team_name", "Team Name", "Organization & Hierarchy"),
    ("report_to_position_code", "Report To Position Code", "Organization & Hierarchy"),
    ("report_to_name", "Report To Name", "Organization & Hierarchy"),
    ("zone_id", "Zone ID", "Geography & Field Status"),
    ("zone_name", "Zone Name", "Geography & Field Status"),
    ("sub_zone", "Sub Zone", "Geography & Field Status"),
    ("alias_zone", "Alias Zone", "Geography & Field Status"),
    ("team_region_id", "Team Region ID", "Geography & Field Status"),
    ("region_name", "Region Name", "Geography & Field Status"),
    ("alias_region", "Alias Region", "Geography & Field Status"),
    ("province", "Province", "Geography & Field Status"),
    ("base_town", "Base Town", "Geography & Field Status"),
    ("base_name", "Base Name", "Geography & Field Status"),
    ("actual_field_force", "Actual Field Force", "Geography & Field Status"),
    ("remarks", "Remarks", "Geography & Field Status"),
]

FF_STAFF_LABEL_MAP: dict[str, str] = {col: label for col, label, _ in FF_STAFF_COLUMN_META}
FF_STAFF_SECTION_MAP: dict[str, str] = {col: section for col, _, section in FF_STAFF_COLUMN_META}
FF_STAFF_ALLOWED_COLUMNS: tuple[str, ...] = tuple(col for col, _, _ in FF_STAFF_COLUMN_META)
FF_STAFF_AUTOCOMPLETE_COLUMNS: tuple[str, ...] = (
    "team_name",
    "designation_name",
    "zone_name",
    "sub_zone",
    "region_name",
    "province",
    "base_town",
    "nature_type",
    "payroll_type",
    "check_status",
    "depatment_name",
    "level",
    "remarks",
)


class FFMasterStaffPayload(BaseModel):
    model_config = ConfigDict(extra="forbid")

    values: dict[str, Any]
    expectedVersion: str | None = None


def verify_ff_write_authorization(request: Request) -> None:
    """Verify write/delete authorization token from request headers."""
    client_token = (request.headers.get("X-FF-Write-Token") or "").strip()
    configured_api_key = (os.getenv("FF_EDITOR_API_KEY") or "").strip()
    if configured_api_key:
        provided_key = (request.headers.get("X-FF-Admin-Key") or client_token).strip()
        if not secrets.compare_digest(provided_key, configured_api_key):
            raise HTTPException(
                status_code=403,
                detail="Write/delete operation requires a valid authorization key.",
            )
        return
    if not client_token or not secrets.compare_digest(client_token, FF_STAFF_WRITE_TOKEN):
        raise HTTPException(
            status_code=403,
            detail="Missing or invalid write authorization token. Refresh the page and try again.",
        )


def inspect_ff_master_staff_columns(connection) -> list[dict[str, Any]]:
    """Dynamically inspect public.ff_master_staff column metadata from information_schema."""
    col_rows = connection.execute(
        text("""
            SELECT
                column_name,
                ordinal_position,
                data_type,
                udt_name,
                character_maximum_length,
                is_nullable,
                column_default
            FROM information_schema.columns
            WHERE table_schema = 'public'
              AND table_name = 'ff_master_staff'
            ORDER BY ordinal_position
        """)
    ).mappings().all()

    if not col_rows:
        raise HTTPException(
            status_code=500,
            detail="Table public.ff_master_staff was not found in the connected PostgreSQL database.",
        )

    constraint_rows = connection.execute(
        text("""
            SELECT
                tc.constraint_name,
                tc.constraint_type,
                kcu.column_name
            FROM information_schema.table_constraints tc
            LEFT JOIN information_schema.key_column_usage kcu
              ON tc.constraint_name = kcu.constraint_name
             AND tc.table_schema = kcu.table_schema
             AND tc.table_name = kcu.table_name
            WHERE tc.table_schema = 'public'
              AND tc.table_name = 'ff_master_staff'
        """)
    ).mappings().all()

    pk_columns = {
        row["column_name"]
        for row in constraint_rows
        if row["constraint_type"] == "PRIMARY KEY" and row["column_name"]
    }
    unique_columns = {
        row["column_name"]
        for row in constraint_rows
        if row["constraint_type"] in {"PRIMARY KEY", "UNIQUE"} and row["column_name"]
    }
    # Verified business primary identifier across populated records:
    unique_columns.add("position_code")

    columns_meta: list[dict[str, Any]] = []
    for row in col_rows:
        col_name = str(row["column_name"])
        data_type = str(row["data_type"])
        max_len = int(row["character_maximum_length"]) if row["character_maximum_length"] is not None else None
        is_identifier = col_name == "position_code" or col_name in pk_columns
        columns_meta.append({
            "name": col_name,
            "label": FF_STAFF_LABEL_MAP.get(col_name, col_name.replace("_", " ").title()),
            "ordinalPosition": int(row["ordinal_position"]),
            "dataType": data_type,
            "udtName": str(row["udt_name"]),
            "maxLength": max_len,
            "nullable": (row["is_nullable"] == "YES") and not is_identifier,
            "required": is_identifier,
            "isPrimaryKey": is_identifier,
            "isUnique": col_name in unique_columns,
            "inputType": "number" if data_type in {"integer", "bigint", "smallint", "numeric"} else "text",
            "section": FF_STAFF_SECTION_MAP.get(col_name, "Additional Fields"),
        })
    return columns_meta


def compute_ff_row_version(row_mapping: dict[str, Any], column_names: list[str]) -> str:
    """Compute a deterministic hash over all column values for optimistic concurrency control."""
    canonical = "|".join(
        f"{col}={'' if row_mapping.get(col) is None else str(row_mapping.get(col))}"
        for col in column_names
    )
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()[:16]


def serialize_ff_staff_row(row_mapping: dict[str, Any], column_names: list[str]) -> dict[str, Any]:
    """Serialize one database row preserving NULL as None and attaching _row_id and _version."""
    pos_code = row_mapping.get("position_code")
    pos_str = str(pos_code).strip() if pos_code is not None else ""
    ctid_str = str(row_mapping.get("_ctid") or "")
    row_id = pos_str if pos_str else f"__ctid__:{ctid_str}"

    record: dict[str, Any] = {
        "_row_id": row_id,
        "_has_position_code": bool(pos_str),
        "_version": compute_ff_row_version(row_mapping, column_names),
    }
    for col in column_names:
        val = row_mapping.get(col)
        if val is None:
            record[col] = None
        elif isinstance(val, int):
            record[col] = val
        else:
            record[col] = display_value(val)
    return record


def validate_and_clean_ff_payload(
    raw_values: dict[str, Any],
    columns_meta: list[dict[str, Any]],
) -> dict[str, Any]:
    """Validate incoming field values against PostgreSQL schema metadata."""
    meta_by_name = {col["name"]: col for col in columns_meta}
    unknown = [key for key in raw_values if key not in meta_by_name]
    if unknown:
        raise HTTPException(
            status_code=400,
            detail=f"Unknown column(s) submitted: {', '.join(unknown)}",
        )

    cleaned: dict[str, Any] = {}
    for col_meta in columns_meta:
        col_name = col_meta["name"]
        raw_val = raw_values.get(col_name)

        if raw_val is None:
            cleaned_val = None
        elif isinstance(raw_val, str):
            trimmed = raw_val.strip()
            cleaned_val = trimmed if trimmed != "" else None
        else:
            trimmed = str(raw_val).strip()
            cleaned_val = trimmed if trimmed != "" else None

        if col_meta["required"] and cleaned_val is None:
            raise HTTPException(
                status_code=400,
                detail=f"'{col_meta['label']}' ({col_name}) is required and cannot be empty.",
            )

        if cleaned_val is not None:
            if col_meta["dataType"] in {"integer", "bigint", "smallint"}:
                try:
                    if isinstance(raw_val, bool):
                        raise ValueError("Boolean is not a valid integer.")
                    cleaned_val = int(str(cleaned_val))
                except ValueError as error:
                    raise HTTPException(
                        status_code=400,
                        detail=f"'{col_meta['label']}' ({col_name}) must be a valid whole number.",
                    ) from error
            else:
                cleaned_val = str(cleaned_val)
                max_len = col_meta["maxLength"]
                if max_len is not None and len(cleaned_val) > max_len:
                    raise HTTPException(
                        status_code=400,
                        detail=(
                            f"'{col_meta['label']}' ({col_name}) exceeds maximum length "
                            f"of {max_len} characters (got {len(cleaned_val)})."
                        ),
                    )

        cleaned[col_name] = cleaned_val

    return cleaned


@app.get("/api/ff-master-staff/schema")
def get_ff_master_staff_schema() -> dict[str, Any]:
    with get_database_engine().connect() as connection:
        columns_meta = inspect_ff_master_staff_columns(connection)
        available_col_names = {col["name"] for col in columns_meta}

        options: dict[str, list[str]] = {}
        for col_name in FF_STAFF_AUTOCOMPLETE_COLUMNS:
            if col_name not in available_col_names:
                continue
            rows = connection.execute(
                text(f"""
                    SELECT DISTINCT BTRIM({col_name}::text) AS val
                    FROM public.ff_master_staff
                    WHERE {col_name} IS NOT NULL
                      AND BTRIM({col_name}::text) <> ''
                    ORDER BY val ASC
                """)
            ).scalars().all()
            options[col_name] = [str(v) for v in rows if v is not None]

        stats_row = connection.execute(
            text("""
                SELECT
                    COUNT(*) FILTER (WHERE position_code IS NOT NULL AND BTRIM(position_code) <> '') AS populated_positions,
                    COUNT(DISTINCT BTRIM(position_code)) FILTER (WHERE position_code IS NOT NULL AND BTRIM(position_code) <> '') AS distinct_positions
                FROM public.ff_master_staff
            """)
        ).mappings().one()

    return {
        "table": "public.ff_master_staff",
        "primaryKey": "position_code",
        "primaryKeyVerifiedUnique": stats_row["populated_positions"] == stats_row["distinct_positions"],
        "writeToken": FF_STAFF_WRITE_TOKEN,
        "requiresAdminKey": bool((os.getenv("FF_EDITOR_API_KEY") or "").strip()),
        "columns": columns_meta,
        "options": options,
    }


@app.get("/api/ff-master-staff")
def list_ff_master_staff(
    page: int = 1,
    page_size: int = 25,
    search: str = "",
    team_name: str = "",
    designation_name: str = "",
    zone_name: str = "",
    region_name: str = "",
    base_town: str = "",
    payroll_type: str = "",
    sort_by: str = "position_code",
    sort_dir: str = "asc",
) -> dict[str, Any]:
    if page < 1:
        page = 1
    if page_size not in {25, 50, 100}:
        page_size = 25

    with get_database_engine().connect() as connection:
        columns_meta = inspect_ff_master_staff_columns(connection)
        column_names = [col["name"] for col in columns_meta]
        allowed_set = set(column_names)

        non_null_Args = ", ".join(column_names)
        where_clauses = [f"num_nonnulls({non_null_Args}) > 0"]
        params: dict[str, Any] = {}

        filter_inputs = {
            "team_name": team_name,
            "designation_name": designation_name,
            "zone_name": zone_name,
            "region_name": region_name,
            "base_town": base_town,
            "payroll_type": payroll_type,
        }
        for col_key, filter_val in filter_inputs.items():
            cleaned_filter = filter_val.strip()
            if cleaned_filter and col_key in allowed_set:
                param_name = f"filter_{col_key}"
                where_clauses.append(f"LOWER(BTRIM(COALESCE({col_key}::text, ''))) = LOWER(:{param_name})")
                params[param_name] = cleaned_filter

        search_query = search.strip()
        if search_query:
            params["search_like"] = f"%{search_query}%"
            searchable_cols = [
                col for col in (
                    "position_code",
                    "employe_name",
                    "last_employee",
                    "e_code",
                    "sas_id",
                    "hawk_emp_id",
                    "designation_name",
                    "team_name",
                    "zone_name",
                    "sub_zone",
                    "region_name",
                    "base_town",
                    "dept_territory_code",
                    "report_to_name",
                    "remarks",
                )
                if col in allowed_set
            ]
            search_or = " OR ".join(f"COALESCE({col}::text, '') ILIKE :search_like" for col in searchable_cols)
            where_clauses.append(f"({search_or})")

        where_sql = " AND ".join(where_clauses)

        sort_column = sort_by if sort_by in allowed_set else "position_code"
        direction = "DESC" if sort_dir.strip().casefold() == "desc" else "ASC"

        total_rows = int(
            connection.execute(
                text(f"SELECT COUNT(*) FROM public.ff_master_staff WHERE {where_sql}"),
                params,
            ).scalar_one()
        )

        total_pages = max((total_rows + page_size - 1) // page_size, 1)
        if page > total_pages:
            page = total_pages
        offset = (page - 1) * page_size

        select_cols = ", ".join(["ctid::text AS _ctid", *column_names])
        order_sql = f"{sort_column} {direction} NULLS LAST, ctid ASC"
        query_sql = text(f"""
            SELECT {select_cols}
            FROM public.ff_master_staff
            WHERE {where_sql}
            ORDER BY {order_sql}
            LIMIT :limit OFFSET :offset
        """)
        rows_raw = connection.execute(
            query_sql,
            {**params, "limit": page_size, "offset": offset},
        ).mappings().all()

    serialized_rows = [serialize_ff_staff_row(dict(row), column_names) for row in rows_raw]
    return {
        "rows": serialized_rows,
        "total": total_rows,
        "page": page,
        "pageSize": page_size,
        "totalPages": total_pages,
        "sortBy": sort_column,
        "sortDir": direction.lower(),
    }


@app.post("/api/ff-master-staff", status_code=201)
def create_ff_master_staff(payload: FFMasterStaffPayload, request: Request) -> dict[str, Any]:
    verify_ff_write_authorization(request)

    try:
        with get_database_engine().begin() as connection:
            columns_meta = inspect_ff_master_staff_columns(connection)
            column_names = [col["name"] for col in columns_meta]
            cleaned = validate_and_clean_ff_payload(payload.values, columns_meta)

            new_pos = cleaned["position_code"]
            duplicate = connection.execute(
                text("""
                    SELECT position_code
                    FROM public.ff_master_staff
                    WHERE position_code IS NOT NULL
                      AND LOWER(BTRIM(position_code)) = LOWER(BTRIM(:pos))
                    LIMIT 1
                """),
                {"pos": new_pos},
            ).scalar_one_or_none()

            if duplicate is not None:
                raise HTTPException(
                    status_code=409,
                    detail=f"A record with Position Code '{new_pos}' already exists in public.ff_master_staff.",
                )

            cols_sql = ", ".join(column_names)
            placeholders_sql = ", ".join(f":{col}" for col in column_names)
            returning_sql = ", ".join(["ctid::text AS _ctid", *column_names])

            inserted = connection.execute(
                text(f"""
                    INSERT INTO public.ff_master_staff ({cols_sql})
                    VALUES ({placeholders_sql})
                    RETURNING {returning_sql}
                """),
                cleaned,
            ).mappings().one()
    except IntegrityError as error:
        raise HTTPException(
            status_code=409,
            detail=f"Database constraint violation while adding record: {error.orig}",
        ) from error

    return {
        "message": f"Employee record '{cleaned['position_code']}' added to public.ff_master_staff.",
        "record": serialize_ff_staff_row(dict(inserted), column_names),
    }


def locate_single_ff_staff_row(connection, record_id: str, column_names: list[str]) -> dict[str, Any]:
    """Locate and lock a single verified row in public.ff_master_staff inside a transaction."""
    decoded_id = unquote(record_id).strip()
    if not decoded_id:
        raise HTTPException(status_code=400, detail="Record identifier is required.")

    select_cols = ", ".join(["ctid::text AS _ctid", *column_names])
    if decoded_id.startswith("__ctid__:"):
        tid_val = decoded_id.removeprefix("__ctid__:").strip()
        if not re.match(r"^\(\d+,\d+\)$", tid_val):
            raise HTTPException(status_code=400, detail="Invalid tuple identifier format.")
        matched = connection.execute(
            text(f"""
                SELECT {select_cols}
                FROM public.ff_master_staff
                WHERE ctid = CAST(:tid AS tid)
                  AND (position_code IS NULL OR BTRIM(position_code) = '')
                FOR UPDATE
            """),
            {"tid": tid_val},
        ).mappings().all()
    else:
        matched = connection.execute(
            text(f"""
                SELECT {select_cols}
                FROM public.ff_master_staff
                WHERE position_code = :pos
                FOR UPDATE
            """),
            {"pos": decoded_id},
        ).mappings().all()

    if not matched:
        raise HTTPException(
            status_code=404,
            detail=f"Record '{decoded_id}' was not found in public.ff_master_staff.",
        )
    if len(matched) > 1:
        raise HTTPException(
            status_code=409,
            detail=(
                f"Identifier '{decoded_id}' matched {len(matched)} records. "
                "Update/delete aborted to prevent unintended multi-row changes."
            ),
        )
    return dict(matched[0])


@app.put("/api/ff-master-staff/{record_id:path}")
def update_ff_master_staff(
    record_id: str,
    payload: FFMasterStaffPayload,
    request: Request,
) -> dict[str, Any]:
    verify_ff_write_authorization(request)

    try:
        with get_database_engine().begin() as connection:
            columns_meta = inspect_ff_master_staff_columns(connection)
            column_names = [col["name"] for col in columns_meta]

            existing_row = locate_single_ff_staff_row(connection, record_id, column_names)
            current_version = compute_ff_row_version(existing_row, column_names)
            if payload.expectedVersion and payload.expectedVersion != current_version:
                raise HTTPException(
                    status_code=409,
                    detail="This record has been modified since you opened it. Please refresh and try again.",
                )

            cleaned = validate_and_clean_ff_payload(payload.values, columns_meta)
            new_pos = cleaned["position_code"]
            old_pos = (existing_row.get("position_code") or "").strip()
            exact_ctid = existing_row["_ctid"]

            if new_pos.casefold() != old_pos.casefold():
                duplicate = connection.execute(
                    text("""
                        SELECT position_code
                        FROM public.ff_master_staff
                        WHERE position_code IS NOT NULL
                          AND LOWER(BTRIM(position_code)) = LOWER(BTRIM(:pos))
                          AND ctid <> CAST(:tid AS tid)
                        LIMIT 1
                    """),
                    {"pos": new_pos, "tid": exact_ctid},
                ).scalar_one_or_none()
                if duplicate is not None:
                    raise HTTPException(
                        status_code=409,
                        detail=f"Another record with Position Code '{new_pos}' already exists.",
                    )

            assignments_sql = ", ".join(f"{col} = :{col}" for col in column_names)
            returning_sql = ", ".join(["ctid::text AS _ctid", *column_names])
            update_result = connection.execute(
                text(f"""
                    UPDATE public.ff_master_staff
                    SET {assignments_sql}
                    WHERE ctid = CAST(:_target_ctid AS tid)
                    RETURNING {returning_sql}
                """),
                {**cleaned, "_target_ctid": exact_ctid},
            ).mappings().all()

            if len(update_result) != 1:
                raise HTTPException(
                    status_code=409,
                    detail="Unexpected affected row count during update; transaction rolled back.",
                )
            updated_row = dict(update_result[0])
    except IntegrityError as error:
        raise HTTPException(
            status_code=409,
            detail=f"Database constraint violation while updating record: {error.orig}",
        ) from error

    return {
        "message": f"Employee record '{cleaned['position_code']}' updated in public.ff_master_staff.",
        "record": serialize_ff_staff_row(updated_row, column_names),
    }


@app.delete("/api/ff-master-staff/{record_id:path}")
def delete_ff_master_staff(
    record_id: str,
    request: Request,
    expected_version: str | None = None,
) -> dict[str, Any]:
    verify_ff_write_authorization(request)

    try:
        with get_database_engine().begin() as connection:
            columns_meta = inspect_ff_master_staff_columns(connection)
            column_names = [col["name"] for col in columns_meta]

            existing_row = locate_single_ff_staff_row(connection, record_id, column_names)
            current_version = compute_ff_row_version(existing_row, column_names)
            if expected_version and expected_version != current_version:
                raise HTTPException(
                    status_code=409,
                    detail="This record was modified since you loaded it. Refresh before deleting.",
                )

            exact_ctid = existing_row["_ctid"]
            deleted_pos = existing_row.get("position_code") or exact_ctid
            delete_result = connection.execute(
                text("""
                    DELETE FROM public.ff_master_staff
                    WHERE ctid = CAST(:_target_ctid AS tid)
                """),
                {"_target_ctid": exact_ctid},
            )
            if delete_result.rowcount != 1:
                raise HTTPException(
                    status_code=409,
                    detail="Unexpected affected row count during delete; transaction rolled back.",
                )
    except IntegrityError as error:
        raise HTTPException(
            status_code=409,
            detail=(
                "Cannot delete this employee record because it is referenced by other database records "
                f"(foreign-key restriction): {error.orig}"
            ),
        ) from error

    return {
        "message": f"Employee record '{deleted_pos}' permanently deleted from public.ff_master_staff.",
        "deletedId": unquote(record_id),
    }


# ---------------------------------------------------------------------------
# IKON Call Performance Dashboard – Analytics & Interactive Tables
# (public.call_attendence & public.live_call_details)
# ---------------------------------------------------------------------------

def _parse_iso_date_optional(value: str | None, param_name: str = "date") -> date | None:
    if not value or not value.strip():
        return None
    try:
        return date.fromisoformat(value.strip())
    except ValueError as error:
        raise HTTPException(
            status_code=400,
            detail=f"Invalid {param_name} '{value}'. Expected YYYY-MM-DD.",
        ) from error


def _week_of_month_bucket(target_date: date) -> tuple[str, str, date, date]:
    """Return a clean weekly bucket within the month (Week 1: 1-7, Week 2: 8-14, etc.)."""
    day = target_date.day
    week_idx = min((day - 1) // 7 + 1, 5)
    days_in_month = calendar.monthrange(target_date.year, target_date.month)[1]
    start_day = (week_idx - 1) * 7 + 1
    end_day = days_in_month if week_idx == 5 else min(week_idx * 7, days_in_month)
    start_dt = date(target_date.year, target_date.month, start_day)
    end_dt = date(target_date.year, target_date.month, end_day)
    month_short = target_date.strftime("%b")
    key = f"{target_date.strftime('%Y-%m')}-W{week_idx}"
    label = f"W{week_idx} ({start_day:02d}–{end_day:02d} {month_short})"
    return key, label, start_dt, end_dt


@app.get("/api/dashboard/performance")
def get_dashboard_performance(
    report_date: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
    team: str = "",
    manager: str = "",
    base_town: str = "",
    employee: str = "",
    call_type: str = "",
    call_nature: str = "",
) -> dict[str, Any]:
    """Aggregate KPI cards, bar/pie/line charts, and employee daily/weekly/monthly call performance."""
    req_report_date = _parse_iso_date_optional(report_date, "report_date")
    req_date_from = _parse_iso_date_optional(date_from, "date_from")
    req_date_to = _parse_iso_date_optional(date_to, "date_to")

    try:
        with get_database_engine().connect() as connection:
            ca_rows = connection.execute(text("""
                SELECT
                    COALESCE(BTRIM(team), '') AS team,
                    COALESCE(BTRIM(ff_id), '') AS ff_id,
                    COALESCE(BTRIM(ff_name), '') AS ff_name,
                    COALESCE(BTRIM(manager), '') AS manager,
                    COALESCE(BTRIM(base_town), '') AS base_town,
                    call_date,
                    UPPER(COALESCE(BTRIM(call_type), '')) AS call_type,
                    COALESCE(total_sum_of_done, 0) AS total_sum_of_done,
                    COALESCE(total_distinct_count_of_ff_id, 0) AS total_distinct_count_of_ff_id
                FROM public.call_attendence
                WHERE call_date IS NOT NULL
            """)).mappings().all()

            lcd_rows = connection.execute(text("""
                SELECT
                    COALESCE(BTRIM(team), '') AS team,
                    COALESCE(BTRIM(ff_id), '') AS ff_id,
                    COALESCE(BTRIM(ff_name), '') AS ff_name,
                    call_date,
                    COALESCE(BTRIM(manager), '') AS manager,
                    COALESCE(BTRIM(base_town), '') AS base_town,
                    COALESCE(BTRIM(pmdc), '') AS pmdc,
                    COALESCE(BTRIM(doctorname), '') AS doctorname,
                    COALESCE(BTRIM(call_nature), '') AS call_nature,
                    UPPER(COALESCE(BTRIM(calltype), '')) AS calltype,
                    COALESCE(BTRIM(ww), '') AS ww,
                    COALESCE(total_sum_noofcalls, 0) AS total_sum_noofcalls,
                    COALESCE(total_sum_noofproducts, 0) AS total_sum_noofproducts,
                    COALESCE(BTRIM(geometer_range), '') AS geometer_range
                FROM public.live_call_details
                WHERE call_date IS NOT NULL
            """)).mappings().all()
    except SQLAlchemyError as error:
        logger.exception("Dashboard performance query failed", exc_info=error)
        raise HTTPException(
            status_code=503,
            detail="Unable to load dashboard analytics from call_attendence and live_call_details.",
        ) from error

    teams_set: set[str] = set()
    managers_set: set[str] = set()
    towns_set: set[str] = set()
    call_types_set: set[str] = set()
    call_natures_set: set[str] = set()
    employees_directory: dict[str, dict[str, str]] = {}
    all_dates: list[date] = []
    ca_call_dates: list[date] = []

    for r in ca_rows:
        dt = r["call_date"]
        if isinstance(dt, datetime):
            dt = dt.date()
        if isinstance(dt, date):
            all_dates.append(dt)
            if r["call_type"] == "CALL" and int(r["total_sum_of_done"] or 0) > 0:
                ca_call_dates.append(dt)
        if r["team"]:
            teams_set.add(r["team"])
        if r["manager"]:
            managers_set.add(r["manager"])
        if r["base_town"]:
            towns_set.add(r["base_town"])
        if r["call_type"]:
            call_types_set.add(r["call_type"])
        emp_key = normalize_activity_employee_id(r["ff_id"])
        if emp_key:
            info = employees_directory.setdefault(emp_key, {
                "employeeKey": emp_key,
                "ff_id": r["ff_id"],
                "ff_name": r["ff_name"],
                "team": r["team"],
                "manager": r["manager"],
                "base_town": r["base_town"],
            })
            if not info["ff_name"] and r["ff_name"]:
                info["ff_name"] = r["ff_name"]
            if not info["team"] and r["team"]:
                info["team"] = r["team"]
            if not info["manager"] and r["manager"]:
                info["manager"] = r["manager"]
            if not info["base_town"] and r["base_town"]:
                info["base_town"] = r["base_town"]

    lcd_dates: list[date] = []
    for r in lcd_rows:
        dt = r["call_date"]
        if isinstance(dt, datetime):
            dt = dt.date()
        if isinstance(dt, date):
            all_dates.append(dt)
            lcd_dates.append(dt)
        if r["team"]:
            teams_set.add(r["team"])
        if r["manager"]:
            managers_set.add(r["manager"])
        if r["base_town"]:
            towns_set.add(r["base_town"])
        if r["calltype"]:
            call_types_set.add(r["calltype"])
        if r["call_nature"]:
            call_natures_set.add(r["call_nature"])
        emp_key = normalize_activity_employee_id(r["ff_id"])
        if emp_key:
            info = employees_directory.setdefault(emp_key, {
                "employeeKey": emp_key,
                "ff_id": r["ff_id"],
                "ff_name": r["ff_name"],
                "team": r["team"],
                "manager": r["manager"],
                "base_town": r["base_town"],
            })
            if not info["ff_name"] and r["ff_name"]:
                info["ff_name"] = r["ff_name"]
            if not info["team"] and r["team"]:
                info["team"] = r["team"]
            if not info["manager"] and r["manager"]:
                info["manager"] = r["manager"]
            if not info["base_town"] and r["base_town"]:
                info["base_town"] = r["base_town"]

    min_available_date = min(all_dates) if all_dates else REPORT_DATE_FALLBACK
    max_available_date = max(all_dates) if all_dates else REPORT_DATE_FALLBACK
    max_lcd_date = max(lcd_dates) if lcd_dates else max_available_date
    selected_date = req_report_date or max_available_date

    # Define Daily, Weekly (7-day rolling ending at selected_date), and Monthly (MTD) windows
    daily_date = selected_date
    lcd_daily_date = min(selected_date, max_lcd_date) if lcd_dates else selected_date
    yesterday_date = selected_date - timedelta(days=1)
    week_start_date = selected_date - timedelta(days=6)
    prev_week_start = week_start_date - timedelta(days=7)
    prev_week_end = week_start_date - timedelta(days=1)
    month_start_date = selected_date.replace(day=1)

    range_start = req_date_from or min(month_start_date, min_available_date)
    range_end = req_date_to or selected_date
    if range_start > range_end:
        range_start, range_end = range_end, range_start

    team_filter = parse_multi_filter_values(team)
    manager_filter = parse_multi_filter_values(manager)
    town_filter = parse_multi_filter_values(base_town)
    call_type_filter = parse_multi_filter_values(call_type)
    call_nature_filter = parse_multi_filter_values(call_nature)
    emp_multi = parse_multi_filter_values(employee)
    emp_ids_set: set[str] = set()
    emp_norm_set: set[str] = set()
    emp_names_set: set[str] = set()
    emp_text_queries: list[str] = []

    for it in emp_multi:
        clean_it = it.strip().casefold()
        if not clean_it:
            continue
        if "·" in it:
            parts = it.split("·", 1)
            raw_id = parts[0].strip()
            raw_name = parts[1].strip()
            emp_ids_set.add(raw_id.casefold())
            norm = normalize_activity_employee_id(raw_id)
            if norm:
                emp_norm_set.add(norm)
            emp_names_set.add(raw_name.casefold())
        else:
            emp_ids_set.add(clean_it)
            norm = normalize_activity_employee_id(it)
            if norm:
                emp_norm_set.add(norm)
            emp_text_queries.append(clean_it)

    def matches_common_filters(row_team: str, row_manager: str, row_town: str, row_ff_id: str, row_ff_name: str) -> bool:
        if team_filter and row_team.casefold() not in team_filter:
            return False
        if manager_filter and row_manager.casefold() not in manager_filter:
            return False
        if town_filter and row_town.casefold() not in town_filter:
            return False
        if emp_multi:
            row_id_clean = (row_ff_id or "").strip().casefold()
            row_norm = normalize_activity_employee_id(row_ff_id)
            row_name_clean = (row_ff_name or "").strip().casefold()
            matched = False
            if row_id_clean and row_id_clean in emp_ids_set:
                matched = True
            elif row_norm and row_norm in emp_norm_set:
                matched = True
            elif row_name_clean and row_name_clean in emp_names_set:
                matched = True
            elif emp_text_queries:
                haystack = f"{row_ff_id} {row_norm} {row_ff_name}".casefold()
                if any(q in haystack for q in emp_text_queries):
                    matched = True
            if not matched:
                return False
        return True

    # Aggregate structures
    emp_stats: dict[str, dict[str, Any]] = {}

    def get_emp_entry(emp_key: str, ff_id: str, ff_name: str, t_name: str, mgr: str, town: str) -> dict[str, Any]:
        entry = emp_stats.get(emp_key)
        if entry is None:
            meta = employees_directory.get(emp_key, {})
            entry = {
                "employeeKey": emp_key,
                "ff_id": ff_id or meta.get("ff_id", emp_key),
                "ff_name": ff_name or meta.get("ff_name", "-"),
                "team": t_name or meta.get("team", "-"),
                "manager": mgr or meta.get("manager", "-"),
                "base_town": town or meta.get("base_town", "-"),
                "dailyCalls": 0,
                "yesterdayCalls": 0,
                "weeklyCalls": 0,
                "monthlyCalls": 0,
                "dailyAtt": 0,
                "weeklyAttDays": set(),
                "monthlyAttDays": set(),
                "dailyNoCalls": 0,
                "weeklyNoCallDays": set(),
                "monthlyNoCallDays": set(),
                "dailyLcdCalls": 0,
                "weeklyLcdCalls": 0,
                "monthlyLcdCalls": 0,
                "dailyProducts": 0,
                "weeklyProducts": 0,
                "monthlyProducts": 0,
                "dailyLeave": 0,
                "weeklyLeaveDays": set(),
                "monthlyLeaveDays": set(),
                "dailyOtherExc": 0,
                "weeklyOtherExc": 0,
                "monthlyOtherExc": 0,
                "doctorsSet": set(),
                "inRangeCalls": 0,
                "totalGeoCalls": 0,
                "activeCallDaysWeek": set(),
                "activeCallDaysMonth": set(),
                "weekBuckets": {"W1": 0, "W2": 0, "W3": 0, "W4": 0, "W5": 0},
                "dailySeries": {},
                "dailyLcdSeries": {},
                "dailyAttSeries": {},
            }
            emp_stats[emp_key] = entry
        else:
            if (not entry["ff_name"] or entry["ff_name"] == "-") and ff_name:
                entry["ff_name"] = ff_name
            if (not entry["team"] or entry["team"] == "-") and t_name:
                entry["team"] = t_name
            if (not entry["manager"] or entry["manager"] == "-") and mgr:
                entry["manager"] = mgr
            if (not entry["base_town"] or entry["base_town"] == "-") and town:
                entry["base_town"] = town
        return entry

    # Trend buckets by date
    daily_trend_map: dict[date, dict[str, Any]] = {}

    def get_daily_bucket(dt: date) -> dict[str, Any]:
        b = daily_trend_map.get(dt)
        if b is None:
            b = {
                "date": dt.isoformat(),
                "label": dt.strftime("%d %b"),
                "callsDone": 0,
                "attendanceCount": 0,
                "noCallsCount": 0,
                "activeCallEmpSet": set(),
                "activeAttEmpSet": set(),
                "lcdProviderCalls": 0,
                "lcdAttendance": 0,
                "lcdLeave": 0,
                "lcdOtherExc": 0,
                "lcdProducts": 0,
                "lcdDoctorsSet": set(),
            }
            daily_trend_map[dt] = b
        return b

    # Team, Manager, Town aggregations
    team_agg: dict[str, dict[str, Any]] = {}
    manager_agg: dict[str, dict[str, Any]] = {}
    town_agg: dict[str, dict[str, Any]] = {}

    def get_group_bucket(store: dict[str, dict[str, Any]], key: str, extra_team: str = "") -> dict[str, Any]:
        clean_key = key or "Unassigned"
        item = store.get(clean_key)
        if item is None:
            item = {
                "name": clean_key,
                "team": extra_team,
                "dailyCalls": 0,
                "weeklyCalls": 0,
                "monthlyCalls": 0,
                "dailyAttSet": set(),
                "weeklyAttSet": set(),
                "monthlyAttSet": set(),
                "dailyEmpSet": set(),
                "weeklyEmpSet": set(),
                "monthlyEmpSet": set(),
                "lcdDailyCalls": 0,
                "lcdWeeklyCalls": 0,
                "lcdMonthlyCalls": 0,
                "dailyProducts": 0,
                "weeklyProducts": 0,
                "monthlyProducts": 0,
                "leaveCount": 0,
                "noCallsCount": 0,
            }
            store[clean_key] = item
        return item

    # Pie chart counters for daily, weekly, monthly
    call_type_pie = {
        "daily": {"CALL": {"rows": 0, "done": 0}, "EXCEPTION": {"rows": 0, "done": 0}, "NO CALLS": {"rows": 0, "done": 0}},
        "weekly": {"CALL": {"rows": 0, "done": 0}, "EXCEPTION": {"rows": 0, "done": 0}, "NO CALLS": {"rows": 0, "done": 0}},
        "monthly": {"CALL": {"rows": 0, "done": 0}, "EXCEPTION": {"rows": 0, "done": 0}, "NO CALLS": {"rows": 0, "done": 0}},
    }
    call_nature_pie = {
        "daily": {"Provider Calls": 0, "Attendance": 0, "LEAVE": 0, "Others Exc Calls": 0},
        "weekly": {"Provider Calls": 0, "Attendance": 0, "LEAVE": 0, "Others Exc Calls": 0},
        "monthly": {"Provider Calls": 0, "Attendance": 0, "LEAVE": 0, "Others Exc Calls": 0},
    }
    geometer_pie = {
        "daily": {},
        "weekly": {},
        "monthly": {},
    }

    # KPI accumulators
    yesterday_calls_total = 0
    prev_weekly_calls_total = 0
    ca_filtered_rows_count = 0
    lcd_filtered_rows_count = 0

    # Process call_attendence rows
    for r in ca_rows:
        dt = r["call_date"]
        if isinstance(dt, datetime):
            dt = dt.date()
        if not isinstance(dt, date):
            continue
        if not matches_common_filters(r["team"], r["manager"], r["base_town"], r["ff_id"], r["ff_name"]):
            continue
        c_type = r["call_type"] or "OTHER"
        if call_type_filter and c_type.casefold() not in call_type_filter:
            continue

        emp_key = normalize_activity_employee_id(r["ff_id"]) or r["ff_name"] or "unknown"
        done = int(r["total_sum_of_done"] or 0)
        in_daily = (dt == daily_date)
        in_yesterday = (dt == yesterday_date)
        in_weekly = (week_start_date <= dt <= daily_date)
        in_prev_weekly = (prev_week_start <= dt <= prev_week_end)
        in_monthly = (month_start_date <= dt <= daily_date)
        in_range = (range_start <= dt <= range_end)

        if in_range:
            ca_filtered_rows_count += 1

        emp = get_emp_entry(emp_key, r["ff_id"], r["ff_name"], r["team"], r["manager"], r["base_town"])
        t_bucket = get_group_bucket(team_agg, r["team"])
        m_bucket = get_group_bucket(manager_agg, r["manager"], r["team"])
        town_bucket = get_group_bucket(town_agg, r["base_town"])

        if in_daily:
            call_type_pie["daily"].setdefault(c_type, {"rows": 0, "done": 0})
            call_type_pie["daily"][c_type]["rows"] += 1
            call_type_pie["daily"][c_type]["done"] += max(done, 1 if c_type == "NO CALLS" else 0)
        if in_weekly:
            call_type_pie["weekly"].setdefault(c_type, {"rows": 0, "done": 0})
            call_type_pie["weekly"][c_type]["rows"] += 1
            call_type_pie["weekly"][c_type]["done"] += max(done, 1 if c_type == "NO CALLS" else 0)
        if in_monthly:
            call_type_pie["monthly"].setdefault(c_type, {"rows": 0, "done": 0})
            call_type_pie["monthly"][c_type]["rows"] += 1
            call_type_pie["monthly"][c_type]["done"] += max(done, 1 if c_type == "NO CALLS" else 0)

        if in_range:
            d_bucket = get_daily_bucket(dt)
            if c_type == "CALL" and done > 0:
                d_bucket["callsDone"] += done
                d_bucket["activeCallEmpSet"].add(emp_key)
            elif c_type == "EXCEPTION" and done > 0:
                d_bucket["attendanceCount"] += done
                d_bucket["activeAttEmpSet"].add(emp_key)
            elif c_type == "NO CALLS":
                d_bucket["noCallsCount"] += 1

        if c_type == "CALL" and done > 0:
            iso_dt = dt.isoformat()
            emp["dailySeries"][iso_dt] = emp["dailySeries"].get(iso_dt, 0) + done
            if in_daily:
                emp["dailyCalls"] += done
                t_bucket["dailyCalls"] += done
                m_bucket["dailyCalls"] += done
                town_bucket["dailyCalls"] += done
                t_bucket["dailyEmpSet"].add(emp_key)
                m_bucket["dailyEmpSet"].add(emp_key)
                town_bucket["dailyEmpSet"].add(emp_key)
            if in_yesterday:
                emp["yesterdayCalls"] += done
                yesterday_calls_total += done
            if in_weekly:
                emp["weeklyCalls"] += done
                emp["activeCallDaysWeek"].add(dt)
                t_bucket["weeklyCalls"] += done
                m_bucket["weeklyCalls"] += done
                town_bucket["weeklyCalls"] += done
                t_bucket["weeklyEmpSet"].add(emp_key)
                m_bucket["weeklyEmpSet"].add(emp_key)
                town_bucket["weeklyEmpSet"].add(emp_key)
            if in_prev_weekly:
                prev_weekly_calls_total += done
            if in_monthly:
                emp["monthlyCalls"] += done
                emp["activeCallDaysMonth"].add(dt)
                w_idx = f"W{min((dt.day - 1) // 7 + 1, 5)}"
                emp["weekBuckets"][w_idx] = emp["weekBuckets"].get(w_idx, 0) + done
                t_bucket["monthlyCalls"] += done
                m_bucket["monthlyCalls"] += done
                town_bucket["monthlyCalls"] += done
                t_bucket["monthlyEmpSet"].add(emp_key)
                m_bucket["monthlyEmpSet"].add(emp_key)
                town_bucket["monthlyEmpSet"].add(emp_key)
        elif c_type == "EXCEPTION" and done > 0:
            iso_dt = dt.isoformat()
            emp["dailyAttSeries"][iso_dt] = 1
            if in_daily:
                emp["dailyAtt"] = 1
                t_bucket["dailyAttSet"].add(emp_key)
                m_bucket["dailyAttSet"].add(emp_key)
                town_bucket["dailyAttSet"].add(emp_key)
            if in_weekly:
                emp["weeklyAttDays"].add(dt)
                t_bucket["weeklyAttSet"].add((emp_key, dt))
                m_bucket["weeklyAttSet"].add((emp_key, dt))
                town_bucket["weeklyAttSet"].add((emp_key, dt))
            if in_monthly:
                emp["monthlyAttDays"].add(dt)
                t_bucket["monthlyAttSet"].add((emp_key, dt))
                m_bucket["monthlyAttSet"].add((emp_key, dt))
                town_bucket["monthlyAttSet"].add((emp_key, dt))
        elif c_type == "NO CALLS":
            if in_daily:
                emp["dailyNoCalls"] += 1
            if in_weekly:
                emp["weeklyNoCallDays"].add(dt)
            if in_monthly:
                emp["monthlyNoCallDays"].add(dt)
                t_bucket["noCallsCount"] += 1
                m_bucket["noCallsCount"] += 1
                town_bucket["noCallsCount"] += 1

    # Process live_call_details rows
    for r in lcd_rows:
        dt = r["call_date"]
        if isinstance(dt, datetime):
            dt = dt.date()
        if not isinstance(dt, date):
            continue
        if not matches_common_filters(r["team"], r["manager"], r["base_town"], r["ff_id"], r["ff_name"]):
            continue
        c_type = r["calltype"] or "OTHER"
        c_nature = r["call_nature"] or "Other"
        if call_type_filter and c_type.casefold() not in call_type_filter:
            continue
        if call_nature_filter and c_nature.casefold() not in call_nature_filter:
            continue

        emp_key = normalize_activity_employee_id(r["ff_id"]) or r["ff_name"] or "unknown"
        noofcalls = int(r["total_sum_noofcalls"] or 0)
        noofproducts = int(r["total_sum_noofproducts"] or 0)
        geo = r["geometer_range"] or "Unspecified"
        doc = r["doctorname"]

        in_daily = (dt == lcd_daily_date)
        in_weekly = (week_start_date <= dt <= daily_date)
        in_monthly = (month_start_date <= dt <= daily_date)
        in_range = (range_start <= dt <= range_end)

        if in_range:
            lcd_filtered_rows_count += 1

        emp = get_emp_entry(emp_key, r["ff_id"], r["ff_name"], r["team"], r["manager"], r["base_town"])
        t_bucket = get_group_bucket(team_agg, r["team"])
        m_bucket = get_group_bucket(manager_agg, r["manager"], r["team"])
        town_bucket = get_group_bucket(town_agg, r["base_town"])

        if in_daily:
            call_nature_pie["daily"][c_nature] = call_nature_pie["daily"].get(c_nature, 0) + max(noofcalls, 1)
            geometer_pie["daily"][geo] = geometer_pie["daily"].get(geo, 0) + 1
        if in_weekly:
            call_nature_pie["weekly"][c_nature] = call_nature_pie["weekly"].get(c_nature, 0) + max(noofcalls, 1)
            geometer_pie["weekly"][geo] = geometer_pie["weekly"].get(geo, 0) + 1
        if in_monthly:
            call_nature_pie["monthly"][c_nature] = call_nature_pie["monthly"].get(c_nature, 0) + max(noofcalls, 1)
            geometer_pie["monthly"][geo] = geometer_pie["monthly"].get(geo, 0) + 1

        if in_range:
            d_bucket = get_daily_bucket(dt)
            d_bucket["lcdProducts"] += noofproducts
            if c_nature == "Provider Calls":
                d_bucket["lcdProviderCalls"] += max(noofcalls, 1)
                if doc:
                    d_bucket["lcdDoctorsSet"].add(doc.casefold())
            elif c_nature == "Attendance":
                d_bucket["lcdAttendance"] += max(noofcalls, 1)
            elif c_nature == "LEAVE":
                d_bucket["lcdLeave"] += max(noofcalls, 1)
            else:
                d_bucket["lcdOtherExc"] += max(noofcalls, 1)

        if in_monthly:
            emp["totalGeoCalls"] += 1
            if geo.casefold() == "in range":
                emp["inRangeCalls"] += 1

        if c_nature == "Provider Calls":
            call_val = max(noofcalls, 1)
            iso_dt = dt.isoformat()
            emp["dailyLcdSeries"][iso_dt] = emp["dailyLcdSeries"].get(iso_dt, 0) + call_val
            if doc:
                emp["doctorsSet"].add(doc.casefold())
            if in_daily:
                emp["dailyLcdCalls"] += call_val
                emp["dailyProducts"] += noofproducts
                t_bucket["lcdDailyCalls"] += call_val
                m_bucket["lcdDailyCalls"] += call_val
                town_bucket["lcdDailyCalls"] += call_val
                t_bucket["dailyProducts"] += noofproducts
                m_bucket["dailyProducts"] += noofproducts
                town_bucket["dailyProducts"] += noofproducts
            if in_weekly:
                emp["weeklyLcdCalls"] += call_val
                emp["weeklyProducts"] += noofproducts
                t_bucket["lcdWeeklyCalls"] += call_val
                m_bucket["lcdWeeklyCalls"] += call_val
                town_bucket["lcdWeeklyCalls"] += call_val
                t_bucket["weeklyProducts"] += noofproducts
                m_bucket["weeklyProducts"] += noofproducts
                town_bucket["weeklyProducts"] += noofproducts
            if in_monthly:
                emp["monthlyLcdCalls"] += call_val
                emp["monthlyProducts"] += noofproducts
                t_bucket["lcdMonthlyCalls"] += call_val
                m_bucket["lcdMonthlyCalls"] += call_val
                town_bucket["lcdMonthlyCalls"] += call_val
                t_bucket["monthlyProducts"] += noofproducts
                m_bucket["monthlyProducts"] += noofproducts
                town_bucket["monthlyProducts"] += noofproducts
        elif c_nature == "LEAVE":
            if in_daily:
                emp["dailyLeave"] += 1
            if in_weekly:
                emp["weeklyLeaveDays"].add(dt)
            if in_monthly:
                emp["monthlyLeaveDays"].add(dt)
                t_bucket["leaveCount"] += 1
                m_bucket["leaveCount"] += 1
                town_bucket["leaveCount"] += 1
        elif c_nature == "Others Exc Calls":
            if in_daily:
                emp["dailyOtherExc"] += max(noofcalls, 1)
            if in_weekly:
                emp["weeklyOtherExc"] += max(noofcalls, 1)
            if in_monthly:
                emp["monthlyOtherExc"] += max(noofcalls, 1)

    # Finalize Daily Trend series across the date range
    daily_trend_list: list[dict[str, Any]] = []
    weekly_buckets_map: dict[str, dict[str, Any]] = {}
    monthly_buckets_map: dict[str, dict[str, Any]] = {}

    cur_dt = range_start
    cumulative_calls = 0
    while cur_dt <= range_end:
        b = daily_trend_map.get(cur_dt)
        calls_done = b["callsDone"] if b else 0
        att_cnt = b["attendanceCount"] if b else 0
        no_calls_cnt = b["noCallsCount"] if b else 0
        active_emp_cnt = len(b["activeCallEmpSet"] | b["activeAttEmpSet"]) if b else 0
        active_call_emp_cnt = len(b["activeCallEmpSet"]) if b else 0
        lcd_prov = b["lcdProviderCalls"] if b else 0
        lcd_att = b["lcdAttendance"] if b else 0
        lcd_leave = b["lcdLeave"] if b else 0
        lcd_other = b["lcdOtherExc"] if b else 0
        lcd_prod = b["lcdProducts"] if b else 0
        cumulative_calls += calls_done

        daily_trend_list.append({
            "date": cur_dt.isoformat(),
            "label": cur_dt.strftime("%d %b"),
            "weekday": cur_dt.strftime("%a"),
            "callsDone": calls_done,
            "attendanceCount": att_cnt,
            "noCallsCount": no_calls_cnt,
            "activeEmployees": active_emp_cnt,
            "activeCallEmployees": active_call_emp_cnt,
            "avgCallsPerEmp": round(calls_done / active_call_emp_cnt, 1) if active_call_emp_cnt else 0,
            "lcdProviderCalls": lcd_prov,
            "lcdAttendance": lcd_att,
            "lcdLeave": lcd_leave,
            "lcdOtherExc": lcd_other,
            "lcdProducts": lcd_prod,
            "cumulativeCalls": cumulative_calls,
        })

        w_key, w_label, w_start, w_end = _week_of_month_bucket(cur_dt)
        wb = weekly_buckets_map.get(w_key)
        if wb is None:
            wb = {
                "weekKey": w_key,
                "label": w_label,
                "startDate": w_start.isoformat(),
                "endDate": w_end.isoformat(),
                "callsDone": 0,
                "attendanceCount": 0,
                "noCallsCount": 0,
                "empSet": set(),
                "callEmpSet": set(),
                "lcdProviderCalls": 0,
                "lcdAttendance": 0,
                "lcdLeave": 0,
                "lcdOtherExc": 0,
                "lcdProducts": 0,
            }
            weekly_buckets_map[w_key] = wb
        wb["callsDone"] += calls_done
        wb["attendanceCount"] += att_cnt
        wb["noCallsCount"] += no_calls_cnt
        wb["lcdProviderCalls"] += lcd_prov
        wb["lcdAttendance"] += lcd_att
        wb["lcdLeave"] += lcd_leave
        wb["lcdOtherExc"] += lcd_other
        wb["lcdProducts"] += lcd_prod
        if b:
            wb["empSet"].update(b["activeCallEmpSet"] | b["activeAttEmpSet"])
            wb["callEmpSet"].update(b["activeCallEmpSet"])

        m_key = cur_dt.strftime("%Y-%m")
        m_label = cur_dt.strftime("%b %Y")
        mb = monthly_buckets_map.get(m_key)
        if mb is None:
            mb = {
                "monthKey": m_key,
                "label": m_label,
                "callsDone": 0,
                "attendanceCount": 0,
                "noCallsCount": 0,
                "empSet": set(),
                "callEmpSet": set(),
                "lcdProviderCalls": 0,
                "lcdAttendance": 0,
                "lcdLeave": 0,
                "lcdOtherExc": 0,
                "lcdProducts": 0,
            }
            monthly_buckets_map[m_key] = mb
        mb["callsDone"] += calls_done
        mb["attendanceCount"] += att_cnt
        mb["noCallsCount"] += no_calls_cnt
        mb["lcdProviderCalls"] += lcd_prov
        mb["lcdAttendance"] += lcd_att
        mb["lcdLeave"] += lcd_leave
        mb["lcdOtherExc"] += lcd_other
        mb["lcdProducts"] += lcd_prod
        if b:
            mb["empSet"].update(b["activeCallEmpSet"] | b["activeAttEmpSet"])
            mb["callEmpSet"].update(b["activeCallEmpSet"])

        cur_dt += timedelta(days=1)

    weekly_trend_list = []
    for wb in weekly_buckets_map.values():
        c_emps = len(wb["callEmpSet"])
        weekly_trend_list.append({
            "weekKey": wb["weekKey"],
            "label": wb["label"],
            "startDate": wb["startDate"],
            "endDate": wb["endDate"],
            "callsDone": wb["callsDone"],
            "attendanceCount": wb["attendanceCount"],
            "noCallsCount": wb["noCallsCount"],
            "activeEmployees": len(wb["empSet"]),
            "activeCallEmployees": c_emps,
            "avgCallsPerEmp": round(wb["callsDone"] / c_emps, 1) if c_emps else 0,
            "lcdProviderCalls": wb["lcdProviderCalls"],
            "lcdAttendance": wb["lcdAttendance"],
            "lcdLeave": wb["lcdLeave"],
            "lcdOtherExc": wb["lcdOtherExc"],
            "lcdProducts": wb["lcdProducts"],
        })

    monthly_trend_list = []
    for mb in monthly_buckets_map.values():
        c_emps = len(mb["callEmpSet"])
        monthly_trend_list.append({
            "monthKey": mb["monthKey"],
            "label": mb["label"],
            "callsDone": mb["callsDone"],
            "attendanceCount": mb["attendanceCount"],
            "noCallsCount": mb["noCallsCount"],
            "activeEmployees": len(mb["empSet"]),
            "activeCallEmployees": c_emps,
            "avgCallsPerEmp": round(mb["callsDone"] / c_emps, 1) if c_emps else 0,
            "lcdProviderCalls": mb["lcdProviderCalls"],
            "lcdAttendance": mb["lcdAttendance"],
            "lcdLeave": mb["lcdLeave"],
            "lcdOtherExc": mb["lcdOtherExc"],
            "lcdProducts": mb["lcdProducts"],
        })

    # Finalize employee performance list
    employee_list: list[dict[str, Any]] = []
    for emp in emp_stats.values():
        w_att = len(emp["weeklyAttDays"])
        m_att = len(emp["monthlyAttDays"])
        w_call_days = len(emp["activeCallDaysWeek"])
        m_call_days = len(emp["activeCallDaysMonth"])
        w_denom = max(w_att, w_call_days, 1)
        m_denom = max(m_att, m_call_days, 1)
        geo_pct = round((emp["inRangeCalls"] / emp["totalGeoCalls"]) * 100, 1) if emp["totalGeoCalls"] else 100.0

        employee_list.append({
            "employeeKey": emp["employeeKey"],
            "ff_id": emp["ff_id"],
            "ff_name": emp["ff_name"],
            "team": emp["team"],
            "manager": emp["manager"],
            "base_town": emp["base_town"],
            "dailyCalls": emp["dailyCalls"],
            "yesterdayCalls": emp["yesterdayCalls"],
            "weeklyCalls": emp["weeklyCalls"],
            "monthlyCalls": emp["monthlyCalls"],
            "dailyAtt": emp["dailyAtt"],
            "weeklyAttDays": w_att,
            "monthlyAttDays": m_att,
            "dailyNoCalls": emp["dailyNoCalls"],
            "weeklyNoCallDays": len(emp["weeklyNoCallDays"]),
            "monthlyNoCallDays": len(emp["monthlyNoCallDays"]),
            "dailyLcdCalls": emp["dailyLcdCalls"],
            "weeklyLcdCalls": emp["weeklyLcdCalls"],
            "monthlyLcdCalls": emp["monthlyLcdCalls"],
            "dailyProducts": emp["dailyProducts"],
            "weeklyProducts": emp["weeklyProducts"],
            "monthlyProducts": emp["monthlyProducts"],
            "dailyLeave": emp["dailyLeave"],
            "weeklyLeaveDays": len(emp["weeklyLeaveDays"]),
            "monthlyLeaveDays": len(emp["monthlyLeaveDays"]),
            "dailyOtherExc": emp["dailyOtherExc"],
            "weeklyOtherExc": emp["weeklyOtherExc"],
            "monthlyOtherExc": emp["monthlyOtherExc"],
            "uniqueDoctors": len(emp["doctorsSet"]),
            "inRangePct": geo_pct,
            "weeklyAvgCalls": round(emp["weeklyCalls"] / w_denom, 1) if emp["weeklyCalls"] else 0.0,
            "monthlyAvgCalls": round(emp["monthlyCalls"] / m_denom, 1) if emp["monthlyCalls"] else 0.0,
            "weekBuckets": emp["weekBuckets"],
            "dailySeries": emp["dailySeries"],
            "dailyLcdSeries": emp["dailyLcdSeries"],
            "dailyAttSeries": emp["dailyAttSeries"],
        })

    employee_list.sort(
        key=lambda item: (item["monthlyCalls"], item["weeklyCalls"], item["dailyCalls"], item["monthlyLcdCalls"]),
        reverse=True,
    )

    def finalize_group_list(store: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
        out = []
        for g in store.values():
            m_emps = len(g["monthlyEmpSet"])
            w_emps = len(g["weeklyEmpSet"])
            d_emps = len(g["dailyEmpSet"])
            out.append({
                "name": g["name"],
                "team": g["team"],
                "dailyCalls": g["dailyCalls"],
                "weeklyCalls": g["weeklyCalls"],
                "monthlyCalls": g["monthlyCalls"],
                "dailyAtt": len(g["dailyAttSet"]),
                "weeklyAtt": len(g["weeklyAttSet"]),
                "monthlyAtt": len(g["monthlyAttSet"]),
                "dailyActiveEmployees": d_emps,
                "weeklyActiveEmployees": w_emps,
                "monthlyActiveEmployees": m_emps,
                "lcdDailyCalls": g["lcdDailyCalls"],
                "lcdWeeklyCalls": g["lcdWeeklyCalls"],
                "lcdMonthlyCalls": g["lcdMonthlyCalls"],
                "dailyProducts": g["dailyProducts"],
                "weeklyProducts": g["weeklyProducts"],
                "monthlyProducts": g["monthlyProducts"],
                "leaveCount": g["leaveCount"],
                "noCallsCount": g["noCallsCount"],
                "avgMonthlyCallsPerEmp": round(g["monthlyCalls"] / m_emps, 1) if m_emps else 0.0,
            })
        out.sort(key=lambda x: (x["monthlyCalls"], x["lcdMonthlyCalls"]), reverse=True)
        return out

    team_performance = finalize_group_list(team_agg)
    manager_performance = finalize_group_list(manager_agg)
    town_performance = finalize_group_list(town_agg)

    # Compute KPI totals across Daily, Weekly, Monthly
    total_daily_calls = sum(e["dailyCalls"] for e in employee_list)
    total_weekly_calls = sum(e["weeklyCalls"] for e in employee_list)
    total_monthly_calls = sum(e["monthlyCalls"] for e in employee_list)

    active_daily_emps = sum(1 for e in employee_list if e["dailyCalls"] > 0 or e["dailyAtt"] > 0 or e["dailyLcdCalls"] > 0)
    active_weekly_emps = sum(1 for e in employee_list if e["weeklyCalls"] > 0 or e["weeklyAttDays"] > 0 or e["weeklyLcdCalls"] > 0)
    active_monthly_emps = sum(1 for e in employee_list if e["monthlyCalls"] > 0 or e["monthlyAttDays"] > 0 or e["monthlyLcdCalls"] > 0)

    call_daily_emps = sum(1 for e in employee_list if e["dailyCalls"] > 0)
    call_weekly_emps = sum(1 for e in employee_list if e["weeklyCalls"] > 0)
    call_monthly_emps = sum(1 for e in employee_list if e["monthlyCalls"] > 0)

    daily_att_count = sum(e["dailyAtt"] for e in employee_list)
    weekly_att_days = sum(e["weeklyAttDays"] for e in employee_list)
    monthly_att_days = sum(e["monthlyAttDays"] for e in employee_list)

    lcd_daily_calls = sum(e["dailyLcdCalls"] for e in employee_list)
    lcd_weekly_calls = sum(e["weeklyLcdCalls"] for e in employee_list)
    lcd_monthly_calls = sum(e["monthlyLcdCalls"] for e in employee_list)

    lcd_daily_products = sum(e["dailyProducts"] for e in employee_list)
    lcd_weekly_products = sum(e["weeklyProducts"] for e in employee_list)
    lcd_monthly_products = sum(e["monthlyProducts"] for e in employee_list)

    daily_no_calls = sum(e["dailyNoCalls"] for e in employee_list)
    weekly_no_calls = sum(e["weeklyNoCallDays"] for e in employee_list)
    monthly_no_calls = sum(e["monthlyNoCallDays"] for e in employee_list)

    daily_leaves = sum(e["dailyLeave"] for e in employee_list)
    weekly_leaves = sum(e["weeklyLeaveDays"] for e in employee_list)
    monthly_leaves = sum(e["monthlyLeaveDays"] for e in employee_list)

    daily_other_exc = sum(e["dailyOtherExc"] for e in employee_list)
    weekly_other_exc = sum(e["weeklyOtherExc"] for e in employee_list)
    monthly_other_exc = sum(e["monthlyOtherExc"] for e in employee_list)

    in_range_total = sum(geometer_pie["monthly"].get(k, 0) for k in geometer_pie["monthly"] if k.casefold() == "in range")
    geo_total = sum(geometer_pie["monthly"].values())

    kpis = {
        "totalEmployees": len(employee_list),
        "caTotalRows": ca_filtered_rows_count,
        "lcdTotalRows": lcd_filtered_rows_count,
        "daily": {
            "label": daily_date.strftime("%d %b %Y"),
            "callsDone": total_daily_calls,
            "prevCallsDone": yesterday_calls_total,
            "activeEmployees": active_daily_emps,
            "callEmployees": call_daily_emps,
            "attendanceCount": daily_att_count,
            "noCallsCount": daily_no_calls,
            "lcdProviderCalls": lcd_daily_calls,
            "lcdProducts": lcd_daily_products,
            "leaveCount": daily_leaves,
            "otherExcCount": daily_other_exc,
            "avgCallsPerEmp": round(total_daily_calls / call_daily_emps, 1) if call_daily_emps else 0.0,
        },
        "weekly": {
            "label": f"{week_start_date.strftime('%d %b')} – {daily_date.strftime('%d %b %Y')}",
            "callsDone": total_weekly_calls,
            "prevCallsDone": prev_weekly_calls_total,
            "activeEmployees": active_weekly_emps,
            "callEmployees": call_weekly_emps,
            "attendanceCount": weekly_att_days,
            "noCallsCount": weekly_no_calls,
            "lcdProviderCalls": lcd_weekly_calls,
            "lcdProducts": lcd_weekly_products,
            "leaveCount": weekly_leaves,
            "otherExcCount": weekly_other_exc,
            "avgCallsPerEmp": round(total_weekly_calls / call_weekly_emps, 1) if call_weekly_emps else 0.0,
        },
        "monthly": {
            "label": f"{month_start_date.strftime('%d %b')} – {daily_date.strftime('%d %b %Y')}",
            "callsDone": total_monthly_calls,
            "prevCallsDone": 0,
            "activeEmployees": active_monthly_emps,
            "callEmployees": call_monthly_emps,
            "attendanceCount": monthly_att_days,
            "noCallsCount": monthly_no_calls,
            "lcdProviderCalls": lcd_monthly_calls,
            "lcdProducts": lcd_monthly_products,
            "leaveCount": monthly_leaves,
            "otherExcCount": monthly_other_exc,
            "avgCallsPerEmp": round(total_monthly_calls / call_monthly_emps, 1) if call_monthly_emps else 0.0,
            "inRangePct": round((in_range_total / geo_total) * 100, 1) if geo_total else 100.0,
        },
    }

    return {
        "selectedDate": selected_date.isoformat(),
        "formattedDate": selected_date.strftime("%d-%b-%Y").upper(),
        "dateFrom": range_start.isoformat(),
        "dateTo": range_end.isoformat(),
        "minDate": min_available_date.isoformat(),
        "maxDate": max_available_date.isoformat(),
        "filters": {
            "teams": sorted(teams_set, key=str.casefold),
            "managers": sorted(managers_set, key=str.casefold),
            "baseTowns": sorted(towns_set, key=str.casefold),
            "callTypes": sorted(call_types_set, key=str.casefold),
            "callNatures": sorted(call_natures_set, key=str.casefold),
            "employees": sorted(
                employees_directory.values(),
                key=lambda item: (item["ff_name"].casefold(), item["ff_id"].casefold()),
            ),
        },
        "kpis": kpis,
        "charts": {
            "dailyTrend": daily_trend_list,
            "weeklyTrend": weekly_trend_list,
            "monthlyTrend": monthly_trend_list,
            "callTypePie": call_type_pie,
            "callNaturePie": call_nature_pie,
            "geometerPie": geometer_pie,
            "teamPerformance": team_performance,
            "managerPerformance": manager_performance,
            "townPerformance": town_performance,
        },
        "employees": employee_list,
    }


def _query_call_attendence_table(
    page: int = 1,
    page_size: int = 25,
    search: str = "",
    team: str = "",
    call_type: str = "",
    manager: str = "",
    base_town: str = "",
    employee: str = "",
    date_from: str | None = None,
    date_to: str | None = None,
    sort_by: str = "call_date",
    sort_dir: str = "desc",
    export_all: bool = False,
) -> dict[str, Any]:
    parsed_from = _parse_iso_date_optional(date_from, "date_from")
    parsed_to = _parse_iso_date_optional(date_to, "date_to")

    allowed_columns = [
        "team",
        "ff_id",
        "ff_name",
        "manager",
        "base_town",
        "call_date",
        "call_type",
        "total_sum_of_done",
        "total_distinct_count_of_ff_id",
    ]
    allowed_set = set(allowed_columns)

    where_clauses: list[str] = ["1=1"]
    params: dict[str, Any] = {}

    if team.strip():
        teams = [t.strip().casefold() for t in team.split("|") if t.strip()]
        if teams:
            where_clauses.append("LOWER(BTRIM(COALESCE(team, ''))) = ANY(:teams)")
            params["teams"] = teams
    if call_type.strip():
        ctypes = [c.strip().casefold() for c in call_type.split("|") if c.strip()]
        if ctypes:
            where_clauses.append("LOWER(BTRIM(COALESCE(call_type, ''))) = ANY(:ctypes)")
            params["ctypes"] = ctypes
    if manager.strip():
        mgrs = [m.strip().casefold() for m in manager.split("|") if m.strip()]
        if mgrs:
            where_clauses.append("LOWER(BTRIM(COALESCE(manager, ''))) = ANY(:mgrs)")
            params["mgrs"] = mgrs
    if base_town.strip():
        towns = [t.strip().casefold() for t in base_town.split("|") if t.strip()]
        if towns:
            where_clauses.append("LOWER(BTRIM(COALESCE(base_town, ''))) = ANY(:towns)")
            params["towns"] = towns
    if employee.strip():
        emp_items = [e.strip() for e in employee.split("|") if e.strip()]
        if len(emp_items) > 1 or "·" in employee:
            emp_ids = []
            emp_norms = []
            emp_names = []
            for it in emp_items:
                if "·" in it:
                    parts = it.split("·", 1)
                    p_id = parts[0].strip()
                    emp_ids.append(p_id.casefold())
                    norm = normalize_activity_employee_id(p_id)
                    if norm:
                        emp_norms.append(norm)
                    emp_names.append(parts[1].strip().casefold())
                else:
                    emp_ids.append(it.casefold())
                    norm = normalize_activity_employee_id(it)
                    if norm:
                        emp_norms.append(norm)
            where_clauses.append("""(
                LOWER(BTRIM(COALESCE(ff_id, ''))) = ANY(:emp_ids)
                OR LTRIM(BTRIM(COALESCE(ff_id, '')), '0') = ANY(:emp_norms)
                OR LOWER(BTRIM(COALESCE(ff_name, ''))) = ANY(:emp_names)
            )""")
            params["emp_ids"] = emp_ids
            params["emp_norms"] = emp_norms
            params["emp_names"] = emp_names
        else:
            emp_q = employee.strip()
            emp_norm = normalize_activity_employee_id(emp_q)
            params["emp_like"] = f"%{emp_q}%"
            params["emp_norm"] = emp_norm
            where_clauses.append(
                "(COALESCE(ff_id, '') ILIKE :emp_like OR COALESCE(ff_name, '') ILIKE :emp_like "
                "OR LTRIM(BTRIM(COALESCE(ff_id, '')), '0') = :emp_norm)"
            )
    if parsed_from:
        where_clauses.append("call_date >= :date_from")
        params["date_from"] = parsed_from
    if parsed_to:
        where_clauses.append("call_date <= :date_to")
        params["date_to"] = parsed_to
    if search.strip():
        params["search_like"] = f"%{search.strip()}%"
        where_clauses.append("""(
            COALESCE(team, '') ILIKE :search_like
            OR COALESCE(ff_id, '') ILIKE :search_like
            OR COALESCE(ff_name, '') ILIKE :search_like
            OR COALESCE(manager, '') ILIKE :search_like
            OR COALESCE(base_town, '') ILIKE :search_like
            OR COALESCE(call_type, '') ILIKE :search_like
            OR COALESCE(CAST(call_date AS text), '') ILIKE :search_like
        )""")

    where_sql = " AND ".join(where_clauses)
    sort_col = sort_by if sort_by in allowed_set else "call_date"
    direction = "ASC" if sort_dir.strip().casefold() == "asc" else "DESC"

    with get_database_engine().connect() as connection:
        summary_row = connection.execute(
            text(f"""
                SELECT
                    COUNT(*) AS total_rows,
                    COALESCE(SUM(total_sum_of_done), 0) AS total_done,
                    COALESCE(SUM(total_sum_of_done) FILTER (WHERE UPPER(BTRIM(call_type)) = 'CALL'), 0) AS call_done,
                    COALESCE(SUM(total_sum_of_done) FILTER (WHERE UPPER(BTRIM(call_type)) = 'EXCEPTION'), 0) AS exception_done,
                    COUNT(*) FILTER (WHERE UPPER(BTRIM(call_type)) = 'NO CALLS') AS no_calls_rows,
                    COUNT(DISTINCT LTRIM(BTRIM(COALESCE(ff_id, '')), '0')) AS distinct_employees,
                    COUNT(DISTINCT BTRIM(COALESCE(team, ''))) AS distinct_teams
                FROM public.call_attendence
                WHERE {where_sql}
            """),
            params,
        ).mappings().one()

        total_rows = int(summary_row["total_rows"] or 0)
        if page < 1:
            page = 1
        if page_size not in {15, 25, 50, 100}:
            page_size = 25
        total_pages = max((total_rows + page_size - 1) // page_size, 1)
        if page > total_pages:
            page = total_pages
        offset = (page - 1) * page_size

        limit_sql = "" if export_all else "LIMIT :limit OFFSET :offset"
        query_params = {**params} if export_all else {**params, "limit": page_size, "offset": offset}
        rows_raw = connection.execute(
            text(f"""
                SELECT
                    team,
                    ff_id,
                    ff_name,
                    manager,
                    base_town,
                    call_date,
                    call_type,
                    total_sum_of_done,
                    total_distinct_count_of_ff_id
                FROM public.call_attendence
                WHERE {where_sql}
                ORDER BY {sort_col} {direction} NULLS LAST, call_date DESC NULLS LAST, ff_id ASC NULLS LAST
                {limit_sql}
            """),
            query_params,
        ).mappings().all()

    rows = [
        {
            "team": display_value(r["team"]),
            "ff_id": display_value(r["ff_id"]),
            "ff_name": display_value(r["ff_name"]),
            "manager": display_value(r["manager"]),
            "base_town": display_value(r["base_town"]),
            "call_date": display_value(r["call_date"]),
            "call_type": display_value(r["call_type"]),
            "total_sum_of_done": int(r["total_sum_of_done"] or 0),
            "total_distinct_count_of_ff_id": int(r["total_distinct_count_of_ff_id"] or 0),
        }
        for r in rows_raw
    ]

    return {
        "table": "public.call_attendence",
        "rows": rows,
        "total": total_rows,
        "page": page,
        "pageSize": page_size,
        "totalPages": total_pages,
        "sortBy": sort_col,
        "sortDir": direction.lower(),
        "summary": {
            "totalRows": total_rows,
            "totalDone": int(summary_row["total_done"] or 0),
            "callDone": int(summary_row["call_done"] or 0),
            "exceptionDone": int(summary_row["exception_done"] or 0),
            "noCallsRows": int(summary_row["no_calls_rows"] or 0),
            "distinctEmployees": int(summary_row["distinct_employees"] or 0),
            "distinctTeams": int(summary_row["distinct_teams"] or 0),
        },
    }


@app.get("/api/dashboard/call-attendence")
def get_dashboard_call_attendence(
    page: int = 1,
    page_size: int = 25,
    search: str = "",
    team: str = "",
    call_type: str = "",
    manager: str = "",
    base_town: str = "",
    employee: str = "",
    date_from: str | None = None,
    date_to: str | None = None,
    sort_by: str = "call_date",
    sort_dir: str = "desc",
) -> dict[str, Any]:
    return _query_call_attendence_table(
        page=page,
        page_size=page_size,
        search=search,
        team=team,
        call_type=call_type,
        manager=manager,
        base_town=base_town,
        employee=employee,
        date_from=date_from,
        date_to=date_to,
        sort_by=sort_by,
        sort_dir=sort_dir,
        export_all=False,
    )


@app.get("/api/dashboard/call-attendence/export")
def export_dashboard_call_attendence(
    search: str = "",
    team: str = "",
    call_type: str = "",
    manager: str = "",
    base_town: str = "",
    employee: str = "",
    date_from: str | None = None,
    date_to: str | None = None,
    sort_by: str = "call_date",
    sort_dir: str = "desc",
) -> StreamingResponse:
    result = _query_call_attendence_table(
        search=search,
        team=team,
        call_type=call_type,
        manager=manager,
        base_town=base_town,
        employee=employee,
        date_from=date_from,
        date_to=date_to,
        sort_by=sort_by,
        sort_dir=sort_dir,
        export_all=True,
    )
    frame = pd.DataFrame(result["rows"])
    csv_bytes = frame.to_csv(index=False).encode("utf-8-sig")
    return StreamingResponse(
        BytesIO(csv_bytes),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=call_attendence_export.csv"},
    )


def _query_live_call_details_table(
    page: int = 1,
    page_size: int = 25,
    search: str = "",
    team: str = "",
    calltype: str = "",
    call_nature: str = "",
    manager: str = "",
    base_town: str = "",
    geometer_range: str = "",
    employee: str = "",
    date_from: str | None = None,
    date_to: str | None = None,
    sort_by: str = "call_date",
    sort_dir: str = "desc",
    export_all: bool = False,
) -> dict[str, Any]:
    parsed_from = _parse_iso_date_optional(date_from, "date_from")
    parsed_to = _parse_iso_date_optional(date_to, "date_to")

    allowed_columns = [
        "team",
        "ff_id",
        "ff_name",
        "call_date",
        "manager",
        "base_town",
        "pmdc",
        "doctorname",
        "call_nature",
        "calltype",
        "ww",
        "lng",
        "lat",
        "total_sum_noofcalls",
        "total_sum_noofproducts",
        "geometer_range",
    ]
    allowed_set = set(allowed_columns)

    where_clauses: list[str] = ["1=1"]
    params: dict[str, Any] = {}

    if team.strip():
        teams = [t.strip().casefold() for t in team.split("|") if t.strip()]
        if teams:
            where_clauses.append("LOWER(BTRIM(COALESCE(team, ''))) = ANY(:teams)")
            params["teams"] = teams
    if calltype.strip():
        ctypes = [c.strip().casefold() for c in calltype.split("|") if c.strip()]
        if ctypes:
            where_clauses.append("LOWER(BTRIM(COALESCE(calltype, ''))) = ANY(:ctypes)")
            params["ctypes"] = ctypes
    if call_nature.strip():
        natures = [n.strip().casefold() for n in call_nature.split("|") if n.strip()]
        if natures:
            where_clauses.append("LOWER(BTRIM(COALESCE(call_nature, ''))) = ANY(:natures)")
            params["natures"] = natures
    if manager.strip():
        mgrs = [m.strip().casefold() for m in manager.split("|") if m.strip()]
        if mgrs:
            where_clauses.append("LOWER(BTRIM(COALESCE(manager, ''))) = ANY(:mgrs)")
            params["mgrs"] = mgrs
    if base_town.strip():
        towns = [t.strip().casefold() for t in base_town.split("|") if t.strip()]
        if towns:
            where_clauses.append("LOWER(BTRIM(COALESCE(base_town, ''))) = ANY(:towns)")
            params["towns"] = towns
    if geometer_range.strip():
        geos = [g.strip().casefold() for g in geometer_range.split("|") if g.strip()]
        if geos:
            where_clauses.append("LOWER(BTRIM(COALESCE(geometer_range, ''))) = ANY(:geos)")
            params["geos"] = geos
    if employee.strip():
        emp_items = [e.strip() for e in employee.split("|") if e.strip()]
        if len(emp_items) > 1 or "·" in employee:
            emp_ids = []
            emp_norms = []
            emp_names = []
            for it in emp_items:
                if "·" in it:
                    parts = it.split("·", 1)
                    p_id = parts[0].strip()
                    emp_ids.append(p_id.casefold())
                    norm = normalize_activity_employee_id(p_id)
                    if norm:
                        emp_norms.append(norm)
                    emp_names.append(parts[1].strip().casefold())
                else:
                    emp_ids.append(it.casefold())
                    norm = normalize_activity_employee_id(it)
                    if norm:
                        emp_norms.append(norm)
            where_clauses.append("""(
                LOWER(BTRIM(COALESCE(ff_id, ''))) = ANY(:emp_ids)
                OR LTRIM(BTRIM(COALESCE(ff_id, '')), '0') = ANY(:emp_norms)
                OR LOWER(BTRIM(COALESCE(ff_name, ''))) = ANY(:emp_names)
            )""")
            params["emp_ids"] = emp_ids
            params["emp_norms"] = emp_norms
            params["emp_names"] = emp_names
        else:
            emp_q = employee.strip()
            emp_norm = normalize_activity_employee_id(emp_q)
            params["emp_like"] = f"%{emp_q}%"
            params["emp_norm"] = emp_norm
            where_clauses.append(
                "(COALESCE(ff_id, '') ILIKE :emp_like OR COALESCE(ff_name, '') ILIKE :emp_like "
                "OR LTRIM(BTRIM(COALESCE(ff_id, '')), '0') = :emp_norm)"
            )
    if parsed_from:
        where_clauses.append("call_date >= :date_from")
        params["date_from"] = parsed_from
    if parsed_to:
        where_clauses.append("call_date <= :date_to")
        params["date_to"] = parsed_to
    if search.strip():
        params["search_like"] = f"%{search.strip()}%"
        where_clauses.append("""(
            COALESCE(team, '') ILIKE :search_like
            OR COALESCE(ff_id, '') ILIKE :search_like
            OR COALESCE(ff_name, '') ILIKE :search_like
            OR COALESCE(manager, '') ILIKE :search_like
            OR COALESCE(base_town, '') ILIKE :search_like
            OR COALESCE(pmdc, '') ILIKE :search_like
            OR COALESCE(doctorname, '') ILIKE :search_like
            OR COALESCE(call_nature, '') ILIKE :search_like
            OR COALESCE(calltype, '') ILIKE :search_like
            OR COALESCE(geometer_range, '') ILIKE :search_like
            OR COALESCE(CAST(call_date AS text), '') ILIKE :search_like
        )""")

    where_sql = " AND ".join(where_clauses)
    sort_col = sort_by if sort_by in allowed_set else "call_date"
    direction = "ASC" if sort_dir.strip().casefold() == "asc" else "DESC"

    with get_database_engine().connect() as connection:
        summary_row = connection.execute(
            text(f"""
                SELECT
                    COUNT(*) AS total_rows,
                    COALESCE(SUM(total_sum_noofcalls), 0) AS total_calls,
                    COALESCE(SUM(total_sum_noofproducts), 0) AS total_products,
                    COUNT(*) FILTER (WHERE BTRIM(call_nature) = 'Provider Calls') AS provider_rows,
                    COUNT(*) FILTER (WHERE BTRIM(call_nature) = 'Attendance') AS attendance_rows,
                    COUNT(*) FILTER (WHERE BTRIM(call_nature) = 'LEAVE') AS leave_rows,
                    COUNT(*) FILTER (WHERE BTRIM(call_nature) = 'Others Exc Calls') AS other_exc_rows,
                    COUNT(DISTINCT LTRIM(BTRIM(COALESCE(ff_id, '')), '0')) AS distinct_employees,
                    COUNT(DISTINCT LOWER(BTRIM(doctorname))) FILTER (WHERE BTRIM(call_nature) = 'Provider Calls' AND BTRIM(COALESCE(doctorname, '')) <> '') AS distinct_doctors,
                    COUNT(*) FILTER (WHERE LOWER(BTRIM(COALESCE(geometer_range, ''))) = 'in range') AS in_range_rows
                FROM public.live_call_details
                WHERE {where_sql}
            """),
            params,
        ).mappings().one()

        total_rows = int(summary_row["total_rows"] or 0)
        if page < 1:
            page = 1
        if page_size not in {15, 25, 50, 100}:
            page_size = 25
        total_pages = max((total_rows + page_size - 1) // page_size, 1)
        if page > total_pages:
            page = total_pages
        offset = (page - 1) * page_size

        limit_sql = "" if export_all else "LIMIT :limit OFFSET :offset"
        query_params = {**params} if export_all else {**params, "limit": page_size, "offset": offset}
        rows_raw = connection.execute(
            text(f"""
                SELECT
                    team,
                    ff_id,
                    ff_name,
                    call_date,
                    manager,
                    base_town,
                    pmdc,
                    doctorname,
                    call_nature,
                    calltype,
                    ww,
                    lng,
                    lat,
                    total_sum_noofcalls,
                    total_sum_noofproducts,
                    geometer_range
                FROM public.live_call_details
                WHERE {where_sql}
                ORDER BY {sort_col} {direction} NULLS LAST, call_date DESC NULLS LAST, ff_id ASC NULLS LAST
                {limit_sql}
            """),
            query_params,
        ).mappings().all()

    rows = [
        {
            "team": display_value(r["team"]),
            "ff_id": display_value(r["ff_id"]),
            "ff_name": display_value(r["ff_name"]),
            "call_date": display_value(r["call_date"]),
            "manager": display_value(r["manager"]),
            "base_town": display_value(r["base_town"]),
            "pmdc": display_value(r["pmdc"]),
            "doctorname": display_value(r["doctorname"]),
            "call_nature": display_value(r["call_nature"]),
            "calltype": display_value(r["calltype"]),
            "ww": display_value(r["ww"]),
            "lng": str(r["lng"]) if r["lng"] is not None else "",
            "lat": str(r["lat"]) if r["lat"] is not None else "",
            "total_sum_noofcalls": int(r["total_sum_noofcalls"] or 0),
            "total_sum_noofproducts": int(r["total_sum_noofproducts"] or 0),
            "geometer_range": display_value(r["geometer_range"]),
        }
        for r in rows_raw
    ]

    return {
        "table": "public.live_call_details",
        "rows": rows,
        "total": total_rows,
        "page": page,
        "pageSize": page_size,
        "totalPages": total_pages,
        "sortBy": sort_col,
        "sortDir": direction.lower(),
        "summary": {
            "totalRows": total_rows,
            "totalCalls": int(summary_row["total_calls"] or 0),
            "totalProducts": int(summary_row["total_products"] or 0),
            "providerRows": int(summary_row["provider_rows"] or 0),
            "attendanceRows": int(summary_row["attendance_rows"] or 0),
            "leaveRows": int(summary_row["leave_rows"] or 0),
            "otherExcRows": int(summary_row["other_exc_rows"] or 0),
            "distinctEmployees": int(summary_row["distinct_employees"] or 0),
            "distinctDoctors": int(summary_row["distinct_doctors"] or 0),
            "inRangeRows": int(summary_row["in_range_rows"] or 0),
        },
    }


@app.get("/api/dashboard/live-call-details")
def get_dashboard_live_call_details(
    page: int = 1,
    page_size: int = 25,
    search: str = "",
    team: str = "",
    calltype: str = "",
    call_nature: str = "",
    manager: str = "",
    base_town: str = "",
    geometer_range: str = "",
    employee: str = "",
    date_from: str | None = None,
    date_to: str | None = None,
    sort_by: str = "call_date",
    sort_dir: str = "desc",
) -> dict[str, Any]:
    return _query_live_call_details_table(
        page=page,
        page_size=page_size,
        search=search,
        team=team,
        calltype=calltype,
        call_nature=call_nature,
        manager=manager,
        base_town=base_town,
        geometer_range=geometer_range,
        employee=employee,
        date_from=date_from,
        date_to=date_to,
        sort_by=sort_by,
        sort_dir=sort_dir,
        export_all=False,
    )


@app.get("/api/dashboard/live-call-details/export")
def export_dashboard_live_call_details(
    search: str = "",
    team: str = "",
    calltype: str = "",
    call_nature: str = "",
    manager: str = "",
    base_town: str = "",
    geometer_range: str = "",
    employee: str = "",
    date_from: str | None = None,
    date_to: str | None = None,
    sort_by: str = "call_date",
    sort_dir: str = "desc",
) -> StreamingResponse:
    result = _query_live_call_details_table(
        search=search,
        team=team,
        calltype=calltype,
        call_nature=call_nature,
        manager=manager,
        base_town=base_town,
        geometer_range=geometer_range,
        employee=employee,
        date_from=date_from,
        date_to=date_to,
        sort_by=sort_by,
        sort_dir=sort_dir,
        export_all=True,
    )
    frame = pd.DataFrame(result["rows"])
    csv_bytes = frame.to_csv(index=False).encode("utf-8-sig")
    return StreamingResponse(
        BytesIO(csv_bytes),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=live_call_details_export.csv"},
    )


# Mount the browser application last so /api routes take priority.
app.mount("/", StaticFiles(directory=PUBLIC_DIR, html=True), name="frontend")


if __name__ == "__main__":
    import uvicorn

    uvicorn.run("main:app", host="127.0.0.1", port=3000, reload=True)
